from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import time
from typing import Any
from uuid import uuid4

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise RuntimeError("Review export requires openpyxl to be installed") from exc

from service.business_store import hydrate_compare_result_summaries, list_compare_results


PROCESSED_REVIEW_STATUSES = (
    "confirmed_high_risk",
    "needs_followup",
    "false_positive",
)
REVIEW_STATUS_LABELS = {
    "pending": "待复核",
    "confirmed_high_risk": "确认高风险",
    "needs_followup": "继续跟进",
    "false_positive": "标记误报",
    "cleared": "已清空",
}
SHEET_TITLE_BY_STATUS = {
    "confirmed_high_risk": "确认高风险",
    "needs_followup": "继续跟进",
    "false_positive": "标记误报",
}
OTHER_PROCESSED_SHEET_TITLE = "其他已处理"
REVIEW_EXPORT_FILENAME_PREFIX = "复核结果导出"


@dataclass(frozen=True)
class ReviewExportColumn:
    header: str
    width: float
    number_format: str | None = None
    wrap_text: bool = False


REVIEW_EXPORT_COLUMNS: tuple[ReviewExportColumn, ...] = (
    ReviewExportColumn("结果ID", 12),
    ReviewExportColumn("任务ID", 38),
    ReviewExportColumn("条目序号", 10),
    ReviewExportColumn("任务状态", 14),
    ReviewExportColumn("检测模式", 12),
    ReviewExportColumn("复核状态", 16),
    ReviewExportColumn("复核人", 14),
    ReviewExportColumn("复核备注", 28, wrap_text=True),
    ReviewExportColumn("复核更新时间", 22),
    ReviewExportColumn("结果创建时间", 22),
    ReviewExportColumn("结果更新时间", 22),
    ReviewExportColumn("任务创建时间", 22),
    ReviewExportColumn("任务完成时间", 22),
    ReviewExportColumn("来源文件", 22),
    ReviewExportColumn("来源标识", 24),
    ReviewExportColumn("短剧名", 22),
    ReviewExportColumn("展示标题", 30, wrap_text=True),
    ReviewExportColumn("集数", 10),
    ReviewExportColumn("作者", 16),
    ReviewExportColumn("平台", 14),
    ReviewExportColumn("源小说名", 22),
    ReviewExportColumn("Excel行号", 12),
    ReviewExportColumn("描述", 36, wrap_text=True),
    ReviewExportColumn("命中书名", 24),
    ReviewExportColumn("命中章节", 24),
    ReviewExportColumn("命中标签", 16),
    ReviewExportColumn("置信标签", 14),
    ReviewExportColumn("语义状态", 18),
    ReviewExportColumn("精排序号", 10),
    ReviewExportColumn("粗排序号", 10),
    ReviewExportColumn("精排分", 12, number_format="0.0000"),
    ReviewExportColumn("粗排分", 12, number_format="0.0000"),
    ReviewExportColumn("耗时(秒)", 12, number_format="0.000"),
    ReviewExportColumn("候选窗口序号", 12),
    ReviewExportColumn("候选起始偏移", 12),
    ReviewExportColumn("候选结束偏移", 12),
    ReviewExportColumn("最长匹配长度", 12),
    ReviewExportColumn("最长匹配比例", 12, number_format="0.0000"),
    ReviewExportColumn("N-Gram召回", 12, number_format="0.0000"),
    ReviewExportColumn("N-Gram精度", 12, number_format="0.0000"),
    ReviewExportColumn("Jaccard", 12, number_format="0.0000"),
    ReviewExportColumn("序列相似度", 12, number_format="0.0000"),
    ReviewExportColumn("精确片段命中", 12),
    ReviewExportColumn("命中片段", 40, wrap_text=True),
    ReviewExportColumn("查询文本预览", 42, wrap_text=True),
    ReviewExportColumn("候选文本预览", 42, wrap_text=True),
    ReviewExportColumn("查询文本证据", 80, wrap_text=True),
    ReviewExportColumn("候选文本证据", 90, wrap_text=True),
    ReviewExportColumn("候选章节全文", 110, wrap_text=True),
)


def build_review_export_xlsx(
    *,
    business_db_path: str | Path,
    retrieval_db_path: str | Path,
    export_root: str | Path,
    owner_user_id: int,
    task_id: str = "",
    item_status: str = "",
    review_status: str = "",
    sort_by: str = "updated_at_desc",
    dedupe_latest: bool = False,
    text_filter: str = "",
    processed_only: bool = True,
    candidate_score_threshold: float | None = None,
) -> tuple[Path, str]:
    file_path, export_filename, _ = build_review_export_xlsx_with_metrics(
        business_db_path=business_db_path,
        retrieval_db_path=retrieval_db_path,
        export_root=export_root,
        owner_user_id=owner_user_id,
        task_id=task_id,
        item_status=item_status,
        review_status=review_status,
        sort_by=sort_by,
        dedupe_latest=dedupe_latest,
        text_filter=text_filter,
        processed_only=processed_only,
        candidate_score_threshold=candidate_score_threshold,
    )
    return file_path, export_filename


def build_review_export_xlsx_with_metrics(
    *,
    business_db_path: str | Path,
    retrieval_db_path: str | Path,
    export_root: str | Path,
    owner_user_id: int,
    task_id: str = "",
    item_status: str = "",
    review_status: str = "",
    sort_by: str = "updated_at_desc",
    dedupe_latest: bool = False,
    text_filter: str = "",
    processed_only: bool = True,
    candidate_score_threshold: float | None = None,
) -> tuple[Path, str, dict[str, Any]]:
    total_started = time.perf_counter()

    stage_started = time.perf_counter()
    summaries = _load_result_summaries(
        business_db_path=business_db_path,
        owner_user_id=owner_user_id,
        task_id=task_id,
        item_status=item_status,
        review_status=review_status,
        sort_by=sort_by,
        dedupe_latest=dedupe_latest,
        text_filter=text_filter,
        processed_only=processed_only,
        candidate_score_threshold=candidate_score_threshold,
    )
    load_result_summaries_seconds = time.perf_counter() - stage_started

    stage_started = time.perf_counter()
    details = hydrate_compare_result_summaries(
        business_db_path,
        summaries,
        retrieval_db_path=retrieval_db_path,
        owner_user_id=owner_user_id,
        review_result_limit=1,
    )
    hydrate_result_details_seconds = time.perf_counter() - stage_started

    stage_started = time.perf_counter()
    workbook = _build_workbook(
        details=details,
        owner_user_id=owner_user_id,
        task_id=task_id,
        item_status=item_status,
        review_status=review_status,
        sort_by=sort_by,
        dedupe_latest=dedupe_latest,
        text_filter=text_filter,
        processed_only=processed_only,
        candidate_score_threshold=candidate_score_threshold,
    )
    build_workbook_seconds = time.perf_counter() - stage_started

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stage_started = time.perf_counter()
    export_dir = (Path(export_root) / "review_exports" / f"user_{owner_user_id}").resolve()
    export_dir.mkdir(parents=True, exist_ok=True)
    file_path = export_dir / f"review_export_{timestamp}_{uuid4().hex[:8]}.xlsx"
    workbook.save(file_path)
    save_workbook_seconds = time.perf_counter() - stage_started
    total_elapsed_seconds = time.perf_counter() - total_started

    metrics = {
        "summary_count": len(summaries),
        "detail_count": len(details),
        "processed_only": bool(processed_only),
        "task_id": str(task_id or ""),
        "owner_user_id": int(owner_user_id),
        "stages": {
            "load_result_summaries_seconds": round(load_result_summaries_seconds, 6),
            "hydrate_result_details_seconds": round(hydrate_result_details_seconds, 6),
            "build_workbook_seconds": round(build_workbook_seconds, 6),
            "save_workbook_seconds": round(save_workbook_seconds, 6),
        },
        "total_elapsed_seconds": round(total_elapsed_seconds, 6),
    }
    return file_path, f"{REVIEW_EXPORT_FILENAME_PREFIX}_{timestamp}.xlsx", metrics


def _load_result_summaries(
    *,
    business_db_path: str | Path,
    owner_user_id: int,
    task_id: str,
    item_status: str,
    review_status: str,
    sort_by: str,
    dedupe_latest: bool,
    text_filter: str,
    processed_only: bool,
    candidate_score_threshold: float | None,
    batch_size: int = 200,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    offset = 0

    while True:
        batch = list_compare_results(
            db_path=business_db_path,
            limit=batch_size,
            offset=offset,
            task_id=task_id,
            item_status=item_status,
            review_status=review_status,
            sort_by=sort_by,
            dedupe_latest=dedupe_latest,
            owner_user_id=owner_user_id,
            text_filter=text_filter,
            exclude_cleared=True,
            candidate_score_threshold=candidate_score_threshold,
        )
        if not batch:
            break

        for item in batch:
            if processed_only and not _is_processed_result(item):
                continue
            items.append(item)

        if len(batch) < batch_size:
            break
        offset += len(batch)

    return items


def _is_processed_result(item: dict[str, Any]) -> bool:
    review = item.get("review")
    review_status = ""
    if isinstance(review, dict):
        review_status = str(review.get("review_status") or "").strip()
    return review_status not in {"", "pending", "cleared"}


def _matches_text_filter(item: dict[str, Any], text_filter: str) -> bool:
    keyword = text_filter.strip().lower()
    if not keyword:
        return True
    for value in (
        item.get("query_text_preview"),
        item.get("source_short_drama"),
        item.get("source_novel_name"),
        item.get("top1_book_name"),
        item.get("top1_chapter_name"),
        item.get("task_id"),
        item.get("top1_review_label"),
    ):
        if keyword in str(value or "").lower():
            return True
    return False


def _build_workbook(
    *,
    details: list[dict[str, Any]],
    owner_user_id: int,
    task_id: str,
    item_status: str,
    review_status: str,
    sort_by: str,
    dedupe_latest: bool,
    text_filter: str,
    processed_only: bool,
    candidate_score_threshold: float | None,
) -> Workbook:
    workbook = Workbook(write_only=True)
    grouped_rows = _group_export_rows(details)

    summary_sheet = workbook.create_sheet(title="导出说明")
    _populate_summary_sheet(
        summary_sheet,
        grouped_rows=grouped_rows,
        owner_user_id=owner_user_id,
        task_id=task_id,
        item_status=item_status,
        review_status=review_status,
        sort_by=sort_by,
        dedupe_latest=dedupe_latest,
        text_filter=text_filter,
        processed_only=processed_only,
        candidate_score_threshold=candidate_score_threshold,
    )

    for status in PROCESSED_REVIEW_STATUSES:
        sheet = workbook.create_sheet(title=SHEET_TITLE_BY_STATUS[status])
        _populate_data_sheet(sheet, grouped_rows.get(status, []))

    extra_rows = grouped_rows.get("__other__", [])
    if extra_rows:
        sheet = workbook.create_sheet(title=OTHER_PROCESSED_SHEET_TITLE)
        _populate_data_sheet(sheet, extra_rows)

    return workbook


def _group_export_rows(details: list[dict[str, Any]]) -> dict[str, list[list[Any]]]:
    grouped: dict[str, list[list[Any]]] = defaultdict(list)
    for detail in details:
        review = detail.get("review") if isinstance(detail.get("review"), dict) else {}
        review_status = str(review.get("review_status") or "").strip()
        if review_status not in PROCESSED_REVIEW_STATUSES:
            grouped["__other__"].append(_build_export_row(detail))
            continue
        grouped[review_status].append(_build_export_row(detail))
    return grouped


def _build_export_row(detail: dict[str, Any]) -> list[Any]:
    fine_payload = detail.get("result_payload")
    fine_section = fine_payload.get("fine") if isinstance(fine_payload, dict) else {}
    fine_results = fine_section.get("results") if isinstance(fine_section, dict) else []
    review_rows = fine_section.get("review_rows") if isinstance(fine_section, dict) else []

    fine_item = fine_results[0] if isinstance(fine_results, list) and fine_results and isinstance(fine_results[0], dict) else {}
    review_row = review_rows[0] if isinstance(review_rows, list) and review_rows and isinstance(review_rows[0], dict) else {}
    best_match = fine_item.get("best_match") if isinstance(fine_item.get("best_match"), dict) else {}
    review = detail.get("review") if isinstance(detail.get("review"), dict) else {}

    query_text_preview = _pick_first_non_empty(
        detail.get("query_text_preview"),
        best_match.get("query_text_preview"),
        review_row.get("query_text_preview"),
    )
    candidate_text_preview = _pick_first_non_empty(
        best_match.get("candidate_text_preview"),
        review_row.get("candidate_text_preview"),
    )
    query_text_evidence = _pick_first_non_empty(
        detail.get("query_text"),
        best_match.get("query_text"),
        review_row.get("query_text"),
        query_text_preview,
    )
    candidate_text_evidence = _pick_first_non_empty(
        best_match.get("candidate_review_context_text"),
        best_match.get("candidate_text_full"),
        best_match.get("candidate_text"),
        review_row.get("candidate_text"),
        candidate_text_preview,
    )
    candidate_text_full = _pick_first_non_empty(
        best_match.get("candidate_text_full"),
        review_row.get("candidate_text"),
        candidate_text_evidence,
    )

    return [
        detail.get("result_id"),
        detail.get("task_id"),
        detail.get("item_order"),
        detail.get("task_status"),
        detail.get("detection_mode"),
        _review_status_label(str(review.get("review_status") or "")),
        review.get("reviewer_name") or "",
        review.get("review_note") or "",
        review.get("updated_at"),
        detail.get("created_at"),
        detail.get("updated_at"),
        detail.get("task_created_at"),
        detail.get("task_finished_at"),
        detail.get("source_file_name"),
        detail.get("source_ref"),
        detail.get("source_short_drama"),
        detail.get("source_display_title"),
        detail.get("source_episode"),
        detail.get("source_author"),
        detail.get("source_platform"),
        detail.get("source_novel_name"),
        detail.get("source_excel_row"),
        detail.get("source_description"),
        _pick_first_non_empty(detail.get("top1_book_name"), fine_item.get("book_name"), review_row.get("book_name")),
        _pick_first_non_empty(detail.get("top1_chapter_name"), fine_item.get("chapter_name"), review_row.get("chapter_name")),
        _pick_first_non_empty(detail.get("top1_review_label"), fine_item.get("review_label"), review_row.get("review_label")),
        _pick_first_non_empty(detail.get("top1_confidence_label"), fine_item.get("confidence_label"), review_row.get("confidence_label")),
        detail.get("semantic_status"),
        _pick_first_non_empty(fine_item.get("fine_rank"), review_row.get("fine_rank")),
        review_row.get("coarse_rank"),
        _pick_first_non_empty(detail.get("top1_fine_score"), fine_item.get("fine_score"), review_row.get("fine_score")),
        review_row.get("coarse_final_score"),
        detail.get("duration_seconds"),
        _pick_first_non_empty(best_match.get("candidate_window_order"), review_row.get("candidate_window_order")),
        _pick_first_non_empty(best_match.get("candidate_start_offset"), review_row.get("candidate_start_offset")),
        _pick_first_non_empty(best_match.get("candidate_end_offset"), review_row.get("candidate_end_offset")),
        _pick_first_non_empty(best_match.get("longest_match_len"), review_row.get("longest_match_len")),
        _pick_first_non_empty(best_match.get("longest_match_ratio"), review_row.get("longest_match_ratio")),
        _pick_first_non_empty(best_match.get("ngram_recall"), review_row.get("ngram_recall")),
        _pick_first_non_empty(best_match.get("ngram_precision"), review_row.get("ngram_precision")),
        _pick_first_non_empty(best_match.get("jaccard"), review_row.get("jaccard")),
        _pick_first_non_empty(best_match.get("sequence_ratio"), review_row.get("sequence_ratio")),
        _pick_first_non_empty(best_match.get("exact_substring_hit"), review_row.get("exact_substring_hit")),
        _pick_first_non_empty(best_match.get("matched_substring"), review_row.get("matched_substring")),
        query_text_preview,
        candidate_text_preview,
        query_text_evidence,
        candidate_text_evidence,
        candidate_text_full,
    ]


def _populate_summary_sheet(
    sheet,
    *,
    grouped_rows: dict[str, list[list[Any]]],
    owner_user_id: int,
    task_id: str,
    item_status: str,
    review_status: str,
    sort_by: str,
    dedupe_latest: bool,
    text_filter: str,
    processed_only: bool,
    candidate_score_threshold: float | None,
) -> None:
    rows = [
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("导出范围", "当前登录账号下的复核结果导出"),
        ("账号ID", owner_user_id),
        ("仅导出已处理结果", "是" if processed_only else "否"),
        ("任务ID筛选", task_id or "全部"),
        ("结果状态筛选", item_status or "全部"),
        ("复核状态筛选", _review_status_label(review_status) if review_status else "全部"),
        ("关键词筛选", text_filter or "无"),
        ("展示阈值", candidate_score_threshold if candidate_score_threshold is not None else "未限制"),
        ("排序方式", sort_by or "updated_at_desc"),
        ("仅保留最新去重结果", "是" if dedupe_latest else "否"),
        ("确认高风险", len(grouped_rows.get("confirmed_high_risk", []))),
        ("继续跟进", len(grouped_rows.get("needs_followup", []))),
        ("标记误报", len(grouped_rows.get("false_positive", []))),
        ("其他已处理", len(grouped_rows.get("__other__", []))),
    ]

    sheet.append(["字段", "值"])
    for row in rows:
        sheet.append(list(row))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{len(rows) + 1}"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 88


def _populate_data_sheet(sheet, rows: list[list[Any]]) -> None:
    sheet.append([column.header for column in REVIEW_EXPORT_COLUMNS])
    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_EXPORT_COLUMNS))}{len(rows) + 1}"

    for index, column in enumerate(REVIEW_EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width


def _pick_first_non_empty(*values: Any) -> Any:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str):
            if value.strip() == "":
                continue
            return value
        return value
    return ""


def _review_status_label(value: str) -> str:
    normalized = value.strip()
    if not normalized:
        return REVIEW_STATUS_LABELS["pending"]
    return REVIEW_STATUS_LABELS.get(normalized, normalized)
