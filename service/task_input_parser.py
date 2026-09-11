from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import csv
import re
from typing import Iterable


TEXT_HEADER_CANDIDATES = (
    "query_text",
    "text",
    "content",
    "subtitle_text",
    "subtitle",
    "subtitle_full_text",
    "文案",
    "文本",
    "内容",
    "字幕",
    "字幕全文",
    "完整字幕",
    "待检测文本",
    "待检测文案",
)

SOURCE_REF_HEADER_CANDIDATES = (
    "source_ref",
    "share_url",
    "url",
    "source_id",
    "row_id",
    "line_no",
    "id",
    "分享链接",
    "视频链接",
    "序号",
    "编号",
    "标题",
    "title",
    "name",
    "名称",
)


SHORT_DRAMA_HEADER_CANDIDATES = (
    "short_drama",
    "shortdrama",
    "drama_name",
    "drama",
    "剧名",
    "短剧名",
)

NOVEL_NAME_HEADER_CANDIDATES = (
    "novel_name",
    "book_name",
    "novel",
    "source_novel",
)

EXCEL_ROW_HEADER_CANDIDATES = (
    "excel_row",
    "row_number",
    "row_no",
    "line_no",
)

EPISODE_HEADER_CANDIDATES = (
    "episode",
    "episode_no",
    "episode_number",
    "集数",
    "第几集",
)

AUTHOR_HEADER_CANDIDATES = (
    "author",
    "作者",
)

PLATFORM_HEADER_CANDIDATES = (
    "platform",
    "平台",
)

DISPLAY_TITLE_HEADER_CANDIDATES = (
    "display_title",
    "展示标题",
    "视频标题",
    "title",
)

DESCRIPTION_HEADER_CANDIDATES = (
    "description",
    "描述",
)


VIDEO_ID_HEADER_CANDIDATES = (
    "source_video_id",
    "video_id",
    "youtube_video_id",
    "视频ID",
)

CHANNEL_HEADER_CANDIDATES = (
    "source_channel",
    "channel",
    "channel_name",
    "uploader",
    "频道",
)

UPLOAD_DATE_HEADER_CANDIDATES = (
    "source_upload_date",
    "upload_date",
    "published_at",
    "published_date",
    "上传日期",
)

CAPTION_LANGUAGE_HEADER_CANDIDATES = (
    "source_caption_language",
    "caption_language",
    "subtitle_language",
    "language_code",
    "字幕语言",
)

CAPTION_SOURCE_HEADER_CANDIDATES = (
    "source_caption_source",
    "caption_source",
    "subtitle_source",
    "字幕来源",
)

SEGMENT_ORDER_HEADER_CANDIDATES = (
    "source_segment_order",
    "segment_order",
    "segment_index",
    "cue_order",
)

TIME_START_HEADER_CANDIDATES = (
    "source_time_start",
    "time_start",
    "start_time",
    "start_seconds",
)

TIME_END_HEADER_CANDIDATES = (
    "source_time_end",
    "time_end",
    "end_time",
    "end_seconds",
)

TIME_RANGE_HEADER_CANDIDATES = (
    "caption_range_seconds",
    "字幕范围（秒）",
    "字幕范围(秒)",
    "字幕范围",
)

ORIGINAL_TEXT_HEADER_CANDIDATES = (
    "source_text_original",
    "original_text",
    "raw_text",
)


PREFERRED_XLSX_SHEET_NAMES = (
    "batch_input",
    "结果总表",
    "Sheet1",
)


@dataclass(frozen=True)
class ParsedTaskInput:
    item_order: int
    source_ref: str
    query_text: str
    source_short_drama: str = ""
    source_novel_name: str = ""
    source_excel_row: str = ""
    source_episode: str = ""
    source_author: str = ""
    source_platform: str = ""
    source_display_title: str = ""
    source_description: str = ""
    source_video_id: str = ""
    source_channel: str = ""
    source_upload_date: str = ""
    source_caption_language: str = ""
    source_caption_source: str = ""
    source_segment_order: str = ""
    source_time_start: str = ""
    source_time_end: str = ""
    source_text_original: str = ""


def _normalize_header(value: str) -> str:
    return re.sub(r"[\s_\-]+", "", value or "").strip().lower()


def _normalized_candidates(values: Iterable[str]) -> set[str]:
    return {_normalize_header(value) for value in values}


def _is_synthetic_header(header: str) -> bool:
    return bool(re.fullmatch(r"column\d+", _normalize_header(header)))


def _load_text_blocks(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8-sig")
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    blocks = [part.strip() for part in re.split(r"\n\s*\n+", normalized) if part.strip()]
    if len(blocks) > 1:
        return blocks
    lines = [line.strip() for line in normalized.split("\n") if line.strip()]
    if len(lines) > 1 and max(len(line) for line in lines) <= 300:
        return lines
    return [normalized]


def _detect_text_column(headers: Iterable[str]) -> str:
    header_list = [header for header in headers if header is not None]
    if not header_list:
        raise ValueError("No headers found in batch file")
    if len(header_list) == 1:
        return header_list[0]
    normalized_map = {_normalize_header(header): header for header in header_list}
    for candidate in TEXT_HEADER_CANDIDATES:
        found = normalized_map.get(_normalize_header(candidate))
        if found:
            return found
    source_ref_headers = _normalized_candidates(SOURCE_REF_HEADER_CANDIDATES)
    for header in header_list:
        if _normalize_header(header) not in source_ref_headers:
            return header
    return header_list[0]


def _detect_source_ref_column(headers: Iterable[str], text_column: str) -> str | None:
    normalized_map = {_normalize_header(header): header for header in headers if header is not None}
    for candidate in SOURCE_REF_HEADER_CANDIDATES:
        found = normalized_map.get(_normalize_header(candidate))
        if found and found != text_column:
            return found
    return None


def _detect_optional_column(
    headers: Iterable[str],
    candidates: Iterable[str],
    excluded: set[str] | None = None,
) -> str | None:
    excluded_headers = excluded or set()
    normalized_map = {_normalize_header(header): header for header in headers if header is not None}
    for candidate in candidates:
        found = normalized_map.get(_normalize_header(candidate))
        if found and found not in excluded_headers:
            return found
    return None


def _score_text_value(value: str) -> int:
    compact = str(value or "").strip()
    if not compact:
        return 0
    score = min(len(compact), 240)
    if compact.isdigit():
        score -= 120
    if re.search(r"[\u4e00-\u9fffA-Za-z]{6,}", compact):
        score += 40
    if re.search(r"[，。！？、“”\"'：；,.!?]", compact):
        score += 20
    if re.search(r"\s", compact):
        score += 10
    return score


def _score_column_as_text(rows: list[dict[str, str]], column_name: str) -> int:
    score = 0
    observed = 0
    for row in rows:
        value = str(row.get(column_name, "") or "").strip()
        if not value:
            continue
        score += _score_text_value(value)
        observed += 1
        if observed >= 20:
            break
    return score


def _detect_text_column_from_rows(rows: list[dict[str, str]]) -> str:
    headers = list(rows[0].keys())
    preferred = _detect_text_column(headers)
    if not _is_synthetic_header(preferred):
        return preferred

    source_ref_headers = _normalized_candidates(SOURCE_REF_HEADER_CANDIDATES)
    scored_headers = [
        header
        for header in headers
        if _normalize_header(header) not in source_ref_headers
    ]
    if not scored_headers:
        return preferred
    return max(
        scored_headers,
        key=lambda header: (_score_column_as_text(rows, header), -headers.index(header)),
    )


def _score_source_ref_value(value: str) -> int:
    compact = str(value or "").strip()
    if not compact:
        return 0
    if re.fullmatch(r"\d+", compact):
        return 100
    if re.match(r"https?://", compact, re.IGNORECASE):
        return 90
    if len(compact) <= 40:
        return 40
    return 0


def _detect_source_ref_column_from_rows(rows: list[dict[str, str]], text_column: str) -> str | None:
    detected = _detect_source_ref_column(rows[0].keys(), text_column=text_column)
    if detected is not None:
        return detected

    headers = [header for header in rows[0].keys() if header != text_column]
    if not headers:
        return None
    scored_headers = []
    for header in headers:
        score = 0
        observed = 0
        for row in rows:
            value = str(row.get(header, "") or "").strip()
            if not value:
                continue
            score += _score_source_ref_value(value)
            observed += 1
            if observed >= 20:
                break
        if score > 0:
            scored_headers.append((score, header))
    if not scored_headers:
        return None
    scored_headers.sort(key=lambda item: item[0], reverse=True)
    return scored_headers[0][1]


def _parse_inline_source_parts(query_text: str) -> tuple[str, str, str] | None:
    compact = str(query_text or "").strip()
    if not compact:
        return None
    parts = [part.strip() for part in compact.split(",", 2)]
    if len(parts) != 3:
        return None
    source_ref, display_title, body = parts
    if not source_ref.isdigit():
        return None
    if not display_title or not body or len(body) < 20:
        return None
    return source_ref, display_title, body


def _split_time_range(value: str) -> tuple[str, str]:
    parts = [part.strip() for part in re.split(r"\s*(?:-|~|至|到)\s*", str(value or ""), maxsplit=1)]
    return (parts[0], parts[1]) if len(parts) == 2 and all(parts) else ("", "")


def _looks_like_header_row(first_row: list[str], second_row: list[str] | None = None) -> bool:
    normalized_headers = {_normalize_header(value) for value in first_row if value}
    known_header_hits = sum(
        1
        for candidate in (
            *TEXT_HEADER_CANDIDATES,
            *SOURCE_REF_HEADER_CANDIDATES,
            *SHORT_DRAMA_HEADER_CANDIDATES,
            *NOVEL_NAME_HEADER_CANDIDATES,
            *EXCEL_ROW_HEADER_CANDIDATES,
            *EPISODE_HEADER_CANDIDATES,
            *AUTHOR_HEADER_CANDIDATES,
            *PLATFORM_HEADER_CANDIDATES,
            *DISPLAY_TITLE_HEADER_CANDIDATES,
            *DESCRIPTION_HEADER_CANDIDATES,
            *VIDEO_ID_HEADER_CANDIDATES,
            *CHANNEL_HEADER_CANDIDATES,
            *UPLOAD_DATE_HEADER_CANDIDATES,
            *CAPTION_LANGUAGE_HEADER_CANDIDATES,
            *CAPTION_SOURCE_HEADER_CANDIDATES,
            *SEGMENT_ORDER_HEADER_CANDIDATES,
            *TIME_START_HEADER_CANDIDATES,
            *TIME_END_HEADER_CANDIDATES,
            *TIME_RANGE_HEADER_CANDIDATES,
            *ORIGINAL_TEXT_HEADER_CANDIDATES,
        )
        if _normalize_header(candidate) in normalized_headers
    )
    if known_header_hits > 0:
        return True

    non_empty = [value for value in first_row if value]
    if not non_empty:
        return False

    if second_row:
        first_has_long_text = any(len(value) >= 80 for value in non_empty)
        second_has_long_text = any(len(str(value or "").strip()) >= 80 for value in second_row)
        if not first_has_long_text and second_has_long_text:
            return True

    return False


def _build_unique_headers(values: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, raw_value in enumerate(values, start=1):
        base = str(raw_value or "").strip() or f"column_{index}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base}_{count + 1}")
    return headers


def _build_rows_from_dicts(rows: list[dict[str, str]]) -> list[ParsedTaskInput]:
    if not rows:
        return []
    text_column = _detect_text_column_from_rows(rows)
    source_ref_column = _detect_source_ref_column_from_rows(rows, text_column=text_column)
    excluded_headers = {text_column}
    if source_ref_column:
        excluded_headers.add(source_ref_column)
    short_drama_column = _detect_optional_column(
        rows[0].keys(),
        SHORT_DRAMA_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if short_drama_column:
        excluded_headers.add(short_drama_column)
    novel_name_column = _detect_optional_column(
        rows[0].keys(),
        NOVEL_NAME_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if novel_name_column:
        excluded_headers.add(novel_name_column)
    excel_row_column = _detect_optional_column(
        rows[0].keys(),
        EXCEL_ROW_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if excel_row_column:
        excluded_headers.add(excel_row_column)
    episode_column = _detect_optional_column(
        rows[0].keys(),
        EPISODE_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if episode_column:
        excluded_headers.add(episode_column)
    author_column = _detect_optional_column(
        rows[0].keys(),
        AUTHOR_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if author_column:
        excluded_headers.add(author_column)
    platform_column = _detect_optional_column(
        rows[0].keys(),
        PLATFORM_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if platform_column:
        excluded_headers.add(platform_column)
    display_title_column = _detect_optional_column(
        rows[0].keys(),
        DISPLAY_TITLE_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if display_title_column:
        excluded_headers.add(display_title_column)
    description_column = _detect_optional_column(
        rows[0].keys(),
        DESCRIPTION_HEADER_CANDIDATES,
        excluded=excluded_headers,
    )
    if description_column:
        excluded_headers.add(description_column)

    def optional_column(candidates: Iterable[str]) -> str | None:
        column = _detect_optional_column(rows[0].keys(), candidates, excluded=excluded_headers)
        if column:
            excluded_headers.add(column)
        return column

    video_id_column = optional_column(VIDEO_ID_HEADER_CANDIDATES)
    channel_column = optional_column(CHANNEL_HEADER_CANDIDATES)
    upload_date_column = optional_column(UPLOAD_DATE_HEADER_CANDIDATES)
    caption_language_column = optional_column(CAPTION_LANGUAGE_HEADER_CANDIDATES)
    caption_source_column = optional_column(CAPTION_SOURCE_HEADER_CANDIDATES)
    segment_order_column = optional_column(SEGMENT_ORDER_HEADER_CANDIDATES)
    time_start_column = optional_column(TIME_START_HEADER_CANDIDATES)
    time_end_column = optional_column(TIME_END_HEADER_CANDIDATES)
    time_range_column = optional_column(TIME_RANGE_HEADER_CANDIDATES)
    original_text_column = optional_column(ORIGINAL_TEXT_HEADER_CANDIDATES)
    parsed: list[ParsedTaskInput] = []
    for row_index, row in enumerate(rows, start=1):
        query_text = str(row.get(text_column, "") or "").strip()
        if not query_text:
            continue
        source_ref = ""
        if source_ref_column:
            source_ref = str(row.get(source_ref_column, "") or "").strip()
        if not source_ref:
            source_ref = f"row_{row_index}"
        source_short_drama = ""
        if short_drama_column:
            source_short_drama = str(row.get(short_drama_column, "") or "").strip()
        source_novel_name = ""
        if novel_name_column:
            source_novel_name = str(row.get(novel_name_column, "") or "").strip()
        source_excel_row = ""
        if excel_row_column:
            source_excel_row = str(row.get(excel_row_column, "") or "").strip()
        source_episode = ""
        if episode_column:
            source_episode = str(row.get(episode_column, "") or "").strip()
        source_author = ""
        if author_column:
            source_author = str(row.get(author_column, "") or "").strip()
        source_platform = ""
        if platform_column:
            source_platform = str(row.get(platform_column, "") or "").strip()
        source_display_title = ""
        if display_title_column:
            source_display_title = str(row.get(display_title_column, "") or "").strip()
        source_description = ""
        if description_column:
            source_description = str(row.get(description_column, "") or "").strip()
        source_video_id = str(row.get(video_id_column, "") or "").strip() if video_id_column else ""
        source_channel = str(row.get(channel_column, "") or "").strip() if channel_column else ""
        source_upload_date = str(row.get(upload_date_column, "") or "").strip() if upload_date_column else ""
        source_caption_language = str(row.get(caption_language_column, "") or "").strip() if caption_language_column else ""
        source_caption_source = str(row.get(caption_source_column, "") or "").strip() if caption_source_column else ""
        source_segment_order = str(row.get(segment_order_column, "") or "").strip() if segment_order_column else ""
        source_time_start = str(row.get(time_start_column, "") or "").strip() if time_start_column else ""
        source_time_end = str(row.get(time_end_column, "") or "").strip() if time_end_column else ""
        if time_range_column and not source_time_start and not source_time_end:
            source_time_start, source_time_end = _split_time_range(str(row.get(time_range_column, "") or ""))
        source_text_original = str(row.get(original_text_column, "") or "").strip() if original_text_column else query_text
        inline_source_parts = _parse_inline_source_parts(query_text)
        if inline_source_parts is not None:
            inline_source_ref, inline_display_title, inline_query_text = inline_source_parts
            if not source_ref or source_ref.startswith("row_"):
                source_ref = inline_source_ref
            if not source_display_title:
                source_display_title = inline_display_title
            query_text = inline_query_text
        parsed.append(
            ParsedTaskInput(
                item_order=len(parsed) + 1,
                source_ref=source_ref,
                query_text=query_text,
                source_short_drama=source_short_drama,
                source_novel_name=source_novel_name,
                source_excel_row=source_excel_row,
                source_episode=source_episode,
                source_author=source_author,
                source_platform=source_platform,
                source_display_title=source_display_title,
                source_description=source_description,
                source_video_id=source_video_id,
                source_channel=source_channel,
                source_upload_date=source_upload_date,
                source_caption_language=source_caption_language,
                source_caption_source=source_caption_source,
                source_segment_order=source_segment_order,
                source_time_start=source_time_start,
                source_time_end=source_time_end,
                source_text_original=source_text_original,
            )
        )
    return parsed


def _load_csv_rows(path: Path) -> list[dict[str, str]]:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:4096]
    delimiter = ","
    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        return list(reader)


def _load_xlsx_rows(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Parsing .xlsx requires openpyxl to be installed") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = _select_xlsx_sheet(workbook)
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    raw_rows = [
        [str(value).strip() if value is not None else "" for value in values]
        for values in rows
    ]
    first_row = raw_rows[0]
    second_row = raw_rows[1] if len(raw_rows) > 1 else None
    if _looks_like_header_row(first_row, second_row):
        headers = _build_unique_headers(first_row)
        data_rows_raw = raw_rows[1:]
    else:
        column_count = max(len(values) for values in raw_rows)
        headers = [f"column_{index}" for index in range(1, column_count + 1)]
        data_rows_raw = raw_rows
    data_rows: list[dict[str, str]] = []
    for values in data_rows_raw:
        row = {
            headers[index]: values[index]
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }
        data_rows.append(row)
    return data_rows


def _select_xlsx_sheet(workbook):
    def _headers_for_sheet(sheet) -> list[str]:
        try:
            first_row = next(sheet.iter_rows(values_only=True, min_row=1, max_row=1), None)
        except TypeError:
            first_row = next(sheet.iter_rows(values_only=True), None)
        if not first_row:
            return []
        return [str(value).strip() if value is not None else "" for value in first_row]

    def _sheet_score(sheet) -> tuple[int, int]:
        headers = _headers_for_sheet(sheet)
        normalized_headers = {_normalize_header(header) for header in headers if header}
        text_hits = sum(
            1 for candidate in TEXT_HEADER_CANDIDATES if _normalize_header(candidate) in normalized_headers
        )
        meta_hits = sum(
            1
            for candidate in (
                *SOURCE_REF_HEADER_CANDIDATES,
                *SHORT_DRAMA_HEADER_CANDIDATES,
                *NOVEL_NAME_HEADER_CANDIDATES,
                *EXCEL_ROW_HEADER_CANDIDATES,
                "分享链接",
            )
            if _normalize_header(candidate) in normalized_headers
        )
        return (text_hits, meta_hits)

    named_sheets = {str(sheet.title).strip(): sheet for sheet in workbook.worksheets}
    for preferred_name in PREFERRED_XLSX_SHEET_NAMES:
        preferred_sheet = named_sheets.get(preferred_name)
        if preferred_sheet is None:
            continue
        if _sheet_score(preferred_sheet)[0] > 0:
            return preferred_sheet

    best_sheet = workbook.active
    best_score = _sheet_score(best_sheet)
    for sheet in workbook.worksheets:
        score = _sheet_score(sheet)
        if score > best_score:
            best_sheet = sheet
            best_score = score
    return best_sheet


def parse_task_input_file(path: str | Path) -> list[ParsedTaskInput]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        return [
            ParsedTaskInput(item_order=index, source_ref=f"block_{index}", query_text=text)
            for index, text in enumerate(_load_text_blocks(file_path), start=1)
            if text.strip()
        ]
    if suffix in {".csv", ".tsv"}:
        return _build_rows_from_dicts(_load_csv_rows(file_path))
    if suffix == ".xlsx":
        return _build_rows_from_dicts(_load_xlsx_rows(file_path))
    raise ValueError(f"Unsupported task input file type: {suffix}")
