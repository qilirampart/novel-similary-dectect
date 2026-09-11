from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from service.drama_subtitle_task_store import get_drama_subtitle_task, list_drama_subtitle_task_items


REVIEW_STATUSES = ("confirmed_high_risk", "needs_followup", "false_positive")
SHEET_TITLES = {
    "confirmed_high_risk": "确认高风险",
    "needs_followup": "继续跟进",
    "false_positive": "标记误报",
}
HEADERS = (
    "source_video_id", "source_channel", "source_upload_date", "source_caption_language", "source_caption_source", "source_segment_order", "source_time_start", "source_time_end", "source_text_original",
    "任务ID", "条目序号", "复核状态", "复核人", "复核备注", "来源平台", "来源标识", "短剧名称", "集数", "作者",
    "输入语言", "输入台词", "检测耗时(秒)", "命中bookId", "命中剧名", "命中集数", "命中字幕语言",
    "证据窗口ID", "候选字幕起始时间", "候选字幕结束时间", "命中字幕证据全文", "检测状态", "最终判定", "判定原因", "候选剧名", "候选集数", "候选语义分数", "创建时间", "复核更新时间",
)


def _payload_evidence(item: dict[str, Any]) -> str:
    try:
        payload = json.loads(str(item.get("result_payload_json") or "{}"))
    except json.JSONDecodeError:
        return ""
    candidates = payload.get("candidates") if isinstance(payload, dict) else []
    top1 = candidates[0] if isinstance(candidates, list) and candidates else {}
    evidence = top1.get("evidence") if isinstance(top1, dict) else {}
    return str(evidence.get("window_text") or evidence.get("window_text_preview") or "") if isinstance(evidence, dict) else ""


def _payload_decision(item: dict[str, Any]) -> dict[str, Any]:
    try:
        payload = json.loads(str(item.get("result_payload_json") or "{}"))
    except json.JSONDecodeError:
        return {}
    decision = payload.get("decision") if isinstance(payload, dict) else {}
    return decision if isinstance(decision, dict) else {}


def _payload_top_candidate(item: dict[str, Any]) -> dict[str, Any]:
    try:
        payload = json.loads(str(item.get("result_payload_json") or "{}"))
    except json.JSONDecodeError:
        return {}
    candidates = payload.get("candidates") if isinstance(payload, dict) else []
    top1 = candidates[0] if isinstance(candidates, list) and candidates else {}
    return top1 if isinstance(top1, dict) else {}


def _row(item: dict[str, Any]) -> list[Any]:
    decision = _payload_decision(item)
    top_candidate = _payload_top_candidate(item)
    return [
        item.get("source_video_id", ""), item.get("source_channel", ""), item.get("source_upload_date", ""), item.get("source_caption_language", ""), item.get("source_caption_source", ""), item.get("source_segment_order", ""), item.get("source_time_start", ""), item.get("source_time_end", ""), item.get("source_text_original", ""),
        item.get("task_id", ""), item.get("item_order", ""), item.get("review_status", "pending"), item.get("reviewer_name", ""), item.get("review_note", ""),
        item.get("source_platform", ""), item.get("source_ref", ""), item.get("source_short_drama") or item.get("source_display_title", ""), item.get("source_episode", ""), item.get("source_author", ""),
        item.get("query_language_code", ""), item.get("query_text", ""), item.get("duration_seconds", ""), item.get("matched_book_id", ""), item.get("matched_book_name", ""),
        item.get("matched_episode_order", ""), item.get("matched_language_code", ""), item.get("evidence_window_uid", ""), item.get("evidence_time_start", ""), item.get("evidence_time_end", ""),
        _payload_evidence(item), item.get("status", ""), decision.get("status", ""), decision.get("reason", ""),
        top_candidate.get("book_name", ""), top_candidate.get("episode_order", ""), top_candidate.get("semantic_score", ""),
        item.get("created_at", ""), item.get("review_updated_at", ""),
    ]


def build_drama_subtitle_review_export_xlsx(
    *,
    business_db_path: str | Path,
    export_root: str | Path,
    owner_user_id: int,
    task_id: str,
) -> tuple[Path, str]:
    task = get_drama_subtitle_task(business_db_path, task_id, owner_user_id=owner_user_id)
    if task is None:
        raise ValueError("subtitle task not found")
    items = list_drama_subtitle_task_items(business_db_path, task_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    grouped = {status: [item for item in items if item.get("review_status") == status] for status in REVIEW_STATUSES}
    for status in REVIEW_STATUSES:
        sheet = workbook.create_sheet(SHEET_TITLES[status])
        sheet.append(list(HEADERS))
        for item in grouped[status]:
            sheet.append(_row(item))
        _format_sheet(sheet)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = (Path(export_root) / "drama_subtitle_exports" / f"user_{owner_user_id}").resolve()
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"drama_subtitle_review_{timestamp}_{uuid4().hex[:8]}.xlsx"
    workbook.save(path)
    return path, f"短剧字幕复核导出_{timestamp}.xlsx"


def _format_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate((20, 24, 16, 18, 18, 14, 16, 16, 80, 36, 10, 16, 14, 28, 14, 38, 28, 10, 16, 10, 54, 14, 18, 28, 12, 14, 28, 18, 18, 90, 14, 22, 22), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
