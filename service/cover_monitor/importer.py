from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Any, Iterator, Optional
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook


HEADER_ALIASES = {
    "operator_name": ("代理商简称", "代理商", "运营商", "operator", "operator_name"),
    "channel_id": ("频道 ID", "频道ID", "channel_id", "channel id"),
    "channel_name": ("频道名", "频道名称", "channel_name", "channel name"),
    "channel_url": ("频道链接", "频道地址", "channel_url", "channel url"),
    "video_title": ("视频标题", "标题", "video_title", "video title"),
    "video_url": ("原视频链接", "视频链接", "video_url", "video url"),
    "video_id": ("视频 ID", "视频ID", "video_id", "video id"),
    "thumbnail_url": ("封面 CDN 地址", "封面地址", "缩略图地址", "thumbnail_url"),
    "overall_risk": ("检测结论", "风险结论", "overall_risk", "result"),
    "risk_tags": ("风险标签", "risk_tags", "tags"),
    "summary": ("摘要", "summary"),
    "evidence": ("可见证据", "证据", "evidence"),
    "confidence": ("置信度", "confidence"),
    "upload_date": ("上传日期", "发布日期", "upload_date"),
}

RISK_VALUES = {"safe", "review", "risk", "unknown"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"[\s_\-]+", "", _text(value)).casefold()


def _extract_channel_id(channel_url: str) -> str:
    if not channel_url:
        return ""
    path = urlparse(channel_url).path.strip("/")
    match = re.search(r"(?:^|/)channel/(UC[A-Za-z0-9_-]+)(?:/|$)", path)
    return match.group(1) if match else ""


def _extract_video_id(video_url: str) -> str:
    if not video_url:
        return ""
    parsed = urlparse(video_url)
    if parsed.hostname in {"youtu.be", "www.youtu.be"}:
        return parsed.path.strip("/").split("/")[0]
    query_id = parse_qs(parsed.query).get("v", [""])[0]
    if query_id:
        return query_id
    match = re.search(r"/(?:shorts|embed)/([A-Za-z0-9_-]+)", parsed.path)
    return match.group(1) if match else ""


def _confidence(value: Any) -> Optional[float]:
    text = _text(value).rstrip("%")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if number > 1 and number <= 100:
        number /= 100
    return number if 0 <= number <= 1 else None


@dataclass
class ParsedImportRow:
    source_sheet: str
    source_row: int
    status: str
    normalized_key: str
    raw: dict[str, Any]
    normalized: dict[str, Any]
    error_message: str = ""
    conflicts: list[dict[str, Any]] = field(default_factory=list)


class WorkbookImportParser:
    def __init__(
        self,
        path: str | Path,
        *,
        import_kind: str,
        sheet_name: str = "",
        existing_channels: Optional[dict[str, dict[str, Any]]] = None,
        existing_videos: Optional[dict[str, dict[str, Any]]] = None,
    ) -> None:
        if import_kind not in {"baseline", "channels", "videos"}:
            raise ValueError("unsupported import kind")
        self.path = Path(path)
        self.import_kind = import_kind
        self.requested_sheet = sheet_name.strip()
        self.existing_channels = existing_channels or {}
        self.existing_videos = existing_videos or {}
        self.sheet_name = ""
        self.mapping: dict[str, str] = {}
        self.stats: dict[str, int] = {
            "total_rows": 0,
            "valid_rows": 0,
            "unique_channels": 0,
            "unique_videos": 0,
            "duplicate_rows": 0,
            "conflict_rows": 0,
            "missing_rows": 0,
            "operator_conflicts": 0,
        }

    def rows(self) -> Iterator[ParsedImportRow]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if self.requested_sheet:
                if self.requested_sheet not in workbook.sheetnames:
                    raise ValueError(f"sheet not found: {self.requested_sheet}")
                worksheet = workbook[self.requested_sheet]
            else:
                worksheet = workbook.worksheets[0]
            self.sheet_name = worksheet.title
            row_iter = worksheet.iter_rows(values_only=True)
            try:
                headers = [_text(value) for value in next(row_iter)]
            except StopIteration as exc:
                raise ValueError("workbook has no header row") from exc
            self.mapping = self._resolve_mapping(headers)
            seen_channels: dict[str, dict[str, Any]] = {}
            seen_videos: dict[str, dict[str, Any]] = {}
            for row_number, values in enumerate(row_iter, start=2):
                raw = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
                if not any(_text(value) for value in values):
                    continue
                self.stats["total_rows"] += 1
                parsed = self._parse_row(raw, row_number)
                channel_key = parsed.normalized.get("channel_id", "")
                video_key = parsed.normalized.get("video_id", "")
                identity_key = video_key if self.import_kind != "channels" else channel_key
                parsed.normalized_key = identity_key

                file_channel = seen_channels.get(channel_key) if channel_key else None
                previous_channel = file_channel or self.existing_channels.get(channel_key)
                if previous_channel:
                    previous_operator = previous_channel.get("operator_name", "")
                    current_operator = parsed.normalized.get("operator_name", "")
                    if previous_operator and current_operator and previous_operator != current_operator:
                        parsed.conflicts.append({
                            "conflict_type": "channel_operator_mismatch",
                            "natural_key": channel_key,
                            "existing": previous_channel,
                            "incoming": parsed.normalized,
                        })
                        self.stats["operator_conflicts"] += 1
                has_channel_conflict = any(
                    item["conflict_type"] == "channel_operator_mismatch"
                    for item in parsed.conflicts
                )
                if file_channel is None and channel_key and not has_channel_conflict:
                    seen_channels[channel_key] = parsed.normalized

                file_video = seen_videos.get(video_key) if video_key else None
                previous_video = file_video or self.existing_videos.get(video_key)
                if previous_video and previous_video.get("channel_id") != channel_key:
                    parsed.conflicts.append({
                        "conflict_type": "video_channel_mismatch",
                        "natural_key": video_key,
                        "existing": previous_video,
                        "incoming": parsed.normalized,
                    })
                has_video_conflict = any(
                    item["conflict_type"] == "video_channel_mismatch"
                    for item in parsed.conflicts
                )
                if file_video is None and video_key and not has_video_conflict:
                    seen_videos[video_key] = parsed.normalized

                if parsed.conflicts:
                    parsed.status = "conflict"
                    self.stats["conflict_rows"] += 1
                elif parsed.status == "valid" and identity_key:
                    previous_identity = (
                        file_channel if self.import_kind == "channels" else file_video
                    )
                    if previous_identity is not None:
                        parsed.status = "duplicate"
                        self.stats["duplicate_rows"] += 1
                    else:
                        self.stats["valid_rows"] += 1
                elif parsed.status == "valid":
                    self.stats["valid_rows"] += 1
                elif parsed.status == "missing":
                    self.stats["missing_rows"] += 1
                yield parsed
            self.stats["unique_channels"] = len(seen_channels)
            self.stats["unique_videos"] = len(seen_videos)
        finally:
            workbook.close()

    def _resolve_mapping(self, headers: list[str]) -> dict[str, str]:
        normalized_headers = {_header_key(header): header for header in headers if header}
        mapping: dict[str, str] = {}
        for field, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                header = normalized_headers.get(_header_key(alias))
                if header:
                    mapping[field] = header
                    break
        required = {"channel_name"}
        if self.import_kind != "channels":
            required.update({"video_title", "thumbnail_url"})
        missing = sorted(field for field in required if field not in mapping)
        if "channel_id" not in mapping and "channel_url" not in mapping:
            missing.append("channel_id/channel_url")
        if self.import_kind != "channels" and "video_id" not in mapping and "video_url" not in mapping:
            missing.append("video_id/video_url")
        if missing:
            raise ValueError("missing required columns: " + ", ".join(missing))
        return mapping

    def _parse_row(self, raw: dict[str, Any], row_number: int) -> ParsedImportRow:
        def value(field: str) -> Any:
            header = self.mapping.get(field)
            return raw.get(header) if header else None

        channel_url = _text(value("channel_url"))
        channel_id = _text(value("channel_id")) or _extract_channel_id(channel_url)
        if not channel_url and channel_id:
            channel_url = f"https://www.youtube.com/channel/{channel_id}"
        video_url = _text(value("video_url"))
        video_id = _text(value("video_id")) or _extract_video_id(video_url)
        if not video_url and video_id:
            video_url = f"https://www.youtube.com/watch?v={video_id}"
        risk = _text(value("overall_risk")).casefold()
        normalized = {
            "platform": "youtube",
            "operator_name": _text(value("operator_name")),
            "channel_id": channel_id,
            "channel_name": _text(value("channel_name")),
            "channel_url": channel_url,
            "video_id": video_id,
            "video_title": _text(value("video_title")),
            "video_url": video_url,
            "thumbnail_url": _text(value("thumbnail_url")),
            "overall_risk": risk if risk in RISK_VALUES else "",
            "risk_tags": _text(value("risk_tags")),
            "summary": _text(value("summary")),
            "evidence": _text(value("evidence")),
            "confidence": _confidence(value("confidence")),
            "upload_date": _text(value("upload_date")),
        }
        required = {
            "channel_id": channel_id,
            "channel_name": normalized["channel_name"],
            "channel_url": channel_url,
        }
        if self.import_kind != "channels":
            required.update({
                "video_id": video_id,
                "video_title": normalized["video_title"],
                "video_url": video_url,
                "thumbnail_url": normalized["thumbnail_url"],
            })
        missing = [field for field, item in required.items() if not item]
        status = "missing" if missing else "valid"
        return ParsedImportRow(
            source_sheet=self.sheet_name,
            source_row=row_number,
            status=status,
            normalized_key="",
            raw=raw,
            normalized=normalized,
            error_message=("missing required values: " + ", ".join(missing)) if missing else "",
        )
