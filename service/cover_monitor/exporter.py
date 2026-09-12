from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPORT_KINDS = {"new-findings", "historical-rectification"}
DETAIL_HEADERS = (
    "频道 ID", "频道名", "代理商", "视频 ID", "视频标题", "原视频链接", "封面 CDN URL",
    "上传日期", "任务原因", "任务状态", "检测结论", "风险标签", "置信度", "模型摘要",
    "可见证据", "案件状态", "整改判定", "图片内容 SHA-256", "图片尺寸", "模型",
    "提示词版本", "检测耗时(秒)", "检测完成时间", "失败类型", "失败原因", "巡检批次",
)


def _safe_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    return f"'{value}" if text.startswith(("=", "+", "-", "@")) else value


def _risk_label(value: Any) -> str:
    return {
        "safe": "安全",
        "review": "待复核",
        "risk": "风险",
        "unknown": "未知/异常",
    }.get(str(value or ""), "未完成")


def _rectification_label(item: dict[str, Any]) -> str:
    if item.get("case_status") == "confirmed_rectified":
        return "已确认整改"
    if item.get("case_status") == "false_positive":
        return "已标记误报"
    if item.get("case_status") == "unavailable":
        return "无法访问"
    return {
        "rectification_candidate": "换图后安全，待人工确认",
        "safe_redetection": "同图结论变化，待复核",
        "risk_detected": "复查仍有风险",
        "review_detected": "复查待人工判断",
        "unknown_detected": "本轮无法判断",
    }.get(str(item.get("latest_case_event") or ""), "无法核验" if item.get("status") != "succeeded" else "无整改案件")


def _detail_row(item: dict[str, Any], run_id: str) -> list[Any]:
    snapshot = item.get("channel_snapshot") if isinstance(item.get("channel_snapshot"), dict) else {}
    return [
        snapshot.get("channel_id", ""),
        snapshot.get("channel_name", ""),
        snapshot.get("operator_name", "") or "未分配",
        item.get("video_id", ""),
        item.get("video_title", ""),
        item.get("video_url", ""),
        item.get("fetched_url") or item.get("thumbnail_url", ""),
        item.get("upload_date", ""),
        item.get("reason", ""),
        item.get("status", ""),
        _risk_label(item.get("overall_risk")),
        "、".join(str(value) for value in item.get("risk_tags") or []),
        item.get("confidence", ""),
        item.get("summary", ""),
        item.get("evidence", ""),
        item.get("case_status", ""),
        _rectification_label(item),
        item.get("content_sha256", ""),
        f"{item.get('width')}×{item.get('height')}" if item.get("width") and item.get("height") else "",
        item.get("model", ""),
        item.get("prompt_version", ""),
        item.get("duration_seconds", ""),
        item.get("detected_at", ""),
        item.get("error_type", ""),
        item.get("error_message", ""),
        run_id,
    ]


def _append_rows(sheet: Any, headers: Iterable[str], rows: Iterable[Iterable[Any]]) -> None:
    header_values = list(headers)
    widths = (22, 24, 18, 22, 42, 42, 42, 14, 18, 14, 14, 22, 12, 42, 52, 18, 28, 34, 15, 24, 24, 15, 23, 18, 42, 38)
    for index, width in enumerate(widths[: len(header_values)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_cells = []
    for value in header_values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cells.append(cell)
    sheet.append(header_cells)
    row_count = 1
    for row in rows:
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=_safe_cell(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cells.append(cell)
        sheet.append(cells)
        row_count += 1
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header_values))}{row_count}"


def _new_summary(items: list[dict[str, Any]]) -> list[list[Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        snapshot = item.get("channel_snapshot") or {}
        grouped[str(snapshot.get("operator_name") or "未分配")].append(item)
    rows: list[list[Any]] = []
    for operator, group in sorted(grouped.items()):
        rows.append([
            operator,
            len(group),
            sum(item.get("overall_risk") == "risk" for item in group),
            sum(item.get("overall_risk") == "review" for item in group),
            sum(item.get("overall_risk") == "safe" for item in group),
            sum(item.get("overall_risk") == "unknown" for item in group),
            sum(item.get("status") != "succeeded" for item in group),
        ])
    return rows


def _rectification_summary(items: list[dict[str, Any]]) -> list[list[Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        snapshot = item.get("channel_snapshot") or {}
        grouped[str(snapshot.get("operator_name") or "未分配")].append(item)
    rows: list[list[Any]] = []
    for operator, group in sorted(grouped.items()):
        confirmed = sum(item.get("case_status") == "confirmed_rectified" for item in group)
        effective = sum(item.get("overall_risk") in {"safe", "review", "risk", "unknown"} for item in group)
        rows.append([
            operator,
            len(group),
            confirmed,
            sum(item.get("latest_case_event") == "rectification_candidate" and item.get("case_status") != "confirmed_rectified" for item in group),
            sum(item.get("latest_case_event") == "risk_detected" for item in group),
            sum(item.get("overall_risk") in {"review", "unknown"} for item in group),
            len(group) - effective,
            effective,
            confirmed / len(group) if group else "",
        ])
    return rows


def build_cover_report_xlsx(
    *,
    report: dict[str, Any],
    export_kind: str,
    export_root: str | Path,
    user_id: int,
) -> tuple[Path, str]:
    if export_kind not in EXPORT_KINDS:
        raise ValueError("unsupported cover export kind")
    run = report.get("run") if isinstance(report.get("run"), dict) else {}
    run_id = str(run.get("run_id") or "")
    all_items = list(report.get("items") or [])
    if export_kind == "new-findings":
        items = [item for item in all_items if item.get("reason") != "historical_risk"]
    else:
        items = [item for item in all_items if item.get("reason") == "historical_risk"]

    workbook = Workbook(write_only=True)
    if export_kind == "new-findings":
        sheets = (
            ("新增视频明细", items),
            ("新增风险", [item for item in items if item.get("overall_risk") == "risk"]),
            ("待复核与失败", [item for item in items if item.get("overall_risk") in {"review", "unknown"} or item.get("status") != "succeeded"]),
        )
        for title, selected in sheets:
            _append_rows(workbook.create_sheet(title), DETAIL_HEADERS, (_detail_row(item, run_id) for item in selected))
        _append_rows(
            workbook.create_sheet("代理商汇总"),
            ("代理商", "检测项", "风险", "待复核", "安全", "未知/异常", "执行失败"),
            _new_summary(items),
        )
        prefix = "封面新增检测报告"
    else:
        _append_rows(
            workbook.create_sheet("历史风险整改明细"),
            DETAIL_HEADERS,
            (_detail_row(item, run_id) for item in items),
        )
        _append_rows(
            workbook.create_sheet("代理商整改汇总"),
            ("代理商", "冻结待整改数", "已确认整改", "待确认整改", "仍有风险", "待复核", "无法核验", "有效复查数", "确认整改率"),
            _rectification_summary(items),
        )
        prefix = "封面历史整改报告"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = (Path(export_root) / f"user_{int(user_id)}").resolve()
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"cover_{export_kind}_{timestamp}_{uuid4().hex[:8]}.xlsx"
    workbook.save(path)
    return path, f"{prefix}_{run_id[:8]}_{timestamp}.xlsx"
