from __future__ import annotations

from importlib import reload
from io import BytesIO
from pathlib import Path
import sys
import threading
import time
from typing import Any, Dict, Optional
from urllib.parse import unquote

from fastapi.testclient import TestClient
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def login_as_admin(client: TestClient) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "admin123"},
    )
    assert response.status_code == 200


def test_fault_diagnostics_registers_signal_handler(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business_fault_diag.sqlite3"
    retrieval_db = tmp_path / "retrieval_fault_diag.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    retrieval_db.touch()

    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_FAULT_DIAGNOSTICS_ENABLED", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_FAULT_DIAGNOSTICS_SIGNAL", "SIGUSR1")

    import api.config as api_config
    import api.app as api_app

    reload(api_config)
    api_app = reload(api_app)

    observed_calls: list[tuple[str, int, object | None]] = []

    def fake_enable(*, all_threads: bool = False) -> None:
        observed_calls.append(("enable", int(all_threads), None))

    def fake_register(signal_num: int, *, all_threads: bool = False, chain: bool = False) -> None:
        observed_calls.append(("register", int(signal_num), (bool(all_threads), bool(chain))))

    def fake_unregister(signal_num: int) -> None:
        observed_calls.append(("unregister", int(signal_num), None))

    monkeypatch.setattr(api_app.signal, "SIGUSR1", 10, raising=False)
    monkeypatch.setattr(api_app.faulthandler, "enable", fake_enable)
    monkeypatch.setattr(api_app.faulthandler, "register", fake_register, raising=False)
    monkeypatch.setattr(api_app.faulthandler, "unregister", fake_unregister, raising=False)

    api_app._FAULT_DIAGNOSTICS_ARMED = False
    api_app._FAULT_DIAGNOSTICS_SIGNAL_NUM = None

    api_app._configure_fault_diagnostics()
    api_app._teardown_fault_diagnostics()

    assert observed_calls == [
        ("enable", 1, None),
        ("unregister", 10, None),
        ("register", 10, (True, False)),
        ("unregister", 10, None),
    ]
    assert api_app._FAULT_DIAGNOSTICS_ARMED is False
    assert api_app._FAULT_DIAGNOSTICS_SIGNAL_NUM is None


def test_start_auto_worker_respects_worker_count(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business_runtime_worker.sqlite3"
    retrieval_db = tmp_path / "retrieval_runtime_worker.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    retrieval_db.touch()

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_COUNT", "2")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_POLL_SECONDS", "0.5")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker

    reload(api_config)
    runtime_worker = reload(runtime_worker)

    observed_worker_names: set[str] = set()
    entered_workers = 0
    entered_lock = threading.Lock()
    both_entered = threading.Event()

    def fake_claim_next_queued_task_globally(*_args: Any, **kwargs: Any) -> Optional[Dict[str, Any]]:
        nonlocal entered_workers
        with entered_lock:
            observed_worker_names.add(str(kwargs.get("worker_name") or ""))
            entered_workers += 1
            if entered_workers >= 2:
                both_entered.set()
        time.sleep(0.05)
        return None

    monkeypatch.setattr(runtime_worker, "claim_next_queued_task_globally", fake_claim_next_queued_task_globally)

    handle = runtime_worker.start_auto_worker()
    assert handle is not None
    assert len(handle.threads) == 2
    assert both_entered.wait(timeout=3)

    runtime_worker.stop_auto_worker(handle)

    assert observed_worker_names == {
        "local-worker-api-auto-1",
        "local-worker-api-auto-2",
    }
    assert all(not thread.is_alive() for thread in handle.threads)


def test_start_auto_worker_runs_periodic_recovery_sweep(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business_runtime_recovery.sqlite3"
    retrieval_db = tmp_path / "retrieval_runtime_recovery.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    retrieval_db.touch()

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_COUNT", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_POLL_SECONDS", "0.5")
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_RECOVERY_SWEEP_SECONDS", "0.05")
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_RECOVERY_STALE_SECONDS", "33")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker

    reload(api_config)
    runtime_worker = reload(runtime_worker)

    observed_calls: list[tuple[str, float]] = []
    recovery_called = threading.Event()

    def fake_claim_next_queued_task_globally(*_args: Any, **kwargs: Any) -> Optional[Dict[str, Any]]:
        time.sleep(0.02)
        return None

    def fake_recover_interrupted_tasks(db_path: str, *, stale_after_seconds: float) -> Dict[str, int]:
        observed_calls.append((db_path, stale_after_seconds))
        recovery_called.set()
        return {
            "running_to_queued": 0,
            "cancel_requested_to_cancelled": 0,
            "pause_requested_to_paused": 0,
            "requeued_item_count": 0,
        }

    monkeypatch.setattr(runtime_worker, "claim_next_queued_task_globally", fake_claim_next_queued_task_globally)
    monkeypatch.setattr(runtime_worker, "recover_interrupted_tasks", fake_recover_interrupted_tasks)

    handle = runtime_worker.start_auto_worker()
    assert handle is not None
    assert recovery_called.wait(timeout=3)

    runtime_worker.stop_auto_worker(handle)

    assert observed_calls
    assert observed_calls[0] == (str(business_db), 33.0)
    if handle.maintenance_thread is not None:
        assert not handle.maintenance_thread.is_alive()


def test_api_auto_worker_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business.sqlite3"
    retrieval_db = tmp_path / "retrieval.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"
    retrieval_db.touch()

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_COUNT", "2")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_POLL_SECONDS", "0.2")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    runtime_worker = reload(runtime_worker)
    api_app = reload(api_app)

    from service.business_store import claim_next_queued_task_globally as real_claim_next_queued_task_globally
    from service.task_executor import ComparePipelineRequest
    from service.task_executor import execute_claimed_task as real_execute_claimed_task

    captured_thresholds: list[float | None] = []
    observed_worker_names: set[str] = set()

    def fake_pipeline(request: ComparePipelineRequest) -> Dict[str, Any]:
        captured_thresholds.append(request.candidate_display_score_threshold)
        return {
            "query_text": request.query_text,
            "detection_mode": request.detection_mode,
            "rewrite_detection": {
                "status": "fallback_lexical_only",
            },
            "fine": {
                "results": [
                    {
                        "book_name": "Smoke Novel",
                        "chapter_name": "Chapter 1",
                        "review_label": "强证据",
                        "confidence_label": "strong",
                        "fine_score": 0.91,
                    }
                ],
                "review_rows": [
                    {
                        "detection_mode": request.detection_mode,
                        "fine_rank": 1,
                        "review_label": "强证据",
                        "confidence_label": "strong",
                        "fine_score": 0.91,
                        "coarse_rank": 1,
                        "coarse_final_score": 0.83,
                        "dataset_key": "smoke_dataset",
                        "book_ext_id": "book-1",
                        "book_name": "Smoke Novel",
                        "chapter_uid": 1,
                        "chapter_ext_id": "chapter-1",
                        "chapter_name": "Chapter 1",
                        "candidate_window_order": 1,
                        "candidate_start_offset": 0,
                        "candidate_end_offset": 128,
                        "exact_substring_hit": True,
                        "longest_match_len": len(request.query_text),
                        "longest_match_ratio": 1.0,
                        "ngram_recall": 1.0,
                        "ngram_precision": 1.0,
                        "jaccard": 1.0,
                        "sequence_ratio": 1.0,
                        "matched_substring": request.query_text,
                        "query_text_preview": request.query_text,
                        "candidate_text_preview": request.query_text,
                        "query_text": request.query_text,
                        "candidate_text": request.query_text,
                    }
                ],
            },
        }

    def claim_with_observed_worker(db_path: str, **kwargs: Any) -> Optional[Dict[str, Any]]:
        observed_worker_names.add(str(kwargs.get("worker_name") or ""))
        return real_claim_next_queued_task_globally(db_path, **kwargs)

    def execute_with_fake_pipeline(**kwargs: Any) -> Optional[Dict[str, Any]]:
        return real_execute_claimed_task(compare_fn=fake_pipeline, **kwargs)

    monkeypatch.setattr(runtime_worker, "claim_next_queued_task_globally", claim_with_observed_worker)
    monkeypatch.setattr(runtime_worker, "execute_claimed_task", execute_with_fake_pipeline)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.txt", b"first query\n\nsecond query", "text/plain")},
            data={
                "detection_mode": "rewrite",
                "candidate_display_score_threshold": "0.35",
            },
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        deadline = time.time() + 10
        last_payload: Optional[Dict[str, Any]] = None
        while time.time() < deadline:
            detail_response = client.get(f"/api/v1/tasks/{task_id}")
            assert detail_response.status_code == 200
            last_payload = detail_response.json()
            if last_payload["task"]["status"] == "completed":
                break
            time.sleep(0.2)

        assert last_payload is not None
        assert last_payload["task"]["status"] == "completed"
        assert last_payload["task"]["counts"]["accepted"] == 2
        assert last_payload["task"]["counts"]["completed"] == 2
        assert last_payload["task"]["counts"]["failed"] == 0
        assert last_payload["task"]["params"]["candidate_display_score_threshold"] == 0.35
        assert last_payload["item_total"] == 2
        assert last_payload["result_stats"]["item_total"] == 2
        assert last_payload["result_stats"]["high_risk_count"] == 2
        assert last_payload["result_stats"]["semantic_fallback_count"] == 2
        assert len(last_payload["items"]) == 2
        assert last_payload["items"][0]["top1_book_name"] == "Smoke Novel"
        assert captured_thresholds == [0.35, 0.35]
        assert observed_worker_names >= {
            "local-worker-api-auto-1",
            "local-worker-api-auto-2",
        }
        assert Path(last_payload["task"]["summary_export_path"]).name == "task_summary.csv"
        assert Path(last_payload["task"]["review_export_path"]).name == "task_review_rows.csv"
        assert Path(last_payload["task"]["result_json_path"]).name == "task_summary.json"

        paged_detail_response = client.get(
            f"/api/v1/tasks/{task_id}",
            params={"item_limit": 1, "item_offset": 1},
        )
        assert paged_detail_response.status_code == 200
        paged_payload = paged_detail_response.json()
        assert paged_payload["item_total"] == 2
        assert paged_payload["result_stats"]["high_risk_count"] == 2
        assert len(paged_payload["items"]) == 1
        assert paged_payload["items"][0]["item_order"] == 2

        summary_response = client.get(f"/api/v1/tasks/{task_id}/exports/summary")
        assert summary_response.status_code == 200
        assert summary_response.headers["content-type"].startswith("text/csv")
        assert 'filename="task_summary.csv"' in summary_response.headers["content-disposition"]

        review_response = client.get(f"/api/v1/tasks/{task_id}/exports/review")
        assert review_response.status_code == 200
        assert review_response.headers["content-type"].startswith("text/csv")
        assert 'filename="task_review_rows.csv"' in review_response.headers["content-disposition"]

        json_response = client.get(f"/api/v1/tasks/{task_id}/exports/json")
        assert json_response.status_code == 200
        assert json_response.headers["content-type"].startswith("application/json")
        assert 'filename="task_summary.json"' in json_response.headers["content-disposition"]
        json_payload = json_response.json()
        assert json_payload["task_id"] == task_id
        assert json_payload["status"] == "completed"
        assert len(json_payload["items"]) == 2

        unsupported_export_response = client.get(f"/api/v1/tasks/{task_id}/exports/unknown")
        assert unsupported_export_response.status_code == 400
        assert unsupported_export_response.json()["detail"] == "unsupported export kind"

        system_status_response = client.get("/api/v1/system/status")
        assert system_status_response.status_code == 200
        system_payload = system_status_response.json()["payload"]
        assert len(system_payload["cards"]) >= 1
        assert len(system_payload["services"]) >= 1
        assert "recent_exceptions" in system_payload
        assert "slow_tasks" in system_payload
        assert any(card["label"] == "Platform Health" and card["value"] == "healthy" for card in system_payload["cards"])
        assert any(card["label"] == "Stored Results" and card["value"] == "2" for card in system_payload["cards"])
        runtime_paths_service = next(service for service in system_payload["services"] if service["title"] == "Runtime Paths")
        runtime_path_items = {label: value for label, value in runtime_paths_service["items"]}
        assert runtime_path_items["export root"] == str(export_root)
        assert runtime_path_items["upload root"] == str(upload_root)
        assert system_payload["recent_exceptions"][0]["type"] == "normal"
        assert any(task["id"] == task_id and task["status"] == "completed" for task in system_payload["slow_tasks"])


def test_api_result_review_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business.sqlite3"
    retrieval_db = tmp_path / "retrieval.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_COUNT", "2")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_POLL_SECONDS", "0.2")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    runtime_worker = reload(runtime_worker)
    api_app = reload(api_app)

    from service.task_executor import ComparePipelineRequest
    from service.task_executor import execute_claimed_task as real_execute_claimed_task

    def fake_pipeline(request: ComparePipelineRequest) -> Dict[str, Any]:
        return {
            "query_text": request.query_text,
            "detection_mode": request.detection_mode,
            "rewrite_detection": {
                "status": "fallback_lexical_only",
            },
            "fine": {
                "results": [
                    {
                        "book_name": "Smoke Novel",
                        "chapter_name": "Chapter 1",
                        "review_label": "强证据",
                        "confidence_label": "strong",
                        "fine_score": 0.91,
                    }
                ],
                "review_rows": [
                    {
                        "detection_mode": request.detection_mode,
                        "fine_rank": 1,
                        "review_label": "强证据",
                        "confidence_label": "strong",
                        "fine_score": 0.91,
                        "coarse_rank": 1,
                        "coarse_final_score": 0.83,
                        "dataset_key": "smoke_dataset",
                        "book_ext_id": "book-1",
                        "book_name": "Smoke Novel",
                        "chapter_uid": 1,
                        "chapter_ext_id": "chapter-1",
                        "chapter_name": "Chapter 1",
                        "candidate_window_order": 1,
                        "candidate_start_offset": 0,
                        "candidate_end_offset": 128,
                        "exact_substring_hit": True,
                        "longest_match_len": len(request.query_text),
                        "longest_match_ratio": 1.0,
                        "ngram_recall": 1.0,
                        "ngram_precision": 1.0,
                        "jaccard": 1.0,
                        "sequence_ratio": 1.0,
                        "matched_substring": request.query_text,
                        "query_text_preview": request.query_text,
                        "candidate_text_preview": request.query_text,
                        "query_text": request.query_text,
                        "candidate_text": request.query_text,
                    }
                ],
            },
        }

    def execute_with_fake_pipeline(**kwargs: Any) -> Optional[Dict[str, Any]]:
        return real_execute_claimed_task(compare_fn=fake_pipeline, **kwargs)

    monkeypatch.setattr(runtime_worker, "execute_claimed_task", execute_with_fake_pipeline)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.txt", b"first query\n\nsecond query", "text/plain")},
            data={"detection_mode": "rewrite"},
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        deadline = time.time() + 10
        while time.time() < deadline:
            detail_response = client.get(f"/api/v1/tasks/{task_id}")
            assert detail_response.status_code == 200
            task_payload = detail_response.json()
            if task_payload["task"]["status"] == "completed":
                break
            time.sleep(0.2)
        assert task_payload["item_total"] == 2
        assert task_payload["result_stats"]["high_risk_count"] == 2
        assert task_payload["result_stats"]["semantic_fallback_count"] == 2

        result_list_response = client.get(
            "/api/v1/results",
            params={"task_id": task_id, "status": "completed"},
        )
        assert result_list_response.status_code == 200
        listed_results = result_list_response.json()["items"]
        assert len(listed_results) == 2
        assert listed_results[0]["review"]["review_status"] == ""

        pending_response = client.get(
            "/api/v1/results",
            params={"task_id": task_id, "review_status": "pending"},
        )
        assert pending_response.status_code == 200
        pending_items = pending_response.json()["items"]
        assert len(pending_items) == 2

        result_id = listed_results[0]["result_id"]
        review_response = client.post(
            f"/api/v1/results/{result_id}/review",
            json={
                "review_status": "confirmed_high_risk",
                "reviewer_name": "pytest",
                "review_note": "smoke review note",
            },
        )
        assert review_response.status_code == 200
        reviewed_result = review_response.json()["result"]
        assert reviewed_result["review"]["review_status"] == "confirmed_high_risk"
        assert reviewed_result["review"]["reviewer_name"] == "管理员"
        assert reviewed_result["review"]["review_note"] == "smoke review note"

        filtered_response = client.get(
            "/api/v1/results",
            params={"task_id": task_id, "review_status": "confirmed_high_risk"},
        )
        assert filtered_response.status_code == 200
        filtered_items = filtered_response.json()["items"]
        assert len(filtered_items) == 1
        assert filtered_items[0]["result_id"] == result_id

        detail_response = client.get(f"/api/v1/results/{result_id}")
        assert detail_response.status_code == 200
        result_detail = detail_response.json()["result"]
        assert result_detail["review"]["review_status"] == "confirmed_high_risk"
        assert result_detail["review"]["reviewer_name"] == "管理员"
        fine_results = result_detail["result_payload"]["fine"]["results"]
        assert fine_results
        first_result = fine_results[0]
        if "best_match" in first_result:
            best_match = first_result["best_match"]
            assert isinstance(best_match["candidate_text"], str)
            assert "candidate_review_context_text" in best_match
            assert "candidate_text_full" in best_match

        pending_review_response = client.post(
            f"/api/v1/results/{listed_results[1]['result_id']}/review",
            json={
                "review_status": "pending",
                "reviewer_name": "pytest",
                "review_note": "kept pending",
            },
        )
        assert pending_review_response.status_code == 200

        pending_after_save_response = client.get(
            "/api/v1/results",
            params={"task_id": task_id, "review_status": "pending"},
        )
        assert pending_after_save_response.status_code == 200
        pending_after_save_items = pending_after_save_response.json()["items"]
        assert len(pending_after_save_items) == 1
        assert pending_after_save_items[0]["result_id"] == listed_results[1]["result_id"]


def test_api_review_export_xlsx_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business.sqlite3"
    retrieval_db = tmp_path / "retrieval.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "1")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_COUNT", "2")
    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOWORKER_POLL_SECONDS", "0.2")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    runtime_worker = reload(runtime_worker)
    api_app = reload(api_app)

    from service.task_executor import ComparePipelineRequest
    from service.task_executor import execute_claimed_task as real_execute_claimed_task

    def fake_pipeline(request: ComparePipelineRequest) -> Dict[str, Any]:
        candidate_context = "这里是命中章节的长证据文本，用于后续侵权复核。" * 6
        query_text = request.query_text
        return {
            "query_text": query_text,
            "detection_mode": request.detection_mode,
            "rewrite_detection": {
                "status": "semantic_ready",
            },
            "fine": {
                "results": [
                    {
                        "fine_rank": 1,
                        "book_name": "云中证据小说",
                        "chapter_name": "第一章",
                        "review_label": "高风险",
                        "confidence_label": "strong",
                        "fine_score": 0.9321,
                        "best_match": {
                            "candidate_window_order": 3,
                            "candidate_start_offset": 128,
                            "candidate_end_offset": 456,
                            "exact_substring_hit": True,
                            "longest_match_len": len(query_text),
                            "longest_match_ratio": 0.96,
                            "ngram_recall": 0.94,
                            "ngram_precision": 0.92,
                            "jaccard": 0.91,
                            "sequence_ratio": 0.95,
                            "matched_substring": query_text[: min(len(query_text), 12)],
                            "query_text_preview": query_text[:40],
                            "candidate_text_preview": candidate_context[:80],
                            "query_text": query_text,
                            "candidate_text": candidate_context,
                            "candidate_text_full": candidate_context + " FULL",
                            "candidate_review_context_text": candidate_context + " CONTEXT",
                        },
                    }
                ],
                "review_rows": [
                    {
                        "detection_mode": request.detection_mode,
                        "fine_rank": 1,
                        "review_label": "高风险",
                        "confidence_label": "strong",
                        "fine_score": 0.9321,
                        "coarse_rank": 1,
                        "coarse_final_score": 0.8512,
                        "dataset_key": "smoke_dataset",
                        "book_ext_id": "book-1",
                        "book_name": "云中证据小说",
                        "chapter_uid": 1,
                        "chapter_ext_id": "chapter-1",
                        "chapter_name": "第一章",
                        "candidate_window_order": 3,
                        "candidate_start_offset": 128,
                        "candidate_end_offset": 456,
                        "exact_substring_hit": True,
                        "longest_match_len": len(query_text),
                        "longest_match_ratio": 0.96,
                        "ngram_recall": 0.94,
                        "ngram_precision": 0.92,
                        "jaccard": 0.91,
                        "sequence_ratio": 0.95,
                        "matched_substring": query_text[: min(len(query_text), 12)],
                        "query_text_preview": query_text[:40],
                        "candidate_text_preview": candidate_context[:80],
                        "query_text": query_text,
                        "candidate_text": candidate_context,
                    }
                ],
            },
        }

    def execute_with_fake_pipeline(**kwargs: Any) -> Optional[Dict[str, Any]]:
        return real_execute_claimed_task(compare_fn=fake_pipeline, **kwargs)

    monkeypatch.setattr(runtime_worker, "execute_claimed_task", execute_with_fake_pipeline)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        batch_csv = (
            "short_drama,episode,author,display_title,description,query_text\n"
            "短剧A,1,作者甲,标题A,描述A,第一条长台词文本用于复核导出\n"
            "短剧B,2,作者乙,标题B,描述B,第二条长台词文本用于误报归档\n"
            "短剧C,3,作者丙,标题C,描述C,第三条长台词文本用于继续跟进\n"
        ).encode("utf-8")

        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.csv", batch_csv, "text/csv")},
            data={"detection_mode": "rewrite"},
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        deadline = time.time() + 10
        listed_results: list[dict[str, Any]] = []
        while time.time() < deadline:
            result_list_response = client.get(
                "/api/v1/results",
                params={"task_id": task_id, "status": "completed"},
            )
            assert result_list_response.status_code == 200
            listed_results = result_list_response.json()["items"]
            if len(listed_results) == 3:
                break
            time.sleep(0.2)

        assert len(listed_results) == 3
        listed_results = sorted(listed_results, key=lambda item: int(item["item_order"]))
        review_states = ["confirmed_high_risk", "false_positive", "needs_followup"]
        for item, review_status in zip(listed_results, review_states):
            review_response = client.post(
                f"/api/v1/results/{item['result_id']}/review",
                json={
                    "review_status": review_status,
                    "review_note": f"note for {review_status}",
                },
            )
            assert review_response.status_code == 200

        export_response = client.get(
            "/api/v1/results/exports/review-xlsx",
            params={
                "task_id": task_id,
                "q": "短剧",
                "dedupe_latest": "true",
            },
        )
        assert export_response.status_code == 200
        assert export_response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        content_disposition = unquote(export_response.headers["content-disposition"])
        assert "复核结果导出" in content_disposition

        workbook = load_workbook(filename=BytesIO(export_response.content), data_only=True)
        assert workbook.sheetnames[:4] == ["导出说明", "确认高风险", "继续跟进", "标记误报"]

        summary_sheet = workbook["导出说明"]
        assert summary_sheet["A1"].value == "字段"
        summary_map = {
            summary_sheet.cell(row=row_index, column=1).value: summary_sheet.cell(row=row_index, column=2).value
            for row_index in range(2, summary_sheet.max_row + 1)
        }
        assert summary_map["导出范围"] == "当前登录账号下的复核结果导出"
        assert summary_map["仅导出已处理结果"] == "是"

        high_risk_sheet = workbook["确认高风险"]
        headers = [cell.value for cell in high_risk_sheet[1]]
        high_risk_rows = [
            {headers[index]: high_risk_sheet.cell(row=row_index, column=index + 1).value for index in range(len(headers))}
            for row_index in range(2, high_risk_sheet.max_row + 1)
        ]
        values = next(row for row in high_risk_rows if row["短剧名"] == "短剧A")
        assert values["短剧名"] == "短剧A"
        assert values["作者"] == "作者甲"
        assert values["复核状态"] == "确认高风险"
        assert values["复核备注"] == "note for confirmed_high_risk"
        assert str(values["查询文本证据"]).startswith("第一条长台词文本")
        assert "命中章节的长证据文本" in str(values["候选文本证据"])

        followup_sheet = workbook["继续跟进"]
        followup_headers = [cell.value for cell in followup_sheet[1]]
        followup_rows = [
            {
                followup_headers[index]: followup_sheet.cell(row=row_index, column=index + 1).value
                for index in range(len(followup_headers))
            }
            for row_index in range(2, followup_sheet.max_row + 1)
        ]
        followup_values = next(row for row in followup_rows if row["短剧名"] == "短剧C")
        assert followup_values["复核状态"] == "继续跟进"
        assert followup_values["短剧名"] == "短剧C"

        false_positive_sheet = workbook["标记误报"]
        false_positive_headers = [cell.value for cell in false_positive_sheet[1]]
        false_positive_rows = [
            {
                false_positive_headers[index]: false_positive_sheet.cell(row=row_index, column=index + 1).value
                for index in range(len(false_positive_headers))
            }
            for row_index in range(2, false_positive_sheet.max_row + 1)
        ]
        false_positive_values = next(row for row in false_positive_rows if row["短剧名"] == "短剧B")
        assert false_positive_values["复核状态"] == "标记误报"
        assert false_positive_values["短剧名"] == "短剧B"


def test_api_cancel_queued_task_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business.sqlite3"
    retrieval_db = tmp_path / "retrieval.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "0")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    reload(runtime_worker)
    api_app = reload(api_app)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.txt", b"first query\n\nsecond query", "text/plain")},
            data={"detection_mode": "rewrite"},
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        cancel_response = client.post(f"/api/v1/tasks/{task_id}/cancel")
        assert cancel_response.status_code == 200
        payload = cancel_response.json()
        assert payload["task"]["task_id"] == task_id
        assert payload["task"]["status"] == "cancelled"

        detail_response = client.get(f"/api/v1/tasks/{task_id}")
        assert detail_response.status_code == 200
        detail = detail_response.json()
        assert detail["task"]["status"] == "cancelled"

        missing_export_response = client.get(f"/api/v1/tasks/{task_id}/exports/summary")
        assert missing_export_response.status_code == 404
        assert missing_export_response.json()["detail"] == "export file not found"


def test_api_retry_cancelled_task_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business.sqlite3"
    retrieval_db = tmp_path / "retrieval.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "0")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    reload(runtime_worker)
    api_app = reload(api_app)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.txt", b"first query\n\nsecond query", "text/plain")},
            data={"detection_mode": "rewrite"},
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        cancel_response = client.post(f"/api/v1/tasks/{task_id}/cancel")
        assert cancel_response.status_code == 200

        retry_response = client.post(f"/api/v1/tasks/{task_id}/retry")
        assert retry_response.status_code == 200
        retry_payload = retry_response.json()
        assert retry_payload["task_id"] != task_id
        assert retry_payload["status"] == "queued"
        assert retry_payload["detection_mode"] == "rewrite"

        new_detail_response = client.get(f"/api/v1/tasks/{retry_payload['task_id']}")
        assert new_detail_response.status_code == 200
        new_detail = new_detail_response.json()
        assert new_detail["task"]["status"] == "queued"
        assert new_detail["task"]["source_file_name"] == "batch.txt"


def test_api_pause_resume_and_delete_task_smoke(tmp_path: Path, monkeypatch) -> None:
    business_db = tmp_path / "web_business_pause_api.sqlite3"
    retrieval_db = tmp_path / "retrieval_pause_api.sqlite3"
    upload_root = tmp_path / "uploads"
    export_root = tmp_path / "exports"

    monkeypatch.setenv("NOVEL_SIMILARITY_AUTOSTART_WORKER", "0")
    monkeypatch.setenv("NOVEL_SIMILARITY_BUSINESS_DB", str(business_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_DB", str(retrieval_db))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_UPLOAD_ROOT", str(upload_root))
    monkeypatch.setenv("NOVEL_SIMILARITY_TASK_EXPORT_ROOT", str(export_root))

    import api.config as api_config
    import api.runtime_worker as runtime_worker
    import api.app as api_app

    reload(api_config)
    reload(runtime_worker)
    api_app = reload(api_app)

    with TestClient(api_app.app) as client:
        login_as_admin(client)
        create_response = client.post(
            "/api/v1/tasks",
            files={"file": ("batch.txt", b"first query\n\nsecond query", "text/plain")},
            data={"detection_mode": "rewrite"},
        )
        assert create_response.status_code == 200
        task_id = create_response.json()["task_id"]

        pause_response = client.post(f"/api/v1/tasks/{task_id}/pause")
        assert pause_response.status_code == 200
        assert pause_response.json()["task"]["status"] == "paused"

        resume_response = client.post(f"/api/v1/tasks/{task_id}/resume")
        assert resume_response.status_code == 200
        assert resume_response.json()["task"]["status"] == "queued"

        delete_response = client.delete(f"/api/v1/tasks/{task_id}")
        assert delete_response.status_code == 200
        assert delete_response.json()["task"]["is_deleted"] is True

        detail_response = client.get(f"/api/v1/tasks/{task_id}")
        assert detail_response.status_code == 404

        list_response = client.get("/api/v1/tasks")
        assert list_response.status_code == 200
        assert list_response.json()["items"] == []
