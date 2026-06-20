from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from service.task_input_parser import parse_task_input_file


def test_parse_standard_batch_input_xlsx(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "batch_input"
    sheet.append(["source_ref", "excel_row", "novel_name", "short_drama", "query_text"])
    sheet.append(["row_a", "12", "原著A", "短剧A", "第一条待检测文本"])
    sheet.append(["row_b", "18", "原著B", "短剧B", "第二条待检测文本"])

    file_path = tmp_path / "standard_batch.xlsx"
    workbook.save(file_path)

    parsed = parse_task_input_file(file_path)

    assert len(parsed) == 2
    assert parsed[0].source_ref == "row_a"
    assert parsed[0].source_excel_row == "12"
    assert parsed[0].source_novel_name == "原著A"
    assert parsed[0].source_short_drama == "短剧A"
    assert parsed[0].source_episode == ""
    assert parsed[0].source_author == ""
    assert parsed[0].query_text == "第一条待检测文本"
    assert parsed[1].source_ref == "row_b"
    assert parsed[1].query_text == "第二条待检测文本"


def test_parse_patrol_export_summary_sheet_xlsx(tmp_path: Path) -> None:
    workbook = Workbook()
    info_sheet = workbook.active
    info_sheet.title = "导出说明"
    info_sheet.append(["字段", "值"])
    info_sheet.append(["导出类型", "当前批次"])

    summary_sheet = workbook.create_sheet("结果总表")
    summary_sheet.append(
        [
            "来源",
            "状态",
            "平台",
            "剧名",
            "集数",
            "展示标题",
            "描述",
            "分享链接",
            "作者",
            "视频标题",
            "视频文件",
            "音频文件",
            "段落数",
            "字幕全文",
            "TXT文件",
            "SRT文件",
            "JSON文件",
            "错误",
            "记录时间",
        ]
    )
    summary_sheet.append(
        [
            "当前批次",
            "success",
            "douyin",
            "短剧甲",
            1,
            "标题甲",
            "描述甲",
            "https://example.com/a",
            "作者甲",
            "视频甲",
            "a.mp4",
            "a.wav",
            12,
            "这是第一条整段字幕全文",
            "a.txt",
            "a.srt",
            "a.json",
            "",
            "2026-05-12 10:00:00",
        ]
    )
    summary_sheet.append(
        [
            "当前批次",
            "success",
            "kuaishou",
            "短剧乙",
            2,
            "标题乙",
            "描述乙",
            "https://example.com/b",
            "作者乙",
            "视频乙",
            "b.mp4",
            "b.wav",
            8,
            "这是第二条整段字幕全文",
            "b.txt",
            "b.srt",
            "b.json",
            "",
            "2026-05-12 10:05:00",
        ]
    )

    detail_sheet = workbook.create_sheet("分段明细")
    detail_sheet.append(["来源", "剧名", "集数", "展示标题", "分享链接", "段号", "开始毫秒", "结束毫秒", "文本", "JSON文件"])
    detail_sheet.append(["当前批次", "短剧甲", 1, "标题甲", "https://example.com/a", 1, 0, 1000, "片段一", "a.json"])

    file_path = tmp_path / "patrol_export.xlsx"
    workbook.save(file_path)

    parsed = parse_task_input_file(file_path)

    assert len(parsed) == 2
    assert parsed[0].source_ref == "https://example.com/a"
    assert parsed[0].source_short_drama == "短剧甲"
    assert parsed[0].source_episode == "1"
    assert parsed[0].source_author == "作者甲"
    assert parsed[0].source_platform == "douyin"
    assert parsed[0].source_display_title == "标题甲"
    assert parsed[0].source_description == "描述甲"
    assert parsed[0].query_text == "这是第一条整段字幕全文"
    assert parsed[1].source_ref == "https://example.com/b"
    assert parsed[1].source_short_drama == "短剧乙"
    assert parsed[1].source_episode == "2"
    assert parsed[1].source_author == "作者乙"
    assert parsed[1].query_text == "这是第二条整段字幕全文"


def test_parse_headerless_inline_text_xlsx(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append([1, "1,Title One,This is the first long query text used to verify headerless inline parsing works correctly."])
    sheet.append([2, "2,Title Two,This is the second long query text used to verify headerless inline parsing works correctly."])
    sheet.append([3, "This is the third long query text without inline metadata, and it should still keep the numeric source reference."])

    file_path = tmp_path / "headerless_inline.xlsx"
    workbook.save(file_path)

    parsed = parse_task_input_file(file_path)

    assert len(parsed) == 3
    assert parsed[0].source_ref == "1"
    assert parsed[0].source_display_title == "Title One"
    assert parsed[0].query_text == "This is the first long query text used to verify headerless inline parsing works correctly."
    assert parsed[1].source_ref == "2"
    assert parsed[1].source_display_title == "Title Two"
    assert parsed[1].query_text == "This is the second long query text used to verify headerless inline parsing works correctly."
    assert parsed[2].source_ref == "3"
    assert parsed[2].source_display_title == ""
    assert parsed[2].query_text == "This is the third long query text without inline metadata, and it should still keep the numeric source reference."


def test_parse_xlsx_with_sequence_and_fragment_content_headers(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["\u5e8f\u53f7", "\u5c0f\u8bf4\u7247\u6bb5\u5185\u5bb9"])
    sheet.append([1, "First fragment body for parser validation."])
    sheet.append([2, "Second fragment body for parser validation."])

    file_path = tmp_path / "fragment_content.xlsx"
    workbook.save(file_path)

    parsed = parse_task_input_file(file_path)

    assert len(parsed) == 2
    assert parsed[0].source_ref == "1"
    assert parsed[0].query_text == "First fragment body for parser validation."
    assert parsed[1].source_ref == "2"
    assert parsed[1].query_text == "Second fragment body for parser validation."
