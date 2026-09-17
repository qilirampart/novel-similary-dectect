from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sqlite3

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from api.cover_routes import build_cover_monitor_router
from service.cover_monitor.store import CoverAccessScope, CoverMonitorStore
from service.cover_monitor.storage import LocalCoverAssetStorage, RoutedCoverAssetStorage


class OssNamedLocalStorage(LocalCoverAssetStorage):
    backend_name = "oss"


def test_cover_overview_uses_independent_empty_database(tmp_path: Path) -> None:
    db_path = tmp_path / "cover-monitor.sqlite3"
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(db_path),
            asset_storage=LocalCoverAssetStorage(tmp_path / "assets"),
            current_user_dependency=lambda: {"user_id": 7, "role": "operator"},
        )
    )

    response = TestClient(app).get("/api/v1/cover-monitor/overview")

    assert response.status_code == 200
    assert response.json() == {
        "channel_count": 0,
        "video_count": 0,
        "risk_count": 0,
        "pending_review_count": 0,
        "risk_distribution": {"safe": 0, "review": 0, "risk": 0, "unknown": 0},
        "latest_run": None,
    }
    results = TestClient(app).get("/api/v1/cover-monitor/results")
    assert results.status_code == 200
    assert results.json() == {
        "items": [],
        "total": 0,
        "limit": 20,
        "offset": 0,
        "counts": {"all": 0, "risk": 0, "review": 0, "unknown": 0},
    }
    assert TestClient(app).get(
        "/api/v1/cover-monitor/results?overall_risk=safe"
    ).status_code == 422
    assert db_path.is_file()


def _channel_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["代理商简称", "频道 ID", "频道名", "频道链接"])
    worksheet.append([
        "代理甲",
        "UC-route-1",
        "接口测试频道",
        "https://www.youtube.com/channel/UC-route-1",
    ])
    workbook.save(path)


def test_cover_import_preview_and_confirm_routes(tmp_path: Path) -> None:
    db_path = tmp_path / "cover-monitor.sqlite3"
    import_root = tmp_path / "imports"
    source = tmp_path / "channels.xlsx"
    _channel_workbook(source)
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(db_path),
            import_root=str(import_root),
            current_user_dependency=lambda: {"user_id": 7, "role": "operator"},
        )
    )
    client = TestClient(app)

    with source.open("rb") as upload:
        preview_response = client.post(
            "/api/v1/cover-monitor/imports/preview",
            data={"import_kind": "channels"},
            files={"file": (source.name, upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert preview["status"] == "previewed"
    assert preview["stats"]["unique_channels"] == 1
    confirm_response = client.post(
        f"/api/v1/cover-monitor/imports/{preview['import_id']}/confirm"
    )
    assert confirm_response.status_code == 200
    assert confirm_response.json()["stats"]["applied_channels"] == 1
    assert client.get("/api/v1/cover-monitor/overview").json()["channel_count"] == 1


def test_cover_channel_route_filters_and_paginates_with_operational_counts(tmp_path: Path) -> None:
    db_path = tmp_path / "cover-monitor.sqlite3"
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(db_path),
            current_user_dependency=lambda: {"user_id": 7, "role": "operator"},
        )
    )
    store = CoverMonitorStore(db_path)
    scope = CoverAccessScope(workspace_key="internal", user_id=7)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO cover_operators (
                workspace_key, external_id, name, active, created_by_user_id, created_at, updated_at
            ) VALUES ('internal', 'agency-a', '代理甲', 1, 7, '2026-09-12T00:00:00+00:00', '2026-09-12T00:00:00+00:00')
            """
        )
        operator_pk = int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])
    alpha = store.upsert_channel(
        scope,
        platform="youtube",
        channel_id="UC-alpha",
        name="Alpha 剧场",
        source_url="https://www.youtube.com/channel/UC-alpha",
        operator_pk=operator_pk,
    )
    store.upsert_channel(
        scope,
        platform="youtube",
        channel_id="UC-beta",
        name="Beta 剧场",
        source_url="https://www.youtube.com/channel/UC-beta",
    )
    store.upsert_video(
        scope,
        channel_pk=alpha["channel_pk"],
        platform="youtube",
        video_id="video-alpha-1",
        title="Alpha video",
        video_url="https://www.youtube.com/watch?v=video-alpha-1",
        thumbnail_url="https://i.ytimg.com/vi/video-alpha-1/hqdefault.jpg",
    )

    response = TestClient(app).get(
        "/api/v1/cover-monitor/channels?keyword=alpha&active=true&limit=1&offset=0"
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 1
    assert payload["limit"] == 1
    assert payload["offset"] == 0
    assert payload["items"][0] == {
        "channel_pk": alpha["channel_pk"],
        "platform": "youtube",
        "channel_id": "UC-alpha",
        "name": "Alpha 剧场",
        "source_url": "https://www.youtube.com/channel/UC-alpha",
        "active": True,
        "operator_pk": operator_pk,
        "operator_name": "代理甲",
        "video_count": 1,
        "open_case_count": 0,
        "last_scan_at": None,
        "latest_scan_status": None,
        "latest_scan_completeness": None,
        "updated_at": alpha["updated_at"],
    }
    selected = TestClient(app).get(
        "/api/v1/cover-monitor/channels/ids?keyword=alpha&active=true"
    )
    assert selected.status_code == 200
    assert selected.json() == {
        "channel_pks": [alpha["channel_pk"]],
        "total": 1,
        "truncated": False,
    }
    assigned = TestClient(app).get(
        f"/api/v1/cover-monitor/channels?operator_pk={operator_pk}"
    )
    assert assigned.status_code == 200
    assert assigned.json()["total"] == 1
    assert assigned.json()["items"][0]["channel_pk"] == alpha["channel_pk"]
    assigned_ids = TestClient(app).get(
        f"/api/v1/cover-monitor/channels/ids?operator_pk={operator_pk}"
    )
    assert assigned_ids.status_code == 200
    assert assigned_ids.json() == {
        "channel_pks": [alpha["channel_pk"]],
        "total": 1,
        "truncated": False,
    }
    unassigned = TestClient(app).get(
        "/api/v1/cover-monitor/channels?operator_pk=-1"
    )
    assert unassigned.status_code == 200
    assert unassigned.json()["total"] == 1
    assert unassigned.json()["items"][0]["channel_id"] == "UC-beta"
    filters = TestClient(app).get(
        f"/api/v1/cover-monitor/filter-options?operator_pk={operator_pk}"
    )
    assert filters.status_code == 200
    assert filters.json()["operators"] == [
        {"operator_pk": operator_pk, "name": "代理甲", "channel_count": 1},
        {"operator_pk": -1, "name": "未分配", "channel_count": 1},
    ]
    assert filters.json()["channels"] == [
        {
            "channel_pk": alpha["channel_pk"],
            "channel_id": "UC-alpha",
            "name": "Alpha 剧场",
            "operator_pk": operator_pk,
        }
    ]
    assert TestClient(app).get(
        "/api/v1/cover-monitor/results?operator_pk=0"
    ).status_code == 422
    assert TestClient(app).get(
        "/api/v1/cover-monitor/channels?operator_pk=0"
    ).status_code == 422

    deactivated = TestClient(app).post(
        "/api/v1/cover-monitor/channels/deactivate",
        json={"channel_pks": [alpha["channel_pk"], alpha["channel_pk"], 999999]},
    )
    assert deactivated.status_code == 200
    assert deactivated.json() == {
        "requested_count": 2,
        "deactivated_count": 1,
        "already_inactive_count": 0,
        "not_found_count": 1,
    }
    assert TestClient(app).get(
        "/api/v1/cover-monitor/channels?active=true"
    ).json()["total"] == 1
    assert TestClient(app).get(
        "/api/v1/cover-monitor/channels?active=false"
    ).json()["items"][0]["channel_pk"] == alpha["channel_pk"]
    with sqlite3.connect(db_path) as conn:
        assert conn.execute(
            "SELECT active FROM cover_channels WHERE channel_pk = ?",
            (alpha["channel_pk"],),
        ).fetchone()[0] == 0
        assert conn.execute(
            "SELECT COUNT(*) FROM cover_videos WHERE channel_pk = ?",
            (alpha["channel_pk"],),
        ).fetchone()[0] == 1


def test_cover_import_rejects_non_xlsx_and_oversized_files(tmp_path: Path) -> None:
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(tmp_path / "cover.sqlite3"),
            import_root=str(tmp_path / "imports"),
            import_max_bytes=4,
            current_user_dependency=lambda: {"user_id": 7},
        )
    )
    client = TestClient(app)

    invalid = client.post(
        "/api/v1/cover-monitor/imports/preview",
        files={"file": ("channels.csv", b"a,b", "text/csv")},
    )
    oversized = client.post(
        "/api/v1/cover-monitor/imports/preview",
        files={"file": ("channels.xlsx", b"12345", "application/octet-stream")},
    )

    assert invalid.status_code == 400
    assert oversized.status_code == 413


def test_cover_import_rejects_corrupt_xlsx_as_client_error(tmp_path: Path) -> None:
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(tmp_path / "cover.sqlite3"),
            import_root=str(tmp_path / "imports"),
            current_user_dependency=lambda: {"user_id": 7},
        )
    )

    response = TestClient(app).post(
        "/api/v1/cover-monitor/imports/preview",
        files={"file": ("broken.xlsx", b"not-an-xlsx", "application/octet-stream")},
    )

    assert response.status_code == 400


def _cover_run_client(
    tmp_path: Path,
    *,
    asset_storage=None,
) -> tuple[TestClient, Path]:
    db_path = tmp_path / "cover.sqlite3"
    app = FastAPI()
    app.include_router(
        build_cover_monitor_router(
            db_path=str(db_path),
            asset_root=str(tmp_path / "assets"),
            asset_storage=asset_storage,
            current_user_dependency=lambda: {"user_id": 7, "role": "operator"},
            vision_provider="vision.example.test",
            vision_model="test-vision-model",
        )
    )
    client = TestClient(app)
    store = CoverMonitorStore(db_path)
    scope = CoverAccessScope(workspace_key="internal", user_id=7)
    for index in range(1, 3):
        store.upsert_channel(
            scope,
            platform="youtube",
            channel_id=f"UC-route-run-{index}",
            name=f"巡检频道 {index}",
            source_url=f"https://www.youtube.com/channel/UC-route-run-{index}",
        )
    return client, db_path


def test_cover_run_routes_create_list_and_get_without_persisting_secrets(tmp_path: Path) -> None:
    client, db_path = _cover_run_client(tmp_path)

    created_response = client.post(
        "/api/v1/cover-monitor/runs",
        json={
            "intensity": "strict",
            "include_shorts": False,
            "force_refresh": True,
            "max_items_per_scope": 25,
        },
    )

    assert created_response.status_code == 200
    created = created_response.json()
    assert created["status"] == "queued"
    assert created["intensity"] == "strict"
    assert created["total_channel_count"] == 2
    listed = client.get("/api/v1/cover-monitor/runs?limit=10&offset=0")
    assert listed.status_code == 200
    assert listed.json()["total"] == 1
    assert listed.json()["items"][0]["run_id"] == created["run_id"]
    detail = client.get(f"/api/v1/cover-monitor/runs/{created['run_id']}")
    assert detail.status_code == 200
    assert detail.json()["run_id"] == created["run_id"]

    with sqlite3.connect(db_path) as conn:
        snapshot = conn.execute(
            "SELECT model_snapshot_json FROM cover_runs WHERE run_id = ?",
            (created["run_id"],),
        ).fetchone()[0]
    assert "test-vision-model" in snapshot
    assert "api_key" not in snapshot.casefold()


def test_cover_run_routes_pause_resume_and_cancel_with_state_conflicts(tmp_path: Path) -> None:
    client, _ = _cover_run_client(tmp_path)
    created = client.post("/api/v1/cover-monitor/runs", json={}).json()
    run_url = f"/api/v1/cover-monitor/runs/{created['run_id']}"

    paused = client.post(f"{run_url}/pause")
    assert paused.status_code == 200
    assert paused.json()["status"] == "paused"
    resumed = client.post(f"{run_url}/resume")
    assert resumed.status_code == 200
    assert resumed.json()["status"] == "queued"
    cancelled = client.post(f"{run_url}/cancel")
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    invalid_resume = client.post(f"{run_url}/resume")
    assert invalid_resume.status_code == 409
    assert client.get("/api/v1/cover-monitor/runs/missing-run").status_code == 404


def test_cover_run_create_validates_channel_scope_and_request_limits(tmp_path: Path) -> None:
    client, _ = _cover_run_client(tmp_path)

    missing_channel = client.post(
        "/api/v1/cover-monitor/runs",
        json={"channel_pks": [999999]},
    )
    invalid_limit = client.post(
        "/api/v1/cover-monitor/runs",
        json={"max_items_per_scope": 10001},
    )

    assert missing_channel.status_code == 409
    assert invalid_limit.status_code == 422


def test_cover_run_detail_returns_channel_progress_and_paginated_items(tmp_path: Path) -> None:
    client, db_path = _cover_run_client(tmp_path)
    created = client.post("/api/v1/cover-monitor/runs", json={}).json()
    store = CoverMonitorStore(db_path)
    scope = CoverAccessScope(workspace_key="internal", user_id=7)
    channel = store.list_channels(scope)[0]
    video = store.upsert_video(
        scope,
        channel_pk=channel["channel_pk"],
        platform="youtube",
        video_id="video-route-detail",
        title="Detail Video",
        video_url="https://www.youtube.com/watch?v=video-route-detail",
        thumbnail_url="https://i.ytimg.com/vi/video-route-detail/hqdefault.jpg",
    )
    claimed = store.claim_next_run(worker_name="route-detail-worker")
    assert claimed is not None
    store.enqueue_task_items(
        created["run_id"],
        claimed["worker_lease_token"],
        [{"video_pk": video["video_pk"], "reason": "manual"}],
    )

    response = client.get(
        f"/api/v1/cover-monitor/runs/{created['run_id']}/detail?item_limit=1&item_offset=0"
    )

    assert response.status_code == 200
    detail = response.json()
    assert detail["run"]["run_id"] == created["run_id"]
    assert detail["item_total"] == 1
    assert detail["item_limit"] == 1
    assert detail["channels"][0]["channel_name"] == "巡检频道 1"
    assert detail["items"][0]["video_id"] == "video-route-detail"
    assert detail["items"][0]["reason"] == "manual"
    assert client.get("/api/v1/cover-monitor/runs/missing/detail").status_code == 404


def _seed_cover_risk_candidate(db_path: Path) -> str:
    store = CoverMonitorStore(db_path)
    scope = CoverAccessScope(workspace_key="internal", user_id=7)
    channel = store.list_channels(scope)[0]
    video = store.upsert_video(
        scope,
        channel_pk=channel["channel_pk"],
        platform="youtube",
        video_id="video-route-risk",
        title="风险案件接口测试",
        video_url="https://www.youtube.com/watch?v=video-route-risk",
        thumbnail_url="https://i.ytimg.com/vi/video-route-risk/hqdefault.jpg",
    )
    for index, overall_risk in enumerate(("risk", "safe"), start=1):
        run = store.create_run(
            scope,
            trigger_type="manual",
            intensity="standard",
            channel_pks=[channel["channel_pk"]],
            params={"force_refresh": index > 1},
            model_snapshot={"provider": "fake", "model": "fake"},
            prompt_version="cover-visible-evidence-v1",
        )
        claimed = store.claim_next_run(worker_name=f"route-risk-worker-{index}")
        assert claimed is not None
        store.enqueue_task_items(
            run["run_id"],
            claimed["worker_lease_token"],
            [{"video_pk": video["video_pk"], "reason": "manual"}],
        )
        item = store.claim_next_task_item(run["run_id"], claimed["worker_lease_token"])
        assert item is not None
        content_sha256 = ("a" if index == 1 else "b") * 64
        asset = store.save_asset(
            scope,
            video_pk=video["video_pk"],
            asset={
                "content_sha256": content_sha256,
                "storage_key": f"video-route-risk/{content_sha256}.jpg",
                "original_url": video["thumbnail_url"],
                "fetched_url": video["thumbnail_url"],
                "mime_type": "image/jpeg",
                "byte_size": 2048,
                "width": 1280,
                "height": 720,
            },
        )
        asset_path = db_path.parent / "assets" / asset["storage_key"]
        asset_path.parent.mkdir(parents=True, exist_ok=True)
        asset_path.write_bytes(f"fake-cover-{index}".encode("ascii"))
        assert store.advance_task_item(
            item["task_item_id"],
            item["worker_lease_token"],
            expected_stage="download",
            next_stage="review",
        )
        store.save_detection_and_case(
            scope,
            task_item_id=item["task_item_id"],
            worker_lease_token=item["worker_lease_token"],
            asset_id=asset["asset_id"],
            detection={
                "overall_risk": overall_risk,
                "risk_tags": ["未成年人"] if overall_risk == "risk" else [],
                "summary": "发现风险元素" if overall_risk == "risk" else "新封面未见风险元素",
                "evidence": "旧封面风险证据" if overall_risk == "risk" else "新封面安全证据",
                "confidence": 0.91,
                "provider": "fake",
                "model": "fake",
                "prompt_version": "cover-visible-evidence-v1",
                "prompt_hash": "c" * 64,
                "intensity": "standard",
                "raw_response": "{}",
                "duration_seconds": 0.2,
            },
        )
        assert store.advance_task_item(
            item["task_item_id"],
            item["worker_lease_token"],
            expected_stage="review",
            next_stage="persist",
        )
        assert store.finish_task_item(
            item["task_item_id"], item["worker_lease_token"], succeeded=True
        )
    return store.list_risk_cases(scope, status="needs_review")["items"][0]["case_id"]


def test_cover_risk_case_routes_list_detail_and_confirm_rectification(tmp_path: Path) -> None:
    local = LocalCoverAssetStorage(tmp_path / "assets")
    oss = OssNamedLocalStorage(tmp_path / "oss-assets")
    routed = RoutedCoverAssetStorage(
        active=oss,
        backends={"local": local, "oss": oss},
    )
    client, db_path = _cover_run_client(tmp_path, asset_storage=routed)
    case_id = _seed_cover_risk_candidate(db_path)

    listed = client.get("/api/v1/cover-monitor/risk-cases?status=needs_review")
    detail = client.get(f"/api/v1/cover-monitor/risk-cases/{case_id}")
    reviewed = client.post(
        f"/api/v1/cover-monitor/risk-cases/{case_id}/review",
        json={
            "action": "confirm_rectified",
            "reason": "人工对照新旧封面，确认风险元素已移除",
        },
    )

    assert listed.status_code == 200
    assert client.get(
        "/api/v1/cover-monitor/risk-cases?status=confirmed_risk"
    ).status_code == 200
    assert listed.json()["total"] == 1
    assert listed.json()["items"][0]["video_id"] == "video-route-risk"
    assert detail.status_code == 200
    assert [item["event_type"] for item in detail.json()["events"]] == [
        "risk_detected",
        "rectification_candidate",
    ]
    assert {item["content_sha256"] for item in detail.json()["events"]} == {"a" * 64, "b" * 64}
    first_asset_id = detail.json()["events"][0]["asset_id"]
    evidence = client.get(
        f"/api/v1/cover-monitor/risk-cases/{case_id}/assets/{first_asset_id}"
    )
    assert evidence.status_code == 200
    assert evidence.content == b"fake-cover-1"
    assert reviewed.status_code == 200
    assert reviewed.json()["case"]["current_status"] == "confirmed_rectified"
    assert reviewed.json()["reviews"][0]["action"] == "confirm_rectified"
    assert client.post(
        f"/api/v1/cover-monitor/risk-cases/{case_id}/review",
        json={"action": "confirm_rectified", "reason": "重复确认应被拒绝"},
    ).status_code == 409
    assert client.get("/api/v1/cover-monitor/risk-cases/missing-case").status_code == 404


def test_cover_run_exports_use_snapshot_and_separate_new_from_rectification(tmp_path: Path) -> None:
    client, db_path = _cover_run_client(tmp_path)
    case_id = _seed_cover_risk_candidate(db_path)
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        run_rows = conn.execute(
            "SELECT run_id FROM cover_runs ORDER BY created_at, rowid"
        ).fetchall()
        new_run_id = str(run_rows[0]["run_id"])
        rectification_run_id = str(run_rows[-1]["run_id"])
        conn.execute(
            "UPDATE cover_task_items SET reason = 'historical_risk' WHERE run_id = ?",
            (rectification_run_id,),
        )
        conn.execute(
            """
            UPDATE cover_videos SET title = '=HYPERLINK("https://invalid.test","风险标题")'
             WHERE video_pk IN (SELECT video_pk FROM cover_task_items WHERE run_id = ?)
            """,
            (rectification_run_id,),
        )
        conn.execute(
            """
            UPDATE cover_run_channels
               SET operator_snapshot_json = '{"operator_name":"冻结代理","channel_id":"UC-snapshot","channel_name":"冻结频道","source_url":"https://example.test/frozen"}'
             WHERE run_id = ?
            """,
            (rectification_run_id,),
        )

    new_response = client.get(
        f"/api/v1/cover-monitor/runs/{new_run_id}/exports/new-findings"
    )
    rectification_response = client.get(
        f"/api/v1/cover-monitor/runs/{rectification_run_id}/exports/historical-rectification"
    )

    assert new_response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in new_response.headers["content-type"]
    new_workbook = load_workbook(BytesIO(new_response.content), read_only=True)
    assert new_workbook.sheetnames == ["新增视频明细", "新增风险", "待复核与失败", "代理商汇总"]
    assert list(next(new_workbook["新增视频明细"].iter_rows(values_only=True)))[:3] == [
        "频道 ID", "频道名", "代理商"
    ]
    new_summary_rows = list(new_workbook["代理商汇总"].iter_rows(values_only=True))
    assert new_summary_rows[1][-1] == 0

    assert rectification_response.status_code == 200
    rectification_workbook = load_workbook(BytesIO(rectification_response.content), read_only=True)
    assert rectification_workbook.sheetnames == ["历史风险整改明细", "代理商整改汇总"]
    detail_rows = list(rectification_workbook["历史风险整改明细"].iter_rows(values_only=True))
    assert detail_rows[1][:3] == ("UC-snapshot", "冻结频道", "冻结代理")
    assert str(detail_rows[1][4]).startswith("'=")
    summary_rows = list(rectification_workbook["代理商整改汇总"].iter_rows(values_only=True))
    assert summary_rows[1][0] == "冻结代理"
    assert summary_rows[1][1] == 1
    assert case_id

    missing = client.get(
        "/api/v1/cover-monitor/runs/missing-run/exports/new-findings"
    )
    assert missing.status_code == 404
