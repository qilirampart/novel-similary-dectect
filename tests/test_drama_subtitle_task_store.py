from __future__ import annotations

from pathlib import Path
import threading
import time

from openpyxl import load_workbook

from service.business_store import init_business_db, list_users
from service.drama_subtitle_task_store import (
    claim_next_drama_subtitle_task,
    create_drama_subtitle_task,
    get_drama_subtitle_task,
    list_drama_subtitle_tasks,
    update_drama_subtitle_task_control,
    upsert_drama_subtitle_task_review,
)
from service.drama_subtitle_export import build_drama_subtitle_review_export_xlsx


def test_drama_subtitle_tasks_are_owned_and_claimed_independently(tmp_path: Path) -> None:
    db_path = tmp_path / "business.sqlite3"
    init_business_db(db_path)
    admin = list_users(db_path)[0]

    task = create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-1",
        source_file_name="input.xlsx",
        source_file_ext=".xlsx",
        source_file_path="runtime/input.xlsx",
        source_file_sha256="abc",
        source_file_size=12,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
        params={"top_k": 10},
    )

    assert task["status"] == "queued"
    assert task["params"] == {"top_k": 10}
    assert list_drama_subtitle_tasks(db_path, owner_user_id=int(admin["user_id"]))[0]["task_id"] == "drama-task-1"
    assert get_drama_subtitle_task(db_path, "drama-task-1", owner_user_id=999) is None

    claimed = claim_next_drama_subtitle_task(db_path, worker_name="test-worker")

    assert claimed is not None
    assert claimed["task_id"] == "drama-task-1"
    assert claimed["status"] == "running"
    assert claim_next_drama_subtitle_task(db_path, worker_name="test-worker") is None


def test_drama_subtitle_task_pause_resume_cancel_and_soft_delete(tmp_path: Path) -> None:
    db_path = tmp_path / "business.sqlite3"
    init_business_db(db_path)
    admin = list_users(db_path)[0]
    create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-control",
        source_file_name="input.txt",
        source_file_ext=".txt",
        source_file_path="runtime/input.txt",
        source_file_sha256="abc",
        source_file_size=1,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
    )

    paused = update_drama_subtitle_task_control(
        db_path=db_path, task_id="drama-task-control", action="pause", owner_user_id=int(admin["user_id"])
    )
    assert paused is not None and paused["status"] == "paused"
    resumed = update_drama_subtitle_task_control(
        db_path=db_path, task_id="drama-task-control", action="resume", owner_user_id=int(admin["user_id"])
    )
    assert resumed is not None and resumed["status"] == "queued"
    cancelled = update_drama_subtitle_task_control(
        db_path=db_path, task_id="drama-task-control", action="cancel", owner_user_id=int(admin["user_id"])
    )
    assert cancelled is not None and cancelled["status"] == "cancelled"
    deleted = update_drama_subtitle_task_control(
        db_path=db_path, task_id="drama-task-control", action="delete", owner_user_id=int(admin["user_id"])
    )
    assert deleted is not None and deleted["is_deleted"] == 1
    assert get_drama_subtitle_task(db_path, "drama-task-control") is None


def test_drama_subtitle_executor_persists_lexical_evidence(tmp_path: Path, monkeypatch) -> None:
    db_path = tmp_path / "business.sqlite3"
    input_path = tmp_path / "subtitle.txt"
    input_path.write_text("大哥这是横店吗你们是不是在拍戏", encoding="utf-8")
    init_business_db(db_path)
    admin = list_users(db_path)[0]
    create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-execute",
        source_file_name="subtitle.txt",
        source_file_ext=".txt",
        source_file_path=str(input_path),
        source_file_sha256="abc",
        source_file_size=input_path.stat().st_size,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
        params={"top_k": 10, "window_limit": 100},
    )

    import service.drama_subtitle_task_executor as executor

    monkeypatch.setattr(
        executor,
        "search_drama_subtitle_hybrid_candidates",
        lambda **_: {
            "query_language_code": "zh",
            "query_language_confidence": 1.0,
            "decision": {"matched": True, "status": "matched", "candidate_rank": 1},
            "candidates": [
                {
                    "book_id": "41000010048",
                    "book_name": "农女致富",
                    "episode_order": 1,
                    "language_code": "zh",
                    "rank": 1,
                    "best_lexical_score": -12.3,
                    "evidence": {
                        "window_uid": "41000010048:1:1-16",
                        "time_start": "00:00:12,720",
                        "time_end": "00:00:51,920",
                    },
                }
            ],
        },
    )

    result = executor.run_next_queued_drama_subtitle_task(
        business_db_path=db_path,
        subtitle_db_path=tmp_path / "subtitle.sqlite3",
        worker_name="test-worker",
    )

    assert result is not None
    assert result["status"] == "completed"
    item = executor.list_drama_subtitle_task_items(db_path, "drama-task-execute")[0]
    assert item["status"] == "completed"
    assert item["matched_book_id"] == "41000010048"
    assert item["matched_episode_order"] == 1
    assert item["evidence_time_start"] == "00:00:12,720"


def test_drama_subtitle_executor_honors_cancel_request_between_items(tmp_path: Path, monkeypatch) -> None:
    db_path = tmp_path / "business.sqlite3"
    input_path = tmp_path / "subtitle.txt"
    input_path.write_text("第一段测试台词\n\n第二段测试台词", encoding="utf-8")
    init_business_db(db_path)
    admin = list_users(db_path)[0]
    create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-cancel",
        source_file_name="subtitle.txt",
        source_file_ext=".txt",
        source_file_path=str(input_path),
        source_file_sha256="abc",
        source_file_size=input_path.stat().st_size,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
    )

    import service.drama_subtitle_task_executor as executor

    calls = 0

    def fake_search(**_):
        nonlocal calls
        calls += 1
        update_drama_subtitle_task_control(db_path=db_path, task_id="drama-task-cancel", action="cancel")
        return {"query_language_code": "zh", "query_language_confidence": 1.0, "candidates": []}

    monkeypatch.setattr(executor, "search_drama_subtitle_hybrid_candidates", fake_search)
    result = executor.run_next_queued_drama_subtitle_task(
        business_db_path=db_path,
        subtitle_db_path=tmp_path / "subtitle.sqlite3",
        worker_name="test-worker",
    )

    assert result is not None and result["status"] == "cancelled", result
    assert calls == 1
    items = executor.list_drama_subtitle_task_items(db_path, "drama-task-cancel")
    assert [item["status"] for item in items] == ["completed", "queued"]


def test_drama_subtitle_review_and_excel_export_include_evidence(tmp_path: Path, monkeypatch) -> None:
    db_path = tmp_path / "business.sqlite3"
    input_path = tmp_path / "subtitle.txt"
    input_path.write_text("需要复核的台词", encoding="utf-8")
    init_business_db(db_path)
    admin = list_users(db_path)[0]
    create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-review",
        source_file_name="subtitle.txt",
        source_file_ext=".txt",
        source_file_path=str(input_path),
        source_file_sha256="abc",
        source_file_size=input_path.stat().st_size,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
    )

    import service.drama_subtitle_task_executor as executor

    monkeypatch.setattr(
        executor,
        "search_drama_subtitle_hybrid_candidates",
        lambda **_: {
            "query_language_code": "zh",
            "query_language_confidence": 1.0,
            "decision": {"matched": True, "status": "matched", "candidate_rank": 1},
            "candidates": [{
                "book_id": "41000010048",
                "book_name": "农女致富",
                "episode_order": 1,
                "language_code": "zh",
                "rank": 1,
                "best_lexical_score": -12.3,
                "evidence": {
                    "window_uid": "41000010048:1:1-16",
                    "time_start": "00:00:12,720",
                    "time_end": "00:00:51,920",
                    "window_text": "这是完整的候选字幕证据。",
                },
            }],
        },
    )
    executor.run_next_queued_drama_subtitle_task(
        business_db_path=db_path,
        subtitle_db_path=tmp_path / "subtitle.sqlite3",
        worker_name="test-worker",
    )
    item = executor.list_drama_subtitle_task_items(db_path, "drama-task-review")[0]
    reviewed = upsert_drama_subtitle_task_review(
        db_path=db_path,
        task_item_id=int(item["task_item_id"]),
        review_status="confirmed_high_risk",
        reviewer_name="reviewer",
        review_note="需要保留证据",
        owner_user_id=int(admin["user_id"]),
    )
    assert reviewed is not None and reviewed["review_status"] == "confirmed_high_risk"

    export_path, _ = build_drama_subtitle_review_export_xlsx(
        business_db_path=db_path,
        export_root=tmp_path / "exports",
        owner_user_id=int(admin["user_id"]),
        task_id="drama-task-review",
    )
    workbook = load_workbook(export_path, read_only=True)
    try:
        sheet = workbook["确认高风险"]
        values = list(next(sheet.iter_rows(min_row=2, values_only=True)))
        assert "需要复核的台词" in values
        assert "这是完整的候选字幕证据。" in values
        assert "需要保留证据" in values
    finally:
        workbook.close()


def test_worker_shutdown_requeues_inflight_drama_task_without_waiting(
    tmp_path: Path,
    monkeypatch,
) -> None:
    db_path = tmp_path / "business_shutdown.sqlite3"
    input_path = tmp_path / "subtitle_shutdown.txt"
    input_path.write_text("第一段测试台词\n\n第二段测试台词", encoding="utf-8")
    init_business_db(db_path)
    admin = list_users(db_path)[0]
    create_drama_subtitle_task(
        db_path=db_path,
        task_id="drama-task-shutdown",
        source_file_name=input_path.name,
        source_file_ext=".txt",
        source_file_path=str(input_path),
        source_file_sha256="shutdown-sha",
        source_file_size=input_path.stat().st_size,
        owner_user_id=int(admin["user_id"]),
        created_by="admin",
    )

    import service.drama_subtitle_task_executor as executor

    started = threading.Event()
    release = threading.Event()
    stop_event = threading.Event()
    result_holder: dict[str, dict[str, object]] = {}

    def matching_payload() -> dict[str, object]:
        return {
            "query_language_code": "zh",
            "query_language_confidence": 1.0,
            "decision": {"matched": True, "status": "matched", "candidate_rank": 1},
            "candidates": [],
        }

    def blocking_search(**_):
        started.set()
        release.wait(timeout=10)
        return matching_payload()

    monkeypatch.setattr(executor, "search_drama_subtitle_hybrid_candidates", blocking_search)

    def run_task() -> None:
        result = executor.run_next_queued_drama_subtitle_task(
            business_db_path=db_path,
            subtitle_db_path=tmp_path / "subtitle.sqlite3",
            worker_name="shutdown-worker",
            stop_event=stop_event,
        )
        assert result is not None
        result_holder["task"] = result

    worker = threading.Thread(target=run_task, daemon=True)
    worker.start()
    assert started.wait(timeout=2)
    stop_started_at = time.perf_counter()
    stop_event.set()
    worker.join(timeout=2)
    stop_elapsed = time.perf_counter() - stop_started_at

    assert not worker.is_alive()
    assert stop_elapsed < 1.5
    assert result_holder["task"]["status"] == "queued"
    items = executor.list_drama_subtitle_task_items(db_path, "drama-task-shutdown")
    assert items
    assert all(item["status"] == "queued" for item in items)

    release.set()
    monkeypatch.setattr(
        executor,
        "search_drama_subtitle_hybrid_candidates",
        lambda **_: matching_payload(),
    )
    finished = executor.run_next_queued_drama_subtitle_task(
        business_db_path=db_path,
        subtitle_db_path=tmp_path / "subtitle.sqlite3",
        worker_name="replacement-worker",
    )
    assert finished is not None
    assert finished["status"] == "completed"
    assert finished["counts"]["completed"] == 2
