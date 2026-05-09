from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_PATH = Path(
    r"E:\点众\小说库相似度比对服务工作区\数据现状说明\01_中台可查_且数据库已有\01_修正版77条结果_补台词_2026-05-09.xlsx"
)
OUTPUT_PATH = Path(
    r"E:\点众\小说库相似度比对服务工作区\data_samples\batch_uploads\batch_report_samples_26rows_2026-05-09.xlsx"
)


def main() -> None:
    wb = load_workbook(SOURCE_PATH, read_only=True, data_only=True)
    try:
        ws = wb["all69_details"]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = "batch_input"

        cols = ["source_ref", "excel_row", "novel_name", "short_drama", "query_text"]
        out_ws.append(cols)

        for row_idx in range(2, ws.max_row + 1):
            status = str(ws.cell(row_idx, headers["台词匹配状态"]).value or "").strip()
            if status != "已回填":
                continue

            source_ref = f"all69_details_row_{row_idx}"
            novel_name = ws.cell(row_idx, headers["原著小说名"]).value or ""
            short_drama = ws.cell(row_idx, headers["短剧名"]).value or ""
            query_text = ws.cell(row_idx, headers["首集台词"]).value or ""

            out_ws.append(
                [
                    source_ref,
                    row_idx,
                    str(novel_name).strip(),
                    str(short_drama).strip(),
                    str(query_text).strip(),
                ]
            )

        header_fill = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
        header_font = Font(name="Arial", bold=True, color="1F3A5F")
        body_font = Font(name="Arial")
        wrap_alignment = Alignment(vertical="top", wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center")

        for cell in out_ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment

        for row in out_ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = wrap_alignment

        out_ws.column_dimensions["A"].width = 24
        out_ws.column_dimensions["B"].width = 10
        out_ws.column_dimensions["C"].width = 26
        out_ws.column_dimensions["D"].width = 26
        out_ws.column_dimensions["E"].width = 180

        for row_idx in range(2, out_ws.max_row + 1):
            out_ws.row_dimensions[row_idx].height = 78

        out_ws.freeze_panes = "A2"
        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        out_wb.save(OUTPUT_PATH)
    finally:
        wb.close()

    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
