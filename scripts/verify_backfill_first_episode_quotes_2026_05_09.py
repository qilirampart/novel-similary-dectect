from pathlib import Path

from openpyxl import load_workbook


PATH = Path(
    r"E:\点众\小说库相似度比对服务工作区\数据现状说明\01_中台可查_且数据库已有\01_修正版77条结果_补台词_2026-05-09.xlsx"
)
QUOTE_PATH = Path(
    r"C:\Users\psk13\Desktop\实验室任务\aps\02_功能包装材料论文项目\03_检索与写作\01_主稿\补充图片\漫剧台词表.xlsx"
)


def main() -> None:
    wb = load_workbook(PATH, data_only=True)
    try:
        ws = wb["all69_details"]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        print("SHEET", ws.title, "ROWS", ws.max_row - 1)
        for row_idx in [2, 3, 4, 5, 6, 10, 11, 12, 13, 14, 15]:
            short_drama = ws.cell(row_idx, headers["短剧名"]).value
            quote = ws.cell(row_idx, headers["首集台词"]).value
            status = ws.cell(row_idx, headers["台词匹配状态"]).value
            source = ws.cell(row_idx, headers["台词来源剧名"]).value
            print(row_idx, short_drama, (quote[:60] if quote else None), status, source)

        review = wb["旧10条复核"]
        print("OLD10")
        for row_idx in range(2, review.max_row + 1):
            print(
                row_idx,
                review.cell(row_idx, 2).value,
                review.cell(row_idx, 4).value,
                review.cell(row_idx, 5).value[:30] if review.cell(row_idx, 5).value else None,
            )
    finally:
        wb.close()

    qwb = load_workbook(QUOTE_PATH, data_only=True)
    try:
        found = 0
        print("SOURCE_CHECK")
        targets = {
            "银心照温庭",
            "离婚后我活成女王",
            "重生丧尸潮，我不救恋爱脑了",
            "真瓷破假缘",
            "产房之外",
        }
        for s in qwb.sheetnames:
            ws = qwb[s]
            headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
            drama_col = headers.get("剧名")
            quote_col = headers.get("首集台词")
            if not drama_col or not quote_col:
                continue
            for row_idx in range(2, ws.max_row + 1):
                drama = ws.cell(row_idx, drama_col).value
                if drama in targets:
                    quote = ws.cell(row_idx, quote_col).value
                    print(s, row_idx, drama, repr(quote[:80] if quote else None))
                    found += 1
        print("SOURCE_CHECK_DONE", found)
    finally:
        qwb.close()


if __name__ == "__main__":
    main()
