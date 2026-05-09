from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


WORKDIR = Path(r"E:\点众\小说库相似度比对服务工作区")
CURRENT_PATH = WORKDIR / "data_samples" / "quality_review" / "real_test_titles_all69_fullsource_v1.xlsx"
AUDIT_PATH = WORKDIR / "data_samples" / "quality_review" / "intersection_audit_2026-05-08.xlsx"
CLUE_PATH = Path(
    r"C:\Users\psk13\Desktop\实验室任务\aps\02_功能包装材料论文项目\03_检索与写作\01_主稿\补充图片\漫剧线索自查表(1).xlsx"
)
OUTPUT_PATH = WORKDIR / "data_samples" / "quality_review" / "real_test_titles_all77_fullsource_v1.xlsx"
REPORT_OUTPUT_PATH = (
    WORKDIR
    / "汇报材料_数据申请准备_2026-05-08"
    / "01_中台可查_且数据库已有"
    / "01_修正版77条结果.xlsx"
)


def normalize_text(value) -> str:
    return "" if value is None else str(value).strip()


def build_header_index(ws) -> dict[str, int]:
    return {normalize_text(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)}


def load_missing_rows() -> list[dict[str, str]]:
    wb = load_workbook(AUDIT_PATH, read_only=True, data_only=True)
    try:
        ws = wb["suspected_missing"]
        headers = [normalize_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        rows: list[dict[str, str]] = []
        for r in range(2, ws.max_row + 1):
            row = {headers[c - 1]: normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
            if any(row.values()):
                rows.append(row)
        return rows
    finally:
        wb.close()


def read_clue_context():
    wb = load_workbook(CLUE_PATH, read_only=True, data_only=True)
    contexts: dict[tuple[str, int], dict[str, str]] = {}
    try:
        for item in load_missing_rows():
            sheet = item["source_sheet"]
            rownum = int(item["source_row"])
            ws = wb[sheet]
            if sheet == "实习生核查表【完】":
                row = [ws.cell(rownum, c).value for c in range(1, 15)]
                contexts[(sheet, rownum)] = {
                    "layout": "wide",
                    "seq": normalize_text(row[0]),
                    "account_info": normalize_text(row[1]),
                    "cover_or_link": normalize_text(row[2]),
                    "drama_name": normalize_text(row[3]),
                    "tag": normalize_text(row[4]),
                    "copyright_owner": normalize_text(row[5]),
                    "producer": normalize_text(row[6]),
                    "has_original": normalize_text(row[7]),
                    "original_name": normalize_text(row[8]),
                    "mid_found": normalize_text(row[9]),
                    "cp_name": normalize_text(row[10]),
                    "assignment": normalize_text(row[11]),
                    "homepage": normalize_text(row[12]),
                    "extra": normalize_text(row[13]),
                }
            else:
                row1 = [ws.cell(1, c).value for c in range(1, 8)]
                row = [ws.cell(rownum, c).value for c in range(1, 8)]
                account_info = normalize_text(row1[0])
                homepage = ""
                if "主页：" in account_info:
                    homepage = account_info.split("主页：", 1)[1].strip()
                elif "主页链接：" in account_info:
                    homepage = account_info.split("主页链接：", 1)[1].strip()
                contexts[(sheet, rownum)] = {
                    "layout": "narrow",
                    "seq": "",
                    "account_info": account_info,
                    "cover_or_link": normalize_text(row[0]),
                    "drama_name": normalize_text(row[1]),
                    "tag": "",
                    "copyright_owner": "",
                    "producer": "",
                    "has_original": normalize_text(row[2]),
                    "original_name": normalize_text(row[3]),
                    "mid_found": normalize_text(row[4]),
                    "cp_name": normalize_text(row[5]),
                    "assignment": "",
                    "homepage": homepage,
                    "extra": normalize_text(row[6]),
                }
        return contexts
    finally:
        wb.close()


def clone_row_style(ws, template_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[template_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(template_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def main() -> None:
    missing_rows = load_missing_rows()
    clue_context = read_clue_context()

    wb = load_workbook(CURRENT_PATH)
    ws = wb.active
    header_idx = build_header_index(ws)
    template_row = ws.max_row
    next_priority = ws.max_row

    for item in missing_rows:
        sheet = item["source_sheet"]
        rownum = int(item["source_row"])
        ctx = clue_context[(sheet, rownum)]
        next_row = ws.max_row + 1
        clone_row_style(ws, template_row, next_row)

        ws.cell(next_row, header_idx["测试优先级"]).value = next_priority
        ws.cell(next_row, header_idx["原著小说名"]).value = item["book_name"]
        ws.cell(next_row, header_idx["短剧名"]).value = ctx["drama_name"]
        ws.cell(next_row, header_idx["来源sheet"]).value = sheet
        ws.cell(next_row, header_idx["小说库数据集"]).value = item["dataset_key"]
        ws.cell(next_row, header_idx["小说book_ext_id"]).value = item["book_ext_id"]
        ws.cell(next_row, header_idx["库内书名"]).value = item["book_name"]
        ws.cell(next_row, header_idx["中台可查"]).value = "是"
        ws.cell(next_row, header_idx["原表布局"]).value = ctx["layout"]
        ws.cell(next_row, header_idx["原表Excel行号"]).value = rownum
        ws.cell(next_row, header_idx["原表_序号"]).value = ctx["seq"]
        ws.cell(next_row, header_idx["原表_账号或页头信息"]).value = ctx["account_info"]
        ws.cell(next_row, header_idx["原表_视频封面或短剧封面"]).value = ctx["cover_or_link"]
        ws.cell(next_row, header_idx["原表_标签"]).value = ctx["tag"]
        ws.cell(next_row, header_idx["原表_版权方"]).value = ctx["copyright_owner"]
        ws.cell(next_row, header_idx["原表_承制方"]).value = ctx["producer"]
        ws.cell(next_row, header_idx["原表_是否有原著小说"]).value = ctx["has_original"]
        ws.cell(next_row, header_idx["原表_原著小说名称"]).value = ctx["original_name"]
        ws.cell(next_row, header_idx["原表_原著小说是否在中台能查找到"]).value = ctx["mid_found"]
        ws.cell(next_row, header_idx["原表_原著小说cp名称"]).value = ctx["cp_name"]
        ws.cell(next_row, header_idx["原表_分工"]).value = ctx["assignment"]
        ws.cell(next_row, header_idx["原表_主页链接"]).value = ctx["homepage"]
        ws.cell(next_row, header_idx["原表_附加列"]).value = ctx["extra"]
        ws.cell(next_row, header_idx["入选原因"]).value = (
            "中台可查=是；与当前小说库严格同名命中；该条在 2026-05-08 自动审计中识别为首次69条结果遗漏，已补录"
        )
        ws.cell(next_row, header_idx["计划用途"]).value = "待下载对应短剧字幕后，用于 reuse / rewrite 两种模式实测"
        ws.cell(next_row, header_idx["当前状态"]).value = "待下载字幕"
        ws.cell(next_row, header_idx["备注"]).value = "2026-05-08 复查补录"

        next_priority += 1

    wb.save(OUTPUT_PATH)
    wb.save(REPORT_OUTPUT_PATH)
    print(OUTPUT_PATH)
    print(REPORT_OUTPUT_PATH)
    print(f"final_rows={ws.max_row - 1}")


if __name__ == "__main__":
    main()
