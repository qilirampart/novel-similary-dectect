from __future__ import annotations

import argparse
import csv
import json
import sys
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from v2_common import connect_db, now_ts
from service.drama_subtitle_language import classify_subtitle_text, ensure_language_columns


DEFAULT_SOURCE_XLSX = "短剧上架状态-20260722.xlsx"
DEFAULT_RAW_ROOT = "data_samples/middle_platform_subtitles_20260722_full/raw_books"
DEFAULT_DB_PATH = "data/drama_subtitle_similarity_v1.sqlite3"
DEFAULT_SCHEMA_PATH = "service/drama_subtitle_schema_v1.sql"
DEFAULT_OUTPUT_ROOT = "data_samples/middle_platform_subtitles_20260722_phase1"
DEFAULT_FIRST_EPISODE_LIMIT = 10
DEFAULT_COMMIT_EVERY = 25

WHITESPACE_RE = re.compile(r"\s+")


@dataclass(frozen=True)
class BookSeed:
    row_index: int
    book_id: str
    title: str
    tag: str


@dataclass(frozen=True)
class ExtractedBook:
    book_row: tuple[Any, ...]
    episode_rows: list[tuple[Any, ...]]
    line_rows: list[tuple[Any, ...]]
    summary_row: dict[str, Any]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build phase-1 normalized drama subtitle corpus from raw middle-platform payloads.",
    )
    parser.add_argument("--source-xlsx", default=DEFAULT_SOURCE_XLSX)
    parser.add_argument("--raw-root", default=DEFAULT_RAW_ROOT)
    parser.add_argument("--db", default=DEFAULT_DB_PATH)
    parser.add_argument("--schema", default=DEFAULT_SCHEMA_PATH)
    parser.add_argument("--output-root", default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--first-episode-limit", type=int, default=DEFAULT_FIRST_EPISODE_LIMIT)
    parser.add_argument("--commit-every", type=int, default=DEFAULT_COMMIT_EVERY)
    parser.add_argument("--limit", type=int, default=0, help="Only process first N payload files.")
    parser.add_argument("--book-id", default="", help="Only process one target book id.")
    parser.add_argument("--overwrite", action="store_true", help="Reset existing drama subtitle tables first.")
    return parser.parse_args()


def resolve_path(raw_path: str) -> Path:
    path = Path(raw_path)
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()
    return path


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def load_seed_rows(path: Path) -> dict[str, BookSeed]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook.active
        result: dict[str, BookSeed] = {}
        for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            book_id = str(row[0] or "").strip()
            if not book_id or book_id in result:
                continue
            result[book_id] = BookSeed(
                row_index=row_index,
                book_id=book_id,
                title=str(row[1] or "").strip(),
                tag=str(row[2] or "").strip(),
            )
        return result
    finally:
        workbook.close()


def load_schema_sql(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def iter_raw_files(raw_root: Path, *, book_id: str = "", limit: int = 0) -> list[Path]:
    if book_id:
        matches = sorted(raw_root.rglob(f"{book_id}.json"))
        return matches[:1]
    files = sorted(raw_root.rglob("*.json"))
    if limit > 0:
        return files[:limit]
    return files


def is_real_subtitle_item(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    text = normalize_text(item.get("text") or "")
    if text:
        return True
    return any(str(item.get(key) or "").strip() for key in ("start", "end", "speaker"))


def normalize_text(text: str) -> str:
    normalized = str(text or "").replace("\u3000", " ").replace("\xa0", " ").strip()
    normalized = WHITESPACE_RE.sub(" ", normalized)
    return normalized


def safe_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def build_episode_uid(book_id: str, episode_order: int) -> str:
    return f"{book_id}:{episode_order}"


def build_line_uid(book_id: str, episode_order: int, line_order: int) -> str:
    return f"{book_id}:{episode_order}:{line_order}"


def extract_payload(
    payload: dict[str, Any],
    *,
    payload_path: Path,
    seed_map: dict[str, BookSeed],
    first_episode_limit: int,
) -> ExtractedBook | None:
    book_id = str(payload.get("bookId") or "").strip()
    if not book_id:
        return None

    book_name = str(payload.get("bookName") or "").strip()
    episodes = payload.get("episodes")
    episode_list = episodes if isinstance(episodes, list) else []
    limited_episodes = episode_list[:first_episode_limit]
    seed = seed_map.get(book_id)
    seed_title = seed.title if seed else ""
    seed_tag = seed.tag if seed else ""

    total_subtitle_line_count = 0
    total_real_subtitle_line_count = 0
    first10_subtitle_line_count = 0
    first10_real_subtitle_line_count = 0
    episode_rows: list[tuple[Any, ...]] = []
    line_rows: list[tuple[Any, ...]] = []
    created_at = now_ts()

    for episode_index, episode in enumerate(episode_list, start=1):
        subtitle_items = episode.get("subtitles") if isinstance(episode, dict) else []
        subtitle_list = subtitle_items if isinstance(subtitle_items, list) else []
        total_subtitle_line_count += len(subtitle_list)
        total_real_subtitle_line_count += sum(1 for item in subtitle_list if is_real_subtitle_item(item))

        if episode_index > first_episode_limit:
            continue

        episode_uid = build_episode_uid(book_id, episode_index)
        chapter_id = safe_int(episode.get("chapterId")) if isinstance(episode, dict) else None
        chapter_name = str(episode.get("chapterName") or "") if isinstance(episode, dict) else ""
        summary = str(episode.get("summary") or "") if isinstance(episode, dict) else ""

        normalized_lines: list[tuple[Any, ...]] = []
        first_subtitle_start = ""
        last_subtitle_end = ""
        kept_line_order = 0
        real_subtitle_count = 0

        for raw_item in subtitle_list:
            if not isinstance(raw_item, dict):
                continue
            text_normalized = normalize_text(raw_item.get("text") or "")
            if not text_normalized:
                continue
            kept_line_order += 1
            real_subtitle_count += 1
            start_time = str(raw_item.get("start") or "").strip()
            end_time = str(raw_item.get("end") or "").strip()
            if kept_line_order == 1:
                first_subtitle_start = start_time
            if end_time:
                last_subtitle_end = end_time
            line_uid = build_line_uid(book_id, episode_index, kept_line_order)
            language = classify_subtitle_text(text_normalized)
            normalized_lines.append(
                (
                    line_uid,
                    book_id,
                    episode_uid,
                    episode_index,
                    chapter_id,
                    kept_line_order,
                    safe_int(raw_item.get("id")),
                    start_time,
                    end_time,
                    str(raw_item.get("speaker") or "").strip(),
                    str(raw_item.get("text") or "").strip(),
                    text_normalized,
                    len(text_normalized),
                    language.language_code,
                    language.confidence,
                    created_at,
                )
            )

        first10_subtitle_line_count += len(subtitle_list)
        first10_real_subtitle_line_count += real_subtitle_count
        episode_rows.append(
            (
                episode_uid,
                book_id,
                episode_index,
                chapter_id,
                chapter_name,
                summary,
                len(subtitle_list),
                real_subtitle_count,
                first_subtitle_start,
                last_subtitle_end,
                created_at,
                created_at,
            )
        )
        line_rows.extend(normalized_lines)

    if first10_real_subtitle_line_count <= 0:
        return None

    book_row = (
        book_id,
        book_name,
        seed_title,
        seed_tag,
        len(episode_list),
        min(len(episode_list), first_episode_limit),
        int(total_real_subtitle_line_count > 0),
        int(first10_real_subtitle_line_count > 0),
        total_subtitle_line_count,
        total_real_subtitle_line_count,
        first10_subtitle_line_count,
        first10_real_subtitle_line_count,
        str(payload_path),
        created_at,
        created_at,
    )
    summary_row = {
        "book_id": book_id,
        "book_name": book_name,
        "source_title": seed_title,
        "source_tag": seed_tag,
        "episode_count": len(episode_list),
        "first10_episode_count": min(len(episode_list), first_episode_limit),
        "total_subtitle_line_count": total_subtitle_line_count,
        "total_real_subtitle_line_count": total_real_subtitle_line_count,
        "first10_subtitle_line_count": first10_subtitle_line_count,
        "first10_real_subtitle_line_count": first10_real_subtitle_line_count,
        "source_payload_path": str(payload_path),
    }
    return ExtractedBook(
        book_row=book_row,
        episode_rows=episode_rows,
        line_rows=line_rows,
        summary_row=summary_row,
    )


def reset_tables(conn) -> None:  # type: ignore[no-untyped-def]
    conn.execute("DELETE FROM drama_subtitle_lines")
    conn.execute("DELETE FROM drama_episodes")
    conn.execute("DELETE FROM drama_books")
    conn.commit()


def flush_batch(
    conn,  # type: ignore[no-untyped-def]
    *,
    delete_book_ids: list[str],
    book_rows: list[tuple[Any, ...]],
    episode_rows: list[tuple[Any, ...]],
    line_rows: list[tuple[Any, ...]],
    allow_replace: bool,
) -> None:
    if not book_rows:
        return

    conn.execute("BEGIN")
    if allow_replace and delete_book_ids:
        conn.executemany(
            "DELETE FROM drama_subtitle_lines WHERE book_id = ?",
            [(book_id,) for book_id in delete_book_ids],
        )
        conn.executemany(
            "DELETE FROM drama_episodes WHERE book_id = ?",
            [(book_id,) for book_id in delete_book_ids],
        )
        conn.executemany(
            "DELETE FROM drama_books WHERE book_id = ?",
            [(book_id,) for book_id in delete_book_ids],
        )

    conn.executemany(
        """
        INSERT INTO drama_books(
            book_id,
            book_name,
            source_title,
            source_tag,
            episode_count,
            first10_episode_count,
            has_any_real_subtitles,
            first10_has_real_subtitles,
            total_subtitle_line_count,
            total_real_subtitle_line_count,
            first10_subtitle_line_count,
            first10_real_subtitle_line_count,
            source_payload_path,
            created_at,
            updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        book_rows,
    )
    conn.executemany(
        """
        INSERT INTO drama_episodes(
            episode_uid,
            book_id,
            episode_order,
            chapter_id,
            chapter_name,
            summary,
            subtitle_line_count,
            real_subtitle_line_count,
            first_subtitle_start,
            last_subtitle_end,
            created_at,
            updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        episode_rows,
    )
    conn.executemany(
        """
        INSERT INTO drama_subtitle_lines(
            line_uid,
            book_id,
            episode_uid,
            episode_order,
            chapter_id,
            line_order,
            source_subtitle_id,
            start_time,
            end_time,
            speaker,
            text,
            text_normalized,
            char_count,
            language_code,
            language_confidence,
            created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        line_rows,
    )
    conn.commit()


def write_summary_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    fieldnames = [
        "book_id",
        "book_name",
        "source_title",
        "source_tag",
        "episode_count",
        "first10_episode_count",
        "total_subtitle_line_count",
        "total_real_subtitle_line_count",
        "first10_subtitle_line_count",
        "first10_real_subtitle_line_count",
        "source_payload_path",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_summary_json(path: Path, summary: dict[str, Any]) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    args = parse_args()
    source_xlsx = resolve_path(args.source_xlsx)
    raw_root = resolve_path(args.raw_root)
    db_path = resolve_path(args.db)
    schema_path = resolve_path(args.schema)
    output_root = resolve_path(args.output_root)

    if not source_xlsx.exists():
        raise FileNotFoundError(f"source workbook not found: {source_xlsx}")
    if not raw_root.exists():
        raise FileNotFoundError(f"raw root not found: {raw_root}")
    if not schema_path.exists():
        raise FileNotFoundError(f"schema file not found: {schema_path}")
    if args.first_episode_limit <= 0:
        raise ValueError("--first-episode-limit must be > 0")
    if args.commit_every <= 0:
        raise ValueError("--commit-every must be > 0")

    ensure_dir(db_path.parent)
    ensure_dir(output_root)
    seed_map = load_seed_rows(source_xlsx)
    raw_files = iter_raw_files(raw_root, book_id=args.book_id.strip(), limit=max(int(args.limit or 0), 0))
    if not raw_files:
        raise FileNotFoundError("no raw subtitle payload files matched current filters")

    conn = connect_db(db_path)
    conn.row_factory = None
    try:
        conn.execute("PRAGMA journal_mode = WAL")
        conn.execute("PRAGMA synchronous = NORMAL")
        conn.execute("PRAGMA temp_store = MEMORY")
        conn.executescript(load_schema_sql(schema_path))
        ensure_language_columns(conn)
        if args.overwrite:
            reset_tables(conn)

        started_at = datetime.now().isoformat(sep=" ", timespec="seconds")
        processed_payloads = 0
        inserted_books = 0
        inserted_episodes = 0
        inserted_lines = 0
        skipped_invalid_payload = 0
        skipped_missing_book_id = 0
        skipped_no_first10_real_subtitles = 0
        summary_rows: list[dict[str, Any]] = []

        pending_delete_book_ids: list[str] = []
        pending_book_rows: list[tuple[Any, ...]] = []
        pending_episode_rows: list[tuple[Any, ...]] = []
        pending_line_rows: list[tuple[Any, ...]] = []

        for raw_index, raw_file in enumerate(raw_files, start=1):
            processed_payloads += 1
            try:
                payload = json.loads(raw_file.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                skipped_invalid_payload += 1
                continue

            if not isinstance(payload, dict):
                skipped_invalid_payload += 1
                continue

            raw_book_id = str(payload.get("bookId") or "").strip()
            if not raw_book_id:
                skipped_missing_book_id += 1
                continue

            extracted = extract_payload(
                payload,
                payload_path=raw_file,
                seed_map=seed_map,
                first_episode_limit=args.first_episode_limit,
            )
            if extracted is None:
                skipped_no_first10_real_subtitles += 1
                continue

            pending_delete_book_ids.append(extracted.summary_row["book_id"])
            pending_book_rows.append(extracted.book_row)
            pending_episode_rows.extend(extracted.episode_rows)
            pending_line_rows.extend(extracted.line_rows)
            summary_rows.append(extracted.summary_row)
            inserted_books += 1
            inserted_episodes += len(extracted.episode_rows)
            inserted_lines += len(extracted.line_rows)

            if len(pending_book_rows) >= args.commit_every:
                flush_batch(
                    conn,
                    delete_book_ids=pending_delete_book_ids,
                    book_rows=pending_book_rows,
                    episode_rows=pending_episode_rows,
                    line_rows=pending_line_rows,
                    allow_replace=not args.overwrite,
                )
                print(
                    f"PROGRESS processed_payloads={processed_payloads} "
                    f"inserted_books={inserted_books} "
                    f"inserted_lines={inserted_lines}"
                )
                pending_delete_book_ids = []
                pending_book_rows = []
                pending_episode_rows = []
                pending_line_rows = []

            if raw_index % 500 == 0:
                print(
                    f"SCAN processed_payloads={processed_payloads} "
                    f"inserted_books={inserted_books} "
                    f"skipped_no_first10_real_subtitles={skipped_no_first10_real_subtitles}"
                )

        flush_batch(
            conn,
            delete_book_ids=pending_delete_book_ids,
            book_rows=pending_book_rows,
            episode_rows=pending_episode_rows,
            line_rows=pending_line_rows,
            allow_replace=not args.overwrite,
        )

        finished_at = datetime.now().isoformat(sep=" ", timespec="seconds")
        run_stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        logs_dir = output_root / "logs"
        summary_csv_path = logs_dir / f"phase1_book_summary_{run_stamp}.csv"
        summary_json_path = logs_dir / f"phase1_build_summary_{run_stamp}.json"
        write_summary_csv(summary_csv_path, summary_rows)
        summary_payload = {
            "started_at": started_at,
            "finished_at": finished_at,
            "raw_root": str(raw_root),
            "source_xlsx": str(source_xlsx),
            "db_path": str(db_path),
            "summary_csv_path": str(summary_csv_path),
            "processed_payloads": processed_payloads,
            "inserted_books": inserted_books,
            "inserted_episodes": inserted_episodes,
            "inserted_lines": inserted_lines,
            "skipped_invalid_payload": skipped_invalid_payload,
            "skipped_missing_book_id": skipped_missing_book_id,
            "skipped_no_first10_real_subtitles": skipped_no_first10_real_subtitles,
            "first_episode_limit": args.first_episode_limit,
            "commit_every": args.commit_every,
            "overwrite": int(args.overwrite),
            "book_filter": args.book_id.strip(),
            "limit": max(int(args.limit or 0), 0),
        }
        write_summary_json(summary_json_path, summary_payload)
    finally:
        conn.close()

    print(f"OK processed_payloads={processed_payloads}")
    print(f"OK inserted_books={inserted_books}")
    print(f"OK inserted_episodes={inserted_episodes}")
    print(f"OK inserted_lines={inserted_lines}")
    print(f"OK skipped_invalid_payload={skipped_invalid_payload}")
    print(f"OK skipped_missing_book_id={skipped_missing_book_id}")
    print(f"OK skipped_no_first10_real_subtitles={skipped_no_first10_real_subtitles}")
    print(f"OK first_episode_limit={args.first_episode_limit}")
    print(f"OK db_path={db_path}")


if __name__ == "__main__":
    main()
