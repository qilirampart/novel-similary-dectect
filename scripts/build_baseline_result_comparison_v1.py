from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


BASELINE_NAME = "分销平台6个频道的推送视频信息0827_含原版剧.xlsx"
RESULT_HEADERS = [
    "来源频道",
    "来源剧名",
    "原视频链接",
    "视频 ID",
    "完整字幕",
    "匹配状态",
    "用户结论",
    "命中剧名",
    "Book ID",
    "命中集数",
    "命中时间范围",
    "匹配原因",
    "确认命中数",
    "待复核数",
    "未命中数",
    "翻译回退",
    "译文",
    "服务端执行",
    "强证据候选",
    "命中证据",
]
BASELINE_HEADERS = [
    "channel_uid",
    "channel_name",
    "material_id",
    "youtube_video_id",
    "drama_id",
    "drama_name",
    "post_type",
    "publish_at",
    "原版剧ID",
    "原版剧名称",
]


def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(iterator)]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if any(value is not None and str(value) != "" for value in values):
            rows.append(dict(zip(headers, values)))
    return headers, rows


def compact_json(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value
    try:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    except (TypeError, ValueError):
        return str(value)


def json_list(value: Any) -> list[Any]:
    try:
        parsed = json.loads(str(value or ""))
    except (TypeError, ValueError, json.JSONDecodeError):
        return []
    return parsed if isinstance(parsed, list) else []


def normalize_id(value: Any) -> str:
    return str(value or "").strip()


def extract_hit_ids(row: dict[str, Any]) -> list[str]:
    identifiers = re.findall(r"\b\d{8,}\b", normalize_id(row.get("Book ID")))
    for field in ("强证据候选", "命中证据"):
        for item in json_list(row.get(field)):
            if isinstance(item, dict):
                candidate_id = normalize_id(item.get("book_id") or item.get("Book ID"))
                if re.fullmatch(r"\d{8,}", candidate_id):
                    identifiers.append(candidate_id)
    return list(dict.fromkeys(identifiers))


def deduplicate_result_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep one result per video and let later input files override older runs."""
    by_video: dict[str, dict[str, Any]] = {}
    no_video_rows: list[dict[str, Any]] = []
    for row in rows:
        video_id = normalize_id(row.get("视频 ID"))
        if not video_id:
            no_video_rows.append(row)
            continue
        by_video.pop(video_id, None)
        by_video[video_id] = row
    return [*by_video.values(), *no_video_rows]


def build_book_inventory(db_path: Path) -> dict[str, dict[str, Any]]:
    import sqlite3

    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    try:
        rows = connection.execute(
            """
            SELECT book_id, book_name, has_any_real_subtitles,
                   first10_has_real_subtitles, total_subtitle_line_count,
                   first10_subtitle_line_count
              FROM drama_books
            """
        ).fetchall()
    finally:
        connection.close()
    return {normalize_id(row["book_id"]): dict(row) for row in rows}


def library_state(book_id: Any, inventory: dict[str, dict[str, Any]]) -> str:
    identifier = normalize_id(book_id)
    if not identifier:
        return "未提供 ID"
    record = inventory.get(identifier)
    if record is None:
        return "未录入字幕库"
    if record.get("first10_has_real_subtitles"):
        return "已录入，前10集有字幕"
    if record.get("has_any_real_subtitles"):
        return "已录入，有字幕但前10集无字幕"
    return "已录入，但无可用字幕"


def state_has_usable_subtitle(state: str) -> bool:
    return state == "已录入，前10集有字幕"


def result_status_is_confirmed(status: str) -> bool:
    return status == "confirmed_match"


def is_result_status(status: str) -> bool:
    return bool(status)


def classify_join(
    *,
    result_status: str,
    hit_ids: list[str],
    baseline: dict[str, Any] | None,
    inventory: dict[str, dict[str, Any]],
) -> tuple[str, str, str, str, str]:
    if baseline is None:
        return (
            "结果表独有，无法与基准表归属",
            "未关联",
            "未关联",
            "未关联",
            "结果记录不在基准表视频 ID 集合中",
        )

    push_id = normalize_id(baseline.get("drama_id"))
    original_id = normalize_id(baseline.get("原版剧ID"))
    push_state = library_state(push_id, inventory)
    original_state = library_state(original_id, inventory)
    hit_set = set(hit_ids)

    if result_status == "no_match":
        if push_state == "未录入字幕库" and original_state in {"未录入字幕库", "未提供 ID"}:
            category = "未命中：推送 ID 与原版 ID 均不在字幕库"
        elif state_has_usable_subtitle(push_state) and state_has_usable_subtitle(original_state):
            category = "未命中：推送 ID 与原版 ID 均有前10集字幕"
        elif state_has_usable_subtitle(original_state):
            category = "未命中：原版 ID 在库且前10集有字幕"
        elif state_has_usable_subtitle(push_state):
            category = "未命中：推送 ID 在库且前10集有字幕"
        elif original_state not in {"未录入字幕库", "未提供 ID"} or push_state not in {"未录入字幕库", "未提供 ID"}:
            category = "未命中：基准 ID 已录入但当前未形成命中"
        else:
            category = "未命中：基准 ID 信息不足"
        return category, push_state, original_state, "未命中", "未达到系统命中阈值"

    if original_id and original_id in hit_set:
        base_category = "原版剧 ID"
    elif push_id and push_id in hit_set:
        base_category = "推送剧 ID"
    elif hit_set:
        base_category = "库内其他版本"
    else:
        base_category = "未解析出命中 ID"

    if result_status_is_confirmed(result_status):
        category = f"确认命中：{base_category}"
        certainty = "确认命中"
    else:
        category = f"候选命中：{base_category}（待复核）"
        certainty = "候选/待复核"
    return category, push_state, original_state, certainty, "命中 ID 与基准字段完成严格比对"


def make_result_rows(
    result_paths: list[Path],
    baseline_by_video: dict[str, list[dict[str, Any]]],
    inventory: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in result_paths:
        headers, source_rows = read_rows(path)
        if "匹配状态" not in headers:
            continue
        for source_row_number, source_row in enumerate(source_rows, start=2):
            video_id = normalize_id(source_row.get("视频 ID"))
            baseline_matches = baseline_by_video.get(video_id, []) if video_id else []
            baseline = baseline_matches[0] if baseline_matches else None
            hit_ids = extract_hit_ids(source_row)
            category, push_state, original_state, certainty, category_reason = classify_join(
                result_status=normalize_id(source_row.get("匹配状态")),
                hit_ids=hit_ids,
                baseline=baseline,
                inventory=inventory,
            )
            row: dict[str, Any] = {
                "记录来源": "结果与基准交集" if baseline else "结果表独有",
                "结果来源文件": path.name,
                "结果源行号": source_row_number,
                "视频 ID": video_id,
                "基准关联行数": len(baseline_matches),
                "命中归属判断": category,
                "判定确定性": certainty,
                "推送 ID 库状态": push_state,
                "原版 ID 库状态": original_state,
                "归属判断说明": category_reason,
            }
            for key in RESULT_HEADERS:
                row[key] = source_row.get(key)
            for key in BASELINE_HEADERS:
                output_key = {
                    "channel_uid": "基准 channel_uid",
                    "channel_name": "基准频道",
                    "material_id": "基准 material_id",
                    "youtube_video_id": "基准视频 ID",
                    "drama_id": "基准推送 drama_id",
                    "drama_name": "基准推送剧名",
                    "post_type": "基准 post_type",
                    "publish_at": "基准发布时间",
                    "原版剧ID": "基准原版剧ID",
                    "原版剧名称": "基准原版剧名称",
                }[key]
                row[output_key] = baseline.get(key) if baseline else None
            row["命中 Book ID（解析后）"] = " | ".join(hit_ids)
            rows.append(row)
    return rows


def make_baseline_only_rows(
    baseline_rows: list[dict[str, Any]],
    result_video_ids: set[str],
    inventory: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in baseline_rows:
        video_id = normalize_id(row.get("youtube_video_id"))
        if video_id in result_video_ids:
            continue
        push_state = library_state(row.get("drama_id"), inventory)
        original_state = library_state(row.get("原版剧ID"), inventory)
        output.append(
            {
                **row,
                "记录状态": "基准表独有，结果表暂无记录",
                "推送 ID 库状态": push_state,
                "原版 ID 库状态": original_state,
                "是否为有效视频 ID": "否" if video_id.startswith("#") or not video_id else "是",
            }
        )
    return output


def write_table_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[dict[str, Any]],
    *,
    freeze_columns: int = 0,
    text_columns: set[str] | None = None,
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = f"{chr(65 + freeze_columns)}2" if freeze_columns else "A2"
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="1F2937")
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="9FBAD0"))
    worksheet.row_dimensions[1].height = 30

    text_columns = text_columns or set()
    for row_cells in worksheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="left" if headers[cell.column - 1] in text_columns else "center",
                vertical="top",
                wrap_text=headers[cell.column - 1] in text_columns,
            )
            cell.border = Border(bottom=thin_gray)
    widths = {
        "来源剧名": 36,
        "完整字幕": 60,
        "命中剧名": 32,
        "匹配原因": 42,
        "用户结论": 42,
        "翻译回退": 46,
        "译文": 60,
        "服务端执行": 32,
        "强证据候选": 48,
        "命中证据": 70,
        "基准推送剧名": 32,
        "基准原版剧名称": 32,
        "归属判断说明": 42,
    }
    for index, header in enumerate(headers, start=1):
        width = widths.get(header, 18)
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    worksheet.auto_filter.ref = worksheet.dimensions
    if rows:
        table_ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
        # Excel requires table display names to be unique across the workbook.
        # Chinese sheet names would otherwise collapse to the same fallback name.
        table = Table(displayName=f"DataTable{len(workbook.worksheets)}", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        worksheet.add_table(table)

    status_column = None
    for index, header in enumerate(headers, start=1):
        if header == "命中归属判断":
            status_column = openpyxl.utils.get_column_letter(index)
            break
    if status_column and rows:
        status_range = f"{status_column}2:{status_column}{len(rows) + 1}"
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'ISNUMBER(SEARCH("未命中",{status_column}2))'], fill=PatternFill("solid", fgColor="FCE4D6")),
        )
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'ISNUMBER(SEARCH("确认命中",{status_column}2))'], fill=PatternFill("solid", fgColor="E2F0D9")),
        )


def write_summary_sheet(
    workbook: Workbook,
    *,
    baseline_rows: list[dict[str, Any]],
    result_rows: list[dict[str, Any]],
    baseline_only_rows: list[dict[str, Any]],
    inventory: dict[str, dict[str, Any]],
    result_file_names: list[str],
) -> None:
    worksheet = workbook.create_sheet("汇总")
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 64
    worksheet["A1"] = "字幕匹配结果与基准表交叉核验汇总"
    worksheet.merge_cells("A1:C1")
    worksheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 32
    worksheet["A3"] = "统计范围"
    worksheet["B3"] = "数量"
    worksheet["C3"] = "说明"
    for cell in worksheet[3]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    baseline_video_ids = {normalize_id(row.get("youtube_video_id")) for row in baseline_rows if normalize_id(row.get("youtube_video_id"))}
    result_video_ids = {normalize_id(row.get("视频 ID")) for row in result_rows if normalize_id(row.get("视频 ID"))}
    counts = [
        ("基准表记录", len(baseline_rows), "分销平台六频道基准表原始记录数"),
        ("基准表唯一有效视频 ID", len({x for x in baseline_video_ids if not x.startswith("#")}), "排除空值和 Excel 错误值后的唯一视频数"),
        ("字幕匹配结果记录", len(result_rows), f"纳入 {len(result_file_names)} 份匹配结果表"),
        ("字幕匹配结果唯一视频 ID", len(result_video_ids), "结果表按视频 ID 去重后无重复"),
        ("结果与基准表交集", len(result_rows) - len([x for x in result_rows if x.get("记录来源") == "结果表独有"]), "按视频 ID 严格关联"),
        ("结果表独有记录", sum(x.get("记录来源") == "结果表独有" for x in result_rows), "结果表中有结果但基准表没有对应视频 ID"),
        ("基准表独有记录", len(baseline_only_rows), "基准表中暂时没有对应匹配结果"),
    ]
    for row_index, values in enumerate(counts, start=4):
        for column_index, value in enumerate(values, start=1):
            worksheet.cell(row_index, column_index, value)
    worksheet["A13"] = "交集结果归属"
    worksheet.merge_cells("A13:C13")
    worksheet["A13"].fill = PatternFill("solid", fgColor="D9EAF7")
    worksheet["A13"].font = Font(name="Microsoft YaHei", size=11, bold=True, color="17365D")
    worksheet["A14"] = "命中归属判断"
    worksheet["B14"] = "数量"
    worksheet["C14"] = "口径"
    for cell in worksheet[14]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    category_counts = Counter(str(row.get("命中归属判断") or "") for row in result_rows if row.get("记录来源") == "结果与基准交集")
    for row_index, (category, count) in enumerate(sorted(category_counts.items()), start=15):
        worksheet.cell(row_index, 1, category)
        worksheet.cell(row_index, 2, count)
        worksheet.cell(row_index, 3, "结果 Book ID 与基准表推送/原版 ID 的严格比对")
    start = 15 + len(category_counts) + 2
    worksheet.cell(start, 1, "原始匹配状态")
    worksheet.merge_cells(start_row=start, start_column=1, end_row=start, end_column=3)
    worksheet.cell(start, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    worksheet.cell(start, 1).font = Font(name="Microsoft YaHei", size=11, bold=True, color="17365D")
    for col, value in enumerate(("匹配状态", "数量", "说明"), start=1):
        cell = worksheet.cell(start + 1, col, value)
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    status_counts = Counter(str(row.get("匹配状态") or "") for row in result_rows)
    for row_index, (status, count) in enumerate(sorted(status_counts.items()), start=start + 2):
        worksheet.cell(row_index, 1, status)
        worksheet.cell(row_index, 2, count)
        worksheet.cell(row_index, 3, "原始结果表字段，不等同于最终业务归属")
    note_row = start + 2 + len(status_counts) + 2
    notes = [
        "关键口径：结果表和基准表按视频 ID 严格关联；不同视频 ID 不因剧名相似而强行合并。",
        "命中归属判断仅比较结果中的 Book ID 与基准表的推送 drama_id、原版剧ID，不代表两个 ID 已完成版本映射。",
        "候选命中、翻译辅助和待复核结果不会被统计为确认命中；请结合“结果总表”和“未命中核查”页复核。",
        "字幕库状态来自 drama_books：已录入不等于一定有字幕，前10集可用字幕状态单独列出。",
        "基准表存在 #NAME? 视频 ID，属于源表错误值，已保留但不作为有效视频 ID 参与唯一性统计。",
        "纳入结果文件：" + "、".join(result_file_names),
        "聚合规则：按视频 ID 去重；同一视频重复时以后加入的结果覆盖先前结果。",
    ]
    worksheet.cell(note_row, 1, "口径与限制")
    worksheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
    worksheet.cell(note_row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    worksheet.cell(note_row, 1).font = Font(name="Microsoft YaHei", size=11, bold=True, color="17365D")
    for index, note in enumerate(notes, start=note_row + 1):
        worksheet.cell(index, 1, note)
        worksheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=3)
        worksheet.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.row_dimensions[index].height = 30
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row not in {1, 3, 13, 14, start, start + 1, note_row}:
                cell.font = Font(name="Microsoft YaHei", size=10, color="1F2937")
                cell.border = Border(bottom=Side(style="thin", color="D9E2F3"))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A4"


def build_workbook(args: argparse.Namespace) -> Path:
    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    baseline_path = Path(args.folder) / BASELINE_NAME
    baseline_headers, baseline_rows = read_rows(baseline_path)
    if baseline_headers[: len(BASELINE_HEADERS)] != BASELINE_HEADERS:
        raise ValueError(f"baseline headers do not match expected schema: {baseline_headers}")
    baseline_by_video: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in baseline_rows:
        baseline_by_video[normalize_id(row.get("youtube_video_id"))].append(row)

    result_paths = [
        Path(args.folder) / "@Bound2Drama.xlsx",
        Path(args.folder) / "AI结果-日语、葡语.xlsx",
        Path(args.folder) / "FlipThatDrama.xlsx",
        Path(args.folder) / "云上剧场.xlsx",
        Path(args.korean_result),
    ]
    if args.extra_result:
        result_paths.append(Path(args.extra_result))
    missing = [str(path) for path in result_paths if not path.exists()]
    if missing:
        raise FileNotFoundError("missing result files: " + ", ".join(missing))
    inventory = build_book_inventory(Path(args.db))
    raw_result_rows = make_result_rows(result_paths, baseline_by_video, inventory)
    result_rows = deduplicate_result_rows(raw_result_rows)
    result_video_ids = {normalize_id(row.get("视频 ID")) for row in result_rows if normalize_id(row.get("视频 ID"))}
    baseline_only_rows = make_baseline_only_rows(baseline_rows, result_video_ids, inventory)
    no_match_rows = [row for row in result_rows if row.get("记录来源") == "结果与基准交集" and row.get("匹配状态") == "no_match"]

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    write_summary_sheet(
        workbook,
        baseline_rows=baseline_rows,
        result_rows=result_rows,
        baseline_only_rows=baseline_only_rows,
        inventory=inventory,
        result_file_names=[path.name for path in result_paths],
    )
    result_headers = [
        "记录来源", "结果来源文件", "结果源行号", "视频 ID", "基准关联行数",
        "命中归属判断", "判定确定性", "推送 ID 库状态", "原版 ID 库状态", "归属判断说明",
        "命中 Book ID（解析后）",
        "来源频道", "来源剧名", "原视频链接", "完整字幕", "匹配状态", "用户结论",
        "命中剧名", "Book ID", "命中集数", "命中时间范围", "匹配原因", "确认命中数",
        "待复核数", "未命中数", "翻译回退", "译文", "服务端执行", "强证据候选", "命中证据",
        "基准 channel_uid", "基准频道", "基准 material_id", "基准视频 ID", "基准推送 drama_id",
        "基准推送剧名", "基准 post_type", "基准发布时间", "基准原版剧ID", "基准原版剧名称",
    ]
    write_table_sheet(
        workbook,
        "结果总表",
        result_headers,
        result_rows,
        freeze_columns=4,
        text_columns={
            "来源剧名", "完整字幕", "用户结论", "命中剧名", "匹配原因", "翻译回退", "译文",
            "强证据候选", "命中证据", "基准推送剧名", "基准原版剧名称", "归属判断说明",
        },
    )
    baseline_only_headers = BASELINE_HEADERS + ["记录状态", "推送 ID 库状态", "原版 ID 库状态", "是否为有效视频 ID"]
    write_table_sheet(
        workbook,
        "基准未出结果",
        baseline_only_headers,
        baseline_only_rows,
        freeze_columns=4,
        text_columns={"channel_name", "drama_name", "原版剧名称", "记录状态", "推送 ID 库状态", "原版 ID 库状态"},
    )
    review_headers = [
        "结果来源文件", "结果源行号", "视频 ID", "来源频道", "来源剧名", "匹配状态", "命中剧名",
        "Book ID", "匹配原因", "命中归属判断", "推送 ID 库状态", "原版 ID 库状态", "基准推送 drama_id",
        "基准推送剧名", "基准原版剧ID", "基准原版剧名称", "完整字幕", "翻译回退", "译文", "强证据候选", "命中证据",
    ]
    review_rows = []
    for row in no_match_rows:
        review_rows.append({header: row.get(header) for header in review_headers})
    write_table_sheet(
        workbook,
        "未命中核查",
        review_headers,
        review_rows,
        freeze_columns=3,
        text_columns={"来源剧名", "命中剧名", "匹配原因", "基准推送剧名", "基准原版剧名称", "完整字幕", "翻译回退", "译文", "强证据候选", "命中证据"},
    )
    field_rows = [
        {"字段": "命中归属判断", "定义": "按结果 Book ID 与基准表推送 drama_id、原版剧ID严格比较后的业务归属。", "注意": "候选/翻译辅助结果会标注待复核，不自动等同确认命中。"},
        {"字段": "结果表独有", "定义": "结果表视频 ID 在基准表中不存在。", "注意": "这类记录是新增采集范围，不能作为未命中或数据缺失处理。"},
        {"字段": "基准表独有", "定义": "基准表视频 ID 在本轮结果表中不存在。", "注意": "可能尚未采集、尚未匹配，或源表视频 ID 为错误值。"},
        {"字段": "前10集有字幕", "定义": "字幕库 drama_books 标记前10集存在可用字幕。", "注意": "ID存在但无字幕时，不能把它当成可检索数据。"},
        {"字段": "未命中核查", "定义": "只保留结果与基准交集中的 no_match 记录。", "注意": "用于核对推送 ID、原版 ID是否在字幕库，以及是否属于库内已有但未命中。"},
    ]
    write_table_sheet(workbook, "字段说明", ["字段", "定义", "注意"], field_rows, text_columns={"定义", "注意"})
    workbook.save(output_path)

    # Export the highest-priority no-match subset for recall debugging.
    original_library_rows = [
        row for row in no_match_rows
        if row.get("原版 ID 库状态") == "已录入，前10集有字幕"
    ]
    focused_headers = [
        "结果来源文件", "结果源行号", "视频 ID", "来源频道", "来源剧名", "完整字幕",
        "匹配状态", "匹配原因", "译文", "命中证据", "强证据候选",
        "基准推送 drama_id", "基准推送剧名", "基准原版剧ID", "基准原版剧名称",
        "推送 ID 库状态", "原版 ID 库状态", "原视频链接", "服务端执行",
    ]
    focused_workbook = Workbook()
    focused_workbook.remove(focused_workbook.active)
    write_table_sheet(
        focused_workbook,
        "原版ID在库未命中",
        focused_headers,
        original_library_rows,
        freeze_columns=3,
        text_columns={
            "来源频道", "来源剧名", "完整字幕", "匹配原因", "译文", "命中证据",
            "强证据候选", "基准推送剧名", "基准原版剧名称", "推送 ID 库状态",
            "原版 ID 库状态",
        },
    )
    focused_output = output_path.parent / "未命中_原版ID在库_前10集有字幕.xlsx"
    focused_workbook.save(focused_output)
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build an auditable comparison workbook for subtitle results and baseline distribution data.")
    parser.add_argument("--folder", default=r"E:\点众\YouTube字幕核验助手工作区\output\结果比对文件夹")
    parser.add_argument("--korean-result", default=r"D:\Dianzhong\YouTube字幕核验助手\output\韩国频道163条.xlsx")
    parser.add_argument("--db", default=r"data\drama_subtitle_similarity_v1.sqlite3")
    parser.add_argument("--output", default=r"E:\点众\YouTube字幕核验助手工作区\output\结果比对文件夹\字幕匹配结果_基准表交叉核验总表.xlsx")
    parser.add_argument("--extra-result", default=r"E:\点众\YouTube字幕核验助手工作区\output\youtube_matching_results_after_160_20260827.xlsx")
    return parser.parse_args()


if __name__ == "__main__":
    print(build_workbook(parse_args()))
