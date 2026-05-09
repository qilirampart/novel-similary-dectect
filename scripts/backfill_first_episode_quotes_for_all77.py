from __future__ import annotations

import argparse
import re
import shutil
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_QUOTE_WORKBOOK = Path(
    r"C:\Users\psk13\Desktop\实验室任务\aps\02_功能包装材料论文项目\03_检索与写作\01_主稿\补充图片\漫剧台词表.xlsx"
)
DEFAULT_TARGET_WORKBOOK = Path(
    r"E:\点众\小说库相似度比对服务工作区\数据现状说明\01_中台可查_且数据库已有\01_修正版77条结果.xlsx"
)
DEFAULT_OLD_RESULT_WORKBOOK = Path(
    r"E:\点众\小说库相似度比对服务工作区\dialogue_quote_compare_results.xlsx"
)
DEFAULT_OUTPUT_WORKBOOK = Path(
    r"E:\点众\小说库相似度比对服务工作区\数据现状说明\01_中台可查_且数据库已有\01_修正版77条结果_补台词_2026-05-09.xlsx"
)

HEADER_FILL = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
HEADER_FONT = Font(name="Arial", bold=True, color="1F3A5F")
BODY_FONT = Font(name="Arial")
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


@dataclass
class QuoteRow:
    source_sheet: str
    source_account: str
    source_account_norm: str
    sheet_account: str
    sheet_account_norm: str
    drama_name: str
    drama_name_norm: str
    quote: str
    quote_valid: bool
    row_index: int


@dataclass
class MatchResult:
    status: str
    method: str
    quote: str
    source_sheet: str
    source_account: str
    source_drama_name: str
    score: float
    note: str
    review_old10: str
    old10_note: str


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return text.strip()


def normalize_compact(value: Any) -> str:
    text = normalize_text(value).lower()
    chars: list[str] = []
    for ch in text:
        if ch.isspace():
            continue
        cat = unicodedata.category(ch)
        if cat.startswith("P") or cat.startswith("S"):
            continue
        chars.append(ch)
    return "".join(chars)


def normalize_title(value: Any) -> str:
    text = normalize_compact(value)
    return text.replace("第1集", "").replace("第一集", "")


def normalize_account_token(value: Any) -> str:
    text = normalize_text(value)
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"主页[:：].*", "", text)
    text = re.sub(r"账号名称?[:：]\s*", "", text)
    text = re.sub(r"^[-:：\s]+", "", text)
    text = re.sub(r"[-:：\s]+$", "", text)
    return normalize_compact(text)


def is_valid_quote(value: Any) -> bool:
    text = normalize_text(value)
    if not text:
        return False
    if text.upper() in {"#N/A", "N/A", "NA", "NULL", "NONE", "<NA>", "NAN"}:
        return False
    return True


def extract_sheet_account(sheet_name: str) -> str:
    text = normalize_text(sheet_name)
    if "-" in text:
        return text.split("-", 1)[1].strip()
    return text


def extract_target_account_aliases(value: Any) -> set[str]:
    text = normalize_text(value)
    aliases: set[str] = set()
    if not text:
        return aliases
    for line in re.split(r"[\r\n]+", text):
        raw = normalize_text(line)
        if not raw:
            continue
        if raw.startswith("主页"):
            continue
        account_match = re.search(r"账号名称?[:：]\s*([^\n]+)", raw)
        if account_match:
            raw = account_match.group(1).strip()
        if "http" in raw.lower():
            continue
        aliases.add(normalize_account_token(raw))
        if "-" in raw:
            left = raw.split("-", 1)[0].strip()
            if left:
                aliases.add(normalize_account_token(left))
    return {item for item in aliases if item}


def account_matches(target_aliases: set[str], quote_row: QuoteRow) -> bool:
    if not target_aliases:
        return False
    candidates = [quote_row.source_account_norm, quote_row.sheet_account_norm]
    for target in target_aliases:
        for candidate in candidates:
            if not target or not candidate:
                continue
            if target in candidate or candidate in target:
                return True
    return False


def safe_cell_value(ws, row: int, col_map: dict[str, int], key: str) -> str:
    col_idx = col_map.get(key)
    if not col_idx:
        return ""
    return normalize_text(ws.cell(row, col_idx).value)


def build_header_map(ws, header_row: int = 1) -> dict[str, int]:
    return {
        normalize_text(ws.cell(header_row, col_idx).value): col_idx
        for col_idx in range(1, ws.max_column + 1)
        if normalize_text(ws.cell(header_row, col_idx).value)
    }


def load_quote_rows(path: Path) -> tuple[list[QuoteRow], dict[str, list[QuoteRow]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[QuoteRow] = []
    by_title: dict[str, list[QuoteRow]] = {}
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers = build_header_map(ws, 1)
            title_col = headers.get("剧名")
            quote_col = headers.get("首集台词")
            account_col = headers.get("账号名") or headers.get("账号名称")
            if not title_col or not quote_col:
                continue
            sheet_account = extract_sheet_account(sheet_name)
            sheet_account_norm = normalize_account_token(sheet_account)
            for row_idx in range(2, ws.max_row + 1):
                drama_name = normalize_text(ws.cell(row_idx, title_col).value)
                quote = normalize_text(ws.cell(row_idx, quote_col).value)
                quote_valid = is_valid_quote(quote)
                if not quote_valid:
                    quote = ""
                source_account = normalize_text(ws.cell(row_idx, account_col).value) if account_col else ""
                if not drama_name:
                    continue
                item = QuoteRow(
                    source_sheet=sheet_name,
                    source_account=source_account or sheet_account,
                    source_account_norm=normalize_account_token(source_account or sheet_account),
                    sheet_account=sheet_account,
                    sheet_account_norm=sheet_account_norm,
                    drama_name=drama_name,
                    drama_name_norm=normalize_title(drama_name),
                    quote=quote,
                    quote_valid=quote_valid,
                    row_index=row_idx,
                )
                rows.append(item)
                by_title.setdefault(item.drama_name_norm, []).append(item)
    finally:
        wb.close()
    return rows, by_title


def load_old10_review_set(path: Path) -> tuple[set[str], list[dict[str, str]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    matched_titles: set[str] = set()
    rows: list[dict[str, str]] = []
    try:
        if "Summary" not in wb.sheetnames:
            return matched_titles, rows
        ws = wb["Summary"]
        headers = build_header_map(ws, 1)
        for row_idx in range(2, ws.max_row + 1):
            short_drama = safe_cell_value(ws, row_idx, headers, "short_drama")
            if not short_drama:
                continue
            matched_titles.add(normalize_title(short_drama))
            rows.append(
                {
                    "priority": safe_cell_value(ws, row_idx, headers, "priority"),
                    "short_drama": short_drama,
                    "quote_preview": safe_cell_value(ws, row_idx, headers, "quote_preview"),
                }
            )
    finally:
        wb.close()
    return matched_titles, rows


def quote_preview_matches(full_quote: str, preview: str) -> bool:
    full_norm = normalize_compact(full_quote)
    preview_norm = normalize_compact(preview)
    if not preview_norm:
        return False
    return preview_norm in full_norm


def rank_exact_matches(matches: list[QuoteRow], target_aliases: set[str]) -> list[QuoteRow]:
    return sorted(
        matches,
        key=lambda item: (
            0 if item.quote_valid else 1,
            0 if account_matches(target_aliases, item) else 1,
            len(item.drama_name),
            item.source_sheet,
            item.row_index,
        ),
    )


def build_candidate_rows(
    target_title_norm: str,
    target_aliases: set[str],
    quote_rows: list[QuoteRow],
    limit: int = 5,
) -> list[tuple[QuoteRow, float, bool, bool]]:
    ranked: list[tuple[QuoteRow, float, bool, bool]] = []
    for row in quote_rows:
        if not row.drama_name_norm:
            continue
        score = SequenceMatcher(None, target_title_norm, row.drama_name_norm).ratio()
        contain = target_title_norm in row.drama_name_norm or row.drama_name_norm in target_title_norm
        same_account = account_matches(target_aliases, row)
        if score < 0.45 and not contain and not same_account:
            continue
        ranked.append((row, score, contain, same_account))
    ranked.sort(
        key=lambda item: (
            0 if item[0].quote else 1,
            0 if item[3] else 1,
            0 if item[2] else 1,
            -item[1],
            len(item[0].drama_name),
        )
    )
    return ranked[:limit]


def pick_match(
    short_drama: str,
    target_aliases: set[str],
    old10_norm_titles: set[str],
    exact_index: dict[str, list[QuoteRow]],
    all_quote_rows: list[QuoteRow],
) -> tuple[MatchResult, list[tuple[QuoteRow, float, bool, bool]]]:
    title_norm = normalize_title(short_drama)
    exact_matches = rank_exact_matches(exact_index.get(title_norm, []), target_aliases)
    review_old10 = "否"
    old10_note = ""

    if title_norm in old10_norm_titles:
        review_old10 = "是"

        if exact_matches:
            best = exact_matches[0]
            unique_sources = {(item.source_sheet, item.row_index) for item in exact_matches}
            note_parts = []
            if len(unique_sources) > 1:
                note_parts.append(f"同名来源{len(unique_sources)}条，已优先选择非空台词/账号更接近项")
            if not best.quote_valid:
                note_parts.append("剧名已命中，但台词源表该条首集台词为占位值")
        if review_old10 == "是":
            old10_note = "已复核到同名来源，旧版若显示较短，多数是截断预览，当前改用完整首集台词"
        status = "已回填" if best.quote_valid else "命中剧名但源表台词为空"
        return (
            MatchResult(
                status=status,
                method="exact",
                quote=best.quote,
                source_sheet=best.source_sheet,
                source_account=best.source_account,
                source_drama_name=best.drama_name,
                score=1.0,
                note="；".join(note_parts),
                review_old10=review_old10,
                old10_note=old10_note,
            ),
            [],
        )

    candidates = build_candidate_rows(title_norm, target_aliases, all_quote_rows)
    if candidates:
        best, score, contain, same_account = candidates[0]
        high_confidence_contain = contain and score >= 0.90
        high_confidence_same_account = same_account and score >= 0.82
        if best.quote_valid and (high_confidence_contain or high_confidence_same_account):
            method = "contain" if high_confidence_contain else "same_account_fuzzy"
            note = f"高可信近似命中，score={score:.4f}"
            return (
                MatchResult(
                    status="已回填",
                    method=method,
                    quote=best.quote,
                    source_sheet=best.source_sheet,
                    source_account=best.source_account,
                    source_drama_name=best.drama_name,
                    score=score,
                    note=note,
                    review_old10=review_old10,
                    old10_note=old10_note,
                ),
                candidates,
            )

        note = (
            f"未自动回填；最佳候选={best.drama_name}；score={score:.4f}"
            f"；same_account={'Y' if same_account else 'N'}；contain={'Y' if contain else 'N'}"
        )
        return (
            MatchResult(
                status="待人工确认",
                method="candidate_only",
                quote="",
                source_sheet="",
                source_account="",
                source_drama_name="",
                score=score,
                note=note,
                review_old10=review_old10,
                old10_note=old10_note,
            ),
            candidates,
        )

    return (
        MatchResult(
            status="未找到可信候选",
            method="unmatched",
            quote="",
            source_sheet="",
            source_account="",
            source_drama_name="",
            score=0.0,
            note="台词表未找到足够可信的同名或近似候选",
            review_old10=review_old10,
            old10_note=old10_note,
        ),
        [],
    )


def ensure_output_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def style_header_row(ws, row_idx: int) -> None:
    for cell in ws[row_idx]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_ALIGNMENT


def apply_body_font(ws, min_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP_ALIGNMENT


def autosize_sheet(ws, max_width: int = 80) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            longest = max((len(line) for line in value.splitlines()), default=0)
            widths[cell.column] = min(max(widths.get(cell.column, 0), longest + 2), max_width)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_table_sheet(wb: Workbook, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    if not rows:
        ws["A1"] = "empty"
        ws["A1"].font = BODY_FONT
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    style_header_row(ws, 1)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    apply_body_font(ws, 2, 1, len(headers))
    ws.freeze_panes = "A2"
    autosize_sheet(ws, max_width=120)


def build_old10_review_rows(
    old10_rows: list[dict[str, str]],
    exact_index: dict[str, list[QuoteRow]],
) -> list[dict[str, Any]]:
    review_rows: list[dict[str, Any]] = []
    for item in old10_rows:
        title_norm = normalize_title(item["short_drama"])
        exact_matches = exact_index.get(title_norm, [])
        if exact_matches:
            best = exact_matches[0]
            matched = quote_preview_matches(best.quote, item["quote_preview"])
            review_rows.append(
                {
                    "priority": item["priority"],
                    "short_drama": item["short_drama"],
                    "旧版台词预览": item["quote_preview"],
                    "复核结果": "对象匹配正常" if matched else "剧名匹配正常，旧版更像截断预览",
                    "当前来源剧名": best.drama_name,
                    "当前来源sheet": best.source_sheet,
                    "当前完整首集台词": best.quote,
                }
            )
        else:
            review_rows.append(
                {
                    "priority": item["priority"],
                    "short_drama": item["short_drama"],
                    "旧版台词预览": item["quote_preview"],
                    "复核结果": "未在新台词表中找到同名来源",
                    "当前来源剧名": "",
                    "当前来源sheet": "",
                    "当前完整首集台词": "",
                }
            )
    return review_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--quotes", default=str(DEFAULT_QUOTE_WORKBOOK))
    parser.add_argument("--target", default=str(DEFAULT_TARGET_WORKBOOK))
    parser.add_argument("--old10", default=str(DEFAULT_OLD_RESULT_WORKBOOK))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT_WORKBOOK))
    args = parser.parse_args()

    quote_path = Path(args.quotes)
    target_path = Path(args.target)
    old10_path = Path(args.old10)
    output_path = Path(args.output)

    ensure_output_parent(output_path)
    shutil.copy2(target_path, output_path)

    quote_rows, exact_index = load_quote_rows(quote_path)
    old10_norm_titles, old10_rows = load_old10_review_set(old10_path)

    wb = load_workbook(output_path)
    ws = wb["all69_details"]
    headers = build_header_map(ws, 1)

    append_headers = [
        "首集台词",
        "台词匹配状态",
        "台词匹配方式",
        "台词来源sheet",
        "台词来源账号",
        "台词来源剧名",
        "台词匹配分数",
        "台词匹配备注",
        "旧10条复核",
        "旧10条复核备注",
    ]
    start_col = ws.max_column + 1
    for offset, header in enumerate(append_headers):
        ws.cell(1, start_col + offset).value = header
    style_header_row(ws, 1)

    audit_rows: list[dict[str, Any]] = []
    unmatched_rows: list[dict[str, Any]] = []
    filled_count = 0
    empty_quote_count = 0
    review_count = 0

    priority_col = headers["测试优先级"]
    novel_col = headers["原著小说名"]
    short_drama_col = headers["短剧名"]
    account_col = headers["原表_账号或页头信息"]

    for row_idx in range(2, ws.max_row + 1):
        priority = normalize_text(ws.cell(row_idx, priority_col).value)
        novel_name = normalize_text(ws.cell(row_idx, novel_col).value)
        short_drama = normalize_text(ws.cell(row_idx, short_drama_col).value)
        account_aliases = extract_target_account_aliases(ws.cell(row_idx, account_col).value)

        match, candidates = pick_match(short_drama, account_aliases, old10_norm_titles, exact_index, quote_rows)

        values = [
            match.quote,
            match.status,
            match.method,
            match.source_sheet,
            match.source_account,
            match.source_drama_name,
            round(match.score, 4) if match.score else "",
            match.note,
            match.review_old10,
            match.old10_note,
        ]
        for offset, value in enumerate(values):
            ws.cell(row_idx, start_col + offset).value = value

        if match.status == "已回填":
            filled_count += 1
        elif match.status == "命中剧名但源表台词为空":
            empty_quote_count += 1

        if match.review_old10 == "是":
            review_count += 1

        audit_rows.append(
            {
                "测试优先级": priority,
                "原著小说名": novel_name,
                "短剧名": short_drama,
                "台词匹配状态": match.status,
                "台词匹配方式": match.method,
                "台词来源剧名": match.source_drama_name,
                "台词来源sheet": match.source_sheet,
                "台词来源账号": match.source_account,
                "台词匹配分数": round(match.score, 4) if match.score else "",
                "旧10条复核": match.review_old10,
                "台词匹配备注": match.note,
            }
        )

        if match.status != "已回填":
            candidate_text = " | ".join(
                [
                    f"{cand.drama_name} @ {cand.source_sheet} score={score:.4f} "
                    f"same_account={'Y' if same_account else 'N'} contain={'Y' if contain else 'N'} "
                    f"quote={'Y' if cand.quote else 'N'}"
                    for cand, score, contain, same_account in candidates
                ]
            )
            unmatched_rows.append(
                {
                    "测试优先级": priority,
                    "原著小说名": novel_name,
                    "短剧名": short_drama,
                    "当前状态": match.status,
                    "未自动回填原因": match.note,
                    "候选列表": candidate_text,
                }
            )

    apply_body_font(ws, 2, start_col, ws.max_column)
    for col_idx in range(start_col, ws.max_column + 1):
        header = normalize_text(ws.cell(1, col_idx).value)
        if header == "首集台词":
            ws.column_dimensions[get_column_letter(col_idx)].width = 120
        elif header in {"台词匹配备注", "旧10条复核备注"}:
            ws.column_dimensions[get_column_letter(col_idx)].width = 40
        else:
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

    old10_review_rows = build_old10_review_rows(old10_rows, exact_index)
    summary_rows = [
        {"指标": "目标表行数", "值": ws.max_row - 1},
        {"指标": "自动回填成功", "值": filled_count},
        {"指标": "命中剧名但源表台词为空", "值": empty_quote_count},
        {"指标": "未自动回填", "值": len(unmatched_rows) - empty_quote_count},
        {"指标": "旧10条复核覆盖数", "值": review_count},
        {"指标": "台词源表总候选数", "值": len(quote_rows)},
    ]

    write_table_sheet(wb, "台词回填汇总", summary_rows)
    write_table_sheet(wb, "台词回填复核", audit_rows)
    write_table_sheet(wb, "未回填清单", unmatched_rows)
    write_table_sheet(wb, "旧10条复核", old10_review_rows)

    wb.save(output_path)

    print(f"output={output_path}")
    print(f"filled={filled_count}")
    print(f"empty_quote={empty_quote_count}")
    print(f"unmatched={len(unmatched_rows)}")
    print(f"old10_reviewed={review_count}")


if __name__ == "__main__":
    main()
