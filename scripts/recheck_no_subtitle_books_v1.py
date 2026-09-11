from __future__ import annotations

import argparse
import csv
import json
import threading
import time
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from requests.adapters import HTTPAdapter


DEFAULT_ENDPOINT = "https://hwzt.ikyuedu.com/content/assets/tmp/batch-episodes-info"
DEFAULT_BATCH_SIZE = 10
DEFAULT_WORKERS = 4
DEFAULT_TIMEOUT = 60
DEFAULT_RETRIES = 2
DEFAULT_SOURCE_CSV = (
    "data_samples/middle_platform_subtitles_20260722_full/logs/no_subtitle_books_20260729_163643.csv"
)

THREAD_LOCAL = threading.local()


@dataclass(frozen=True)
class SeedRow:
    row_index: int
    book_id: str
    seed_title: str
    seed_tag: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Recheck books previously classified as no-subtitle against the middle-platform API.",
    )
    parser.add_argument("--source-csv", default=DEFAULT_SOURCE_CSV)
    parser.add_argument("--endpoint", default=DEFAULT_ENDPOINT)
    parser.add_argument("--output-root", default="data_samples/middle_platform_subtitles_20260722_recheck")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES)
    return parser.parse_args()


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def read_seed_rows(path: Path, *, limit: int = 0) -> list[SeedRow]:
    rows: list[SeedRow] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(
                SeedRow(
                    row_index=int(row["row_index"]),
                    book_id=str(row["book_id"]),
                    seed_title=str(row["seed_title"]),
                    seed_tag=str(row["seed_tag"]),
                )
            )
            if limit > 0 and len(rows) >= limit:
                break
    return rows


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "Content-Type": "application/json",
            "User-Agent": "novel-similarity-middle-platform-recheck/1.0",
            "Connection": "keep-alive",
        }
    )
    adapter = HTTPAdapter(pool_connections=1, pool_maxsize=1, max_retries=0, pool_block=True)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session


def get_session() -> requests.Session:
    session = getattr(THREAD_LOCAL, "session", None)
    if session is None:
        session = build_session()
        THREAD_LOCAL.session = session
    return session


def is_real_subtitle_item(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    return any(
        str(item.get(key) or "").strip()
        for key in ("text", "start", "end", "speaker")
    )


def summarize_payload(seed: SeedRow, payload: dict[str, Any] | None, error_message: str = "") -> dict[str, Any]:
    base = {
        "row_index": seed.row_index,
        "book_id": seed.book_id,
        "seed_title": seed.seed_title,
        "seed_tag": seed.seed_tag,
        "rechecked_at": datetime.now().isoformat(sep=" ", timespec="seconds"),
        "error_message": error_message,
    }
    if not isinstance(payload, dict):
        return {
            **base,
            "found": 0,
            "api_book_name": "",
            "episode_count": 0,
            "total_real_subtitle_lines": 0,
            "first5_real_subtitle_lines": 0,
            "has_any_real_subtitles": 0,
            "first5_has_real_subtitles": 0,
            "status": "missing",
        }

    episodes = payload.get("episodes")
    episode_list = episodes if isinstance(episodes, list) else []
    total_real_subtitle_lines = 0
    first5_real_subtitle_lines = 0
    for idx, episode in enumerate(episode_list):
        subtitle_list = episode.get("subtitles") if isinstance(episode, dict) else []
        subtitle_list = subtitle_list if isinstance(subtitle_list, list) else []
        real_count = sum(1 for item in subtitle_list if is_real_subtitle_item(item))
        total_real_subtitle_lines += real_count
        if idx < 5:
            first5_real_subtitle_lines += real_count

    return {
        **base,
        "found": 1,
        "api_book_name": str(payload.get("bookName") or ""),
        "episode_count": len(episode_list),
        "total_real_subtitle_lines": total_real_subtitle_lines,
        "first5_real_subtitle_lines": first5_real_subtitle_lines,
        "has_any_real_subtitles": int(total_real_subtitle_lines > 0),
        "first5_has_real_subtitles": int(first5_real_subtitle_lines > 0),
        "status": "ok",
    }


def fetch_batch(endpoint: str, seeds: list[SeedRow], *, timeout: int, retries: int) -> tuple[dict[str, Any], str]:
    body = {"bookIds": [seed.book_id for seed in seeds]}
    last_error = ""
    session = get_session()
    for attempt in range(retries + 1):
        try:
            response = session.post(endpoint, json=body, timeout=timeout)
            response.raise_for_status()
            return response.json(), ""
        except (requests.RequestException, ValueError) as exc:
            last_error = str(exc)
            if attempt < retries:
                time.sleep(min(2**attempt, 5))
    return {}, last_error


def chunked(items: list[SeedRow], size: int) -> list[list[SeedRow]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def write_summary_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    if not rows:
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def autosize_worksheet(ws) -> None:  # type: ignore[no-untyped-def]
    widths: dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            text = "" if value is None else str(value)
            widths[idx] = max(widths.get(idx, 0), min(len(text) + 2, 60))
    for idx, width in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width


def write_result_workbook(path: Path, all_rows: list[dict[str, Any]], recovered: list[dict[str, Any]], confirmed: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    workbook = Workbook()
    all_sheet = workbook.active
    all_sheet.title = "二次复查全量"
    recovered_sheet = workbook.create_sheet("补回有字幕")
    confirmed_sheet = workbook.create_sheet("确认无字幕")
    headers = list(all_rows[0].keys()) if all_rows else [
        "row_index", "book_id", "seed_title", "seed_tag", "rechecked_at", "error_message",
        "found", "api_book_name", "episode_count", "total_real_subtitle_lines",
        "first5_real_subtitle_lines", "has_any_real_subtitles", "first5_has_real_subtitles", "status",
    ]
    for sheet in (all_sheet, recovered_sheet, confirmed_sheet):
        sheet.append(headers)
    for row in all_rows:
        all_sheet.append([row.get(header, "") for header in headers])
    for row in recovered:
        recovered_sheet.append([row.get(header, "") for header in headers])
    for row in confirmed:
        confirmed_sheet.append([row.get(header, "") for header in headers])
    for sheet in (all_sheet, recovered_sheet, confirmed_sheet):
        autosize_worksheet(sheet)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    source_csv = Path(args.source_csv)
    if not source_csv.is_absolute():
        source_csv = (Path.cwd() / source_csv).resolve()
    if not source_csv.exists():
        raise FileNotFoundError(f"Source CSV not found: {source_csv}")

    seeds = read_seed_rows(source_csv, limit=max(int(args.limit or 0), 0))
    batch_size = max(int(args.batch_size or DEFAULT_BATCH_SIZE), 1)
    if batch_size > 10:
        raise ValueError("batch-size cannot exceed 10.")
    workers = max(int(args.workers or DEFAULT_WORKERS), 1)
    timeout = max(int(args.timeout or DEFAULT_TIMEOUT), 1)
    retries = max(int(args.retries or DEFAULT_RETRIES), 0)

    output_root = Path(args.output_root)
    if not output_root.is_absolute():
        output_root = (Path.cwd() / output_root).resolve()
    ensure_dir(output_root)
    raw_root = output_root / "raw_books"
    logs_root = output_root / "logs"
    ensure_dir(raw_root)
    ensure_dir(logs_root)

    batches = chunked(seeds, batch_size)
    total_batches = len(batches)
    started_at = time.perf_counter()
    progress_lock = threading.Lock()
    completed_batches = 0
    completed_books = 0
    all_rows: list[dict[str, Any]] = []

    def run_one(batch: list[SeedRow]) -> list[dict[str, Any]]:
        nonlocal completed_batches
        nonlocal completed_books
        api_response, error_message = fetch_batch(args.endpoint, batch, timeout=timeout, retries=retries)
        response_data = api_response.get("data") if isinstance(api_response, dict) else {}
        if not isinstance(response_data, dict):
            response_data = {}
        rows: list[dict[str, Any]] = []
        for seed in batch:
            item_payload = response_data.get(seed.book_id)
            if isinstance(item_payload, dict):
                prefix = seed.book_id[:4] if len(seed.book_id) >= 4 else "misc"
                raw_path = raw_root / prefix / f"{seed.book_id}.json"
                ensure_dir(raw_path.parent)
                raw_path.write_text(json.dumps(item_payload, ensure_ascii=False, indent=2), encoding="utf-8")
            rows.append(
                summarize_payload(
                    seed,
                    item_payload if isinstance(item_payload, dict) else None,
                    error_message="" if item_payload is not None else error_message,
                )
            )
        with progress_lock:
            completed_batches += 1
            completed_books += len(batch)
            if completed_batches % 20 == 0 or completed_batches == total_batches:
                elapsed = time.perf_counter() - started_at
                rate = 0.0 if elapsed <= 0 else completed_books / elapsed
                print(
                    f"progress batches={completed_batches}/{total_batches} "
                    f"books={completed_books}/{len(seeds)} rate={rate:.2f} books/s"
                )
        return rows

    in_flight: dict[Future[list[dict[str, Any]]], list[SeedRow]] = {}
    batch_index = 0
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="subtitle-recheck") as executor:
        while batch_index < total_batches and len(in_flight) < workers:
            batch = batches[batch_index]
            in_flight[executor.submit(run_one, batch)] = batch
            batch_index += 1
        while in_flight:
            done, _ = wait(set(in_flight.keys()), return_when=FIRST_COMPLETED)
            for future in done:
                in_flight.pop(future)
                all_rows.extend(future.result())
                if batch_index < total_batches:
                    batch = batches[batch_index]
                    in_flight[executor.submit(run_one, batch)] = batch
                    batch_index += 1

    all_rows.sort(key=lambda row: (int(row["row_index"]), str(row["book_id"])))
    recovered = [row for row in all_rows if int(row["has_any_real_subtitles"]) > 0]
    confirmed = [row for row in all_rows if int(row["has_any_real_subtitles"]) <= 0]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    all_csv = logs_root / f"recheck_all_{timestamp}.csv"
    recovered_csv = logs_root / f"recheck_recovered_subtitles_{timestamp}.csv"
    confirmed_csv = logs_root / f"recheck_confirmed_no_subtitles_{timestamp}.csv"
    workbook_path = logs_root / f"recheck_result_{timestamp}.xlsx"
    summary_path = logs_root / f"recheck_summary_{timestamp}.txt"

    write_summary_csv(all_csv, all_rows)
    write_summary_csv(recovered_csv, recovered)
    write_summary_csv(confirmed_csv, confirmed)
    write_result_workbook(workbook_path, all_rows, recovered, confirmed)
    with summary_path.open("w", encoding="utf-8") as f:
        f.write(f"generated_at={datetime.now().isoformat(sep=' ', timespec='seconds')}\n")
        f.write(f"seed_count={len(seeds)}\n")
        f.write(f"batch_count={total_batches}\n")
        f.write(f"recovered_count={len(recovered)}\n")
        f.write(f"confirmed_no_subtitle_count={len(confirmed)}\n")
        f.write(f"recovered_first5_count={sum(int(row['first5_has_real_subtitles']) for row in recovered)}\n")

    elapsed = time.perf_counter() - started_at
    print(f"OK seed_count={len(seeds)} batch_count={total_batches} elapsed_seconds={elapsed:.2f}")
    print(f"OK all_csv={all_csv}")
    print(f"OK recovered_csv={recovered_csv}")
    print(f"OK confirmed_csv={confirmed_csv}")
    print(f"OK workbook={workbook_path}")
    print(f"OK summary={summary_path}")


if __name__ == "__main__":
    main()
