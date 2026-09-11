from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a deduplicated Book ID delta for subtitle collection.")
    parser.add_argument("--current-xlsx", required=True)
    parser.add_argument("--previous-coverage-csv", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--output-csv", required=True)
    parser.add_argument("--summary", required=True)
    return parser.parse_args()


def normalize_id(value: object) -> str:
    return str(value or "").strip()


def load_previous_ids(path: Path) -> set[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {
            normalize_id(row.get("book_id"))
            for row in csv.DictReader(handle)
            if normalize_id(row.get("book_id"))
        }


def load_current_rows(path: Path) -> tuple[list[dict[str, object]], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: list[dict[str, object]] = []
        seen: set[str] = set()
        duplicate_count = 0
        for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            book_id = normalize_id(values[0] if len(values) > 0 else "")
            if not book_id:
                continue
            if book_id in seen:
                duplicate_count += 1
                continue
            seen.add(book_id)
            rows.append(
                {
                    "短剧ID": book_id,
                    "短剧名称": str(values[1] or "").strip() if len(values) > 1 else "",
                    "短剧别名": str(values[2] or "").strip() if len(values) > 2 else "",
                    "源表行号": source_row,
                }
            )
        return rows, duplicate_count
    finally:
        workbook.close()


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["短剧ID", "短剧名称", "短剧别名", "源表行号"])
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "video"
    headers = ["短剧ID", "短剧名称", "短剧别名", "源表行号"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row[column] for column in headers])
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 48
    sheet.column_dimensions["C"].width = 58
    sheet.column_dimensions["D"].width = 12
    workbook.save(path)


def main() -> None:
    args = parse_args()
    current_path = Path(args.current_xlsx).resolve()
    previous_path = Path(args.previous_coverage_csv).resolve()
    previous_ids = load_previous_ids(previous_path)
    current_rows, duplicate_count = load_current_rows(current_path)
    current_ids = {str(row["短剧ID"]) for row in current_rows}
    delta_rows = [row for row in current_rows if str(row["短剧ID"]) not in previous_ids]
    old_only = previous_ids - current_ids

    write_xlsx(Path(args.output_xlsx).resolve(), delta_rows)
    write_csv(Path(args.output_csv).resolve(), delta_rows)
    summary_path = Path(args.summary).resolve()
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(
        "\n".join(
            [
                f"current_unique_ids={len(current_ids)}",
                f"current_duplicate_rows={duplicate_count}",
                f"previous_unique_ids={len(previous_ids)}",
                f"intersection={len(current_ids & previous_ids)}",
                f"delta_unique_ids={len(delta_rows)}",
                f"previous_only_ids={len(old_only)}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"OK delta_unique_ids={len(delta_rows)}")
    print(f"OK output_xlsx={Path(args.output_xlsx).resolve()}")
    print(f"OK output_csv={Path(args.output_csv).resolve()}")
    print(f"OK summary={summary_path}")


if __name__ == "__main__":
    main()
