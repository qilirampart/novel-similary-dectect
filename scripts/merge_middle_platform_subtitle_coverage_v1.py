from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Merge previous and delta subtitle coverage reports.")
    parser.add_argument("--previous-csv", required=True)
    parser.add_argument("--delta-csv", required=True)
    parser.add_argument("--previous-raw-root", required=True)
    parser.add_argument("--delta-raw-root", required=True)
    parser.add_argument("--output-csv", required=True)
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--summary", required=True)
    return parser.parse_args()


def read_rows(path: Path, source_batch: str, raw_root: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        book_id = str(row.get("book_id") or "").strip()
        prefix = book_id[:4] if len(book_id) >= 4 else "misc"
        payload_path = raw_root / prefix / f"{book_id}.json"
        row["source_batch"] = source_batch
        row["raw_payload_path"] = str(payload_path) if payload_path.exists() else ""
    return rows


def integer(row: dict[str, str], key: str) -> int:
    try:
        return int(row.get(key) or 0)
    except (TypeError, ValueError):
        return 0


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=False)
    summary = workbook.active
    summary.title = "汇总"
    detail = workbook.create_sheet("全量覆盖明细")
    first5 = workbook.create_sheet("前5集有真实字幕")
    no_subtitle = workbook.create_sheet("无真实字幕")

    metrics = [
        ("唯一Book ID", len(rows)),
        ("接口可查", sum(integer(row, "found") for row in rows)),
        ("接口无返回", sum(integer(row, "found") == 0 for row in rows)),
        ("全剧有真实字幕", sum(integer(row, "has_any_real_subtitles") for row in rows)),
        ("前5集有真实字幕", sum(integer(row, "first5_has_real_subtitles") for row in rows)),
        ("仅第6集以后有真实字幕", sum(integer(row, "has_any_real_subtitles") == 1 and integer(row, "first5_has_real_subtitles") == 0 for row in rows)),
        ("无真实字幕", sum(integer(row, "no_subtitle_file") for row in rows)),
        ("真实字幕总行数", sum(integer(row, "total_real_subtitle_lines") for row in rows)),
        ("前5集真实字幕总行数", sum(integer(row, "first5_real_subtitle_lines") for row in rows)),
    ]
    summary.append(["指标", "数量"])
    for metric in metrics:
        summary.append(metric)

    for sheet in (detail, first5, no_subtitle):
        sheet.append(fields)
    for row in rows:
        values = [row.get(field, "") for field in fields]
        detail.append(values)
        if integer(row, "first5_has_real_subtitles") > 0:
            first5.append(values)
        if integer(row, "no_subtitle_file") > 0:
            no_subtitle.append(values)

    header_fill = PatternFill("solid", fgColor="17365D")
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18
    for sheet in (detail, first5, no_subtitle):
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 42
        sheet.column_dimensions["D"].width = 52
        sheet.column_dimensions["Y"].width = 20
        sheet.column_dimensions["Z"].width = 80
    workbook.save(path)


def main() -> None:
    args = parse_args()
    previous_rows = read_rows(Path(args.previous_csv), "20260722_full", Path(args.previous_raw_root))
    delta_rows = read_rows(Path(args.delta_csv), "20260901_delta", Path(args.delta_raw_root))
    merged_by_id: dict[str, dict[str, str]] = {}
    for row in previous_rows + delta_rows:
        book_id = str(row.get("book_id") or "").strip()
        if book_id:
            merged_by_id[book_id] = row
    rows = sorted(merged_by_id.values(), key=lambda row: str(row.get("book_id") or ""))
    fields = list(rows[0].keys())
    write_csv(Path(args.output_csv), rows, fields)
    write_xlsx(Path(args.output_xlsx), rows, fields)

    summary_lines = [
        f"previous_rows={len(previous_rows)}",
        f"delta_rows={len(delta_rows)}",
        f"merged_unique_book_ids={len(rows)}",
        f"api_found={sum(integer(row, 'found') for row in rows)}",
        f"api_missing={sum(integer(row, 'found') == 0 for row in rows)}",
        f"first5_has_real_subtitles={sum(integer(row, 'first5_has_real_subtitles') for row in rows)}",
        f"any_real_subtitles={sum(integer(row, 'has_any_real_subtitles') for row in rows)}",
        f"later_only_has_real_subtitles={sum(integer(row, 'has_any_real_subtitles') == 1 and integer(row, 'first5_has_real_subtitles') == 0 for row in rows)}",
        f"no_real_subtitles={sum(integer(row, 'no_subtitle_file') for row in rows)}",
        f"total_real_subtitle_lines={sum(integer(row, 'total_real_subtitle_lines') for row in rows)}",
        f"first5_real_subtitle_lines={sum(integer(row, 'first5_real_subtitle_lines') for row in rows)}",
    ]
    summary_path = Path(args.summary)
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    print("\n".join(summary_lines))


if __name__ == "__main__":
    main()
