from __future__ import annotations

import argparse
import csv
import json
import math
import threading
import time
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook, load_workbook
from requests.adapters import HTTPAdapter


DEFAULT_ENDPOINT = "https://hwzt.ikyuedu.com/content/assets/tmp/batch-episodes-info"
DEFAULT_BATCH_SIZE = 10
DEFAULT_WORKERS = 4
DEFAULT_TIMEOUT = 60
DEFAULT_RETRIES = 2

THREAD_LOCAL = threading.local()


@dataclass(frozen=True)
class BookSeed:
    row_index: int
    book_id: str
    title: str
    tag: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Collect full-episode subtitle payloads from middle-platform batch episode API.",
    )
    parser.add_argument(
        "--xlsx",
        default="短剧上架状态-20260722.xlsx",
        help="Workbook containing 短剧 ID / 短剧名 / 标记 columns.",
    )
    parser.add_argument(
        "--endpoint",
        default=DEFAULT_ENDPOINT,
        help="Middle-platform batch episode API endpoint.",
    )
    parser.add_argument("--output-root", default="data_samples/middle_platform_subtitles_20260722")
    parser.add_argument("--limit", type=int, default=0, help="Only process first N seed rows.")
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE)
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def load_seed_rows(path: Path, *, limit: int = 0) -> list[BookSeed]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook.active
        rows: list[BookSeed] = []
        seen: set[str] = set()
        for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            book_id = str(row[0] or "").strip()
            if not book_id or book_id in seen:
                continue
            seen.add(book_id)
            rows.append(
                BookSeed(
                    row_index=row_index,
                    book_id=book_id,
                    title=str(row[1] or "").strip(),
                    tag=str(row[2] or "").strip(),
                )
            )
            if limit > 0 and len(rows) >= limit:
                break
        return rows
    finally:
        workbook.close()


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "Content-Type": "application/json",
            "User-Agent": "novel-similarity-middle-platform-collector/1.0",
            "Connection": "keep-alive",
        }
    )
    adapter = HTTPAdapter(pool_connections=1, pool_maxsize=1, max_retries=0, pool_block=True)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session


def get_session() -> requests.Session:
    session = getattr(THREAD_LOCAL, "session", None)
    if session is None:
        session = build_session()
        THREAD_LOCAL.session = session
    return session


def payload_path(raw_root: Path, book_id: str) -> Path:
    prefix = book_id[:4] if len(book_id) >= 4 else "misc"
    return raw_root / prefix / f"{book_id}.json"


def is_real_subtitle_item(item: Any) -> bool:
    if not isinstance(item, dict):
        return False
    text = str(item.get("text") or "").strip()
    start = str(item.get("start") or "").strip()
    end = str(item.get("end") or "").strip()
    speaker = str(item.get("speaker") or "").strip()
    return bool(text or start or end or speaker)


def analyze_payload(
    seed: BookSeed,
    payload: dict[str, Any] | None,
    *,
    from_cache: bool,
    error_message: str = "",
) -> dict[str, Any]:
    base = {
        "row_index": seed.row_index,
        "book_id": seed.book_id,
        "seed_title": seed.title,
        "seed_tag": seed.tag,
        "from_cache": int(from_cache),
        "error_message": error_message,
        "collected_at": datetime.now().isoformat(sep=" ", timespec="seconds"),
    }
    if not isinstance(payload, dict):
        return {
            **base,
            "found": 0,
            "api_book_name": "",
            "episode_count": 0,
            "episodes_with_subtitles": 0,
            "episodes_with_real_subtitles": 0,
            "total_subtitle_slots": 0,
            "total_real_subtitle_lines": 0,
            "first5_episode_count": 0,
            "first5_episodes_with_subtitles": 0,
            "first5_episodes_with_real_subtitles": 0,
            "first5_subtitle_slots": 0,
            "first5_real_subtitle_lines": 0,
            "first5_placeholder_only_episodes": 0,
            "first5_has_real_subtitles": 0,
            "has_any_real_subtitles": 0,
            "no_subtitle_file": 1,
            "status": "missing",
        }

    episodes = payload.get("episodes")
    episode_list = episodes if isinstance(episodes, list) else []
    total_subtitle_slots = 0
    total_real_subtitle_lines = 0
    episodes_with_subtitles = 0
    episodes_with_real_subtitles = 0
    first5 = episode_list[:5]
    first5_subtitle_slots = 0
    first5_real_subtitle_lines = 0
    first5_episodes_with_subtitles = 0
    first5_episodes_with_real_subtitles = 0
    first5_placeholder_only_episodes = 0

    for idx, episode in enumerate(episode_list):
        subtitles = episode.get("subtitles") if isinstance(episode, dict) else []
        subtitle_list = subtitles if isinstance(subtitles, list) else []
        real_count = sum(1 for item in subtitle_list if is_real_subtitle_item(item))
        slot_count = len(subtitle_list)
        total_subtitle_slots += slot_count
        total_real_subtitle_lines += real_count
        if slot_count > 0:
            episodes_with_subtitles += 1
        if real_count > 0:
            episodes_with_real_subtitles += 1
        if idx < 5:
            first5_subtitle_slots += slot_count
            first5_real_subtitle_lines += real_count
            if slot_count > 0:
                first5_episodes_with_subtitles += 1
            if real_count > 0:
                first5_episodes_with_real_subtitles += 1
            if slot_count > 0 and real_count == 0:
                first5_placeholder_only_episodes += 1

    return {
        **base,
        "found": 1,
        "api_book_name": str(payload.get("bookName") or ""),
        "episode_count": len(episode_list),
        "episodes_with_subtitles": episodes_with_subtitles,
        "episodes_with_real_subtitles": episodes_with_real_subtitles,
        "total_subtitle_slots": total_subtitle_slots,
        "total_real_subtitle_lines": total_real_subtitle_lines,
        "first5_episode_count": len(first5),
        "first5_episodes_with_subtitles": first5_episodes_with_subtitles,
        "first5_episodes_with_real_subtitles": first5_episodes_with_real_subtitles,
        "first5_subtitle_slots": first5_subtitle_slots,
        "first5_real_subtitle_lines": first5_real_subtitle_lines,
        "first5_placeholder_only_episodes": first5_placeholder_only_episodes,
        "first5_has_real_subtitles": int(first5_real_subtitle_lines > 0),
        "has_any_real_subtitles": int(total_real_subtitle_lines > 0),
        "no_subtitle_file": int(total_real_subtitle_lines == 0),
        "status": "ok",
    }


def write_json(path: Path, payload: dict[str, Any]) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def fetch_batch(
    endpoint: str,
    seeds: list[BookSeed],
    *,
    timeout: int,
    retries: int,
) -> tuple[dict[str, Any], str]:
    body = {"bookIds": [seed.book_id for seed in seeds]}
    last_error = ""
    session = get_session()
    for attempt in range(retries + 1):
        try:
            response = session.post(endpoint, json=body, timeout=timeout)
            response.raise_for_status()
            data = response.json()
            return data, ""
        except (requests.RequestException, ValueError) as exc:
            last_error = str(exc)
            if attempt < retries:
                time.sleep(min(2**attempt, 5))
    return {}, last_error


def process_batch(
    endpoint: str,
    seeds: list[BookSeed],
    raw_root: Path,
    *,
    timeout: int,
    retries: int,
    overwrite: bool,
) -> list[dict[str, Any]]:
    cached_rows: list[dict[str, Any]] = []
    missing_seeds: list[BookSeed] = []
    cached_payloads: dict[str, dict[str, Any] | None] = {}
    for seed in seeds:
        path = payload_path(raw_root, seed.book_id)
        if path.exists() and not overwrite:
            try:
                cached_payloads[seed.book_id] = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                missing_seeds.append(seed)
            else:
                continue
        else:
            missing_seeds.append(seed)

    for seed in seeds:
        if seed.book_id in cached_payloads:
            cached_rows.append(
                analyze_payload(
                    seed,
                    cached_payloads.get(seed.book_id),
                    from_cache=True,
                )
            )

    fetched_rows: list[dict[str, Any]] = []
    if not missing_seeds:
        return cached_rows

    api_response, error_message = fetch_batch(
        endpoint,
        missing_seeds,
        timeout=timeout,
        retries=retries,
    )
    response_data = api_response.get("data") if isinstance(api_response, dict) else {}
    if not isinstance(response_data, dict):
        response_data = {}

    for seed in missing_seeds:
        item_payload = response_data.get(seed.book_id)
        if isinstance(item_payload, dict):
            write_json(payload_path(raw_root, seed.book_id), item_payload)
        fetched_rows.append(
            analyze_payload(
                seed,
                item_payload if isinstance(item_payload, dict) else None,
                from_cache=False,
                error_message="" if item_payload is not None else error_message,
            )
        )
    return cached_rows + fetched_rows


def chunked(items: list[BookSeed], size: int) -> list[list[BookSeed]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def write_summary_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    if not rows:
        raise ValueError("No rows to write")
    fieldnames = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_run_summary(path: Path, rows: list[dict[str, Any]], *, seed_count: int, batch_count: int) -> None:
    found_count = sum(int(row["found"]) for row in rows)
    first5_real_count = sum(int(row["first5_has_real_subtitles"]) for row in rows)
    placeholder_only_count = sum(
        1
        for row in rows
        if int(row["first5_episodes_with_subtitles"]) > 0 and int(row["first5_real_subtitle_lines"]) == 0
    )
    total_real_lines = sum(int(row["total_real_subtitle_lines"]) for row in rows)
    first5_real_lines = sum(int(row["first5_real_subtitle_lines"]) for row in rows)
    with path.open("w", encoding="utf-8") as f:
        f.write(f"generated_at={datetime.now().isoformat(sep=' ', timespec='seconds')}\n")
        f.write(f"seed_count={seed_count}\n")
        f.write(f"batch_count={batch_count}\n")
        f.write(f"found_count={found_count}\n")
        f.write(f"first5_has_real_subtitles_count={first5_real_count}\n")
        f.write(f"first5_placeholder_only_count={placeholder_only_count}\n")
        f.write(f"total_real_subtitle_lines={total_real_lines}\n")
        f.write(f"first5_real_subtitle_lines={first5_real_lines}\n")


def autosize_worksheet(ws) -> None:  # type: ignore[no-untyped-def]
    widths: dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            text = "" if value is None else str(value)
            widths[idx] = max(widths.get(idx, 0), min(len(text) + 2, 60))
    for idx, width in widths.items():
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else ws.cell(row=1, column=idx).column_letter].width = width


def write_result_workbook(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "全量结果"
    no_subtitle_sheet = workbook.create_sheet("无字幕剧单")

    ordered_columns = [
        "row_index",
        "book_id",
        "seed_title",
        "seed_tag",
        "found",
        "api_book_name",
        "episode_count",
        "episodes_with_subtitles",
        "episodes_with_real_subtitles",
        "total_subtitle_slots",
        "total_real_subtitle_lines",
        "first5_episode_count",
        "first5_episodes_with_subtitles",
        "first5_episodes_with_real_subtitles",
        "first5_subtitle_slots",
        "first5_real_subtitle_lines",
        "first5_placeholder_only_episodes",
        "first5_has_real_subtitles",
        "has_any_real_subtitles",
        "no_subtitle_file",
        "from_cache",
        "status",
        "error_message",
        "collected_at",
    ]
    summary_sheet.append(ordered_columns)
    no_subtitle_sheet.append(ordered_columns)

    for row in rows:
        values = [row.get(column, "") for column in ordered_columns]
        summary_sheet.append(values)
        if int(row.get("no_subtitle_file", 0)) > 0:
            no_subtitle_sheet.append(values)

    autosize_worksheet(summary_sheet)
    autosize_worksheet(no_subtitle_sheet)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.xlsx)
    if not workbook_path.is_absolute():
        workbook_path = (Path.cwd() / workbook_path).resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    seeds = load_seed_rows(workbook_path, limit=max(int(args.limit or 0), 0))
    batch_size = max(int(args.batch_size or DEFAULT_BATCH_SIZE), 1)
    if batch_size > 10:
        raise ValueError("batch-size cannot exceed 10 because the API only accepts at most 10 bookIds.")
    workers = max(int(args.workers or DEFAULT_WORKERS), 1)
    timeout = max(int(args.timeout or DEFAULT_TIMEOUT), 1)
    retries = max(int(args.retries or DEFAULT_RETRIES), 0)

    output_root = Path(args.output_root)
    if not output_root.is_absolute():
        output_root = (Path.cwd() / output_root).resolve()
    raw_root = output_root / "raw_books"
    logs_root = output_root / "logs"
    ensure_dir(raw_root)
    ensure_dir(logs_root)

    batches = chunked(seeds, batch_size)
    total_batches = len(batches)
    started_at = time.perf_counter()
    progress_lock = threading.Lock()
    completed_batches = 0
    completed_books = 0
    all_rows: list[dict[str, Any]] = []

    def run_one(batch: list[BookSeed]) -> list[dict[str, Any]]:
        nonlocal completed_batches
        nonlocal completed_books
        rows = process_batch(
            args.endpoint,
            batch,
            raw_root,
            timeout=timeout,
            retries=retries,
            overwrite=bool(args.overwrite),
        )
        with progress_lock:
            completed_batches += 1
            completed_books += len(batch)
            if completed_batches % 20 == 0 or completed_batches == total_batches:
                elapsed = time.perf_counter() - started_at
                rate = 0.0 if elapsed <= 0 else completed_books / elapsed
                print(
                    f"progress batches={completed_batches}/{total_batches} "
                    f"books={completed_books}/{len(seeds)} rate={rate:.2f} books/s"
                )
        return rows

    in_flight: dict[Future[list[dict[str, Any]]], list[BookSeed]] = {}
    batch_index = 0
    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="subtitle-collect") as executor:
        while batch_index < total_batches and len(in_flight) < workers:
            batch = batches[batch_index]
            in_flight[executor.submit(run_one, batch)] = batch
            batch_index += 1

        while in_flight:
            done, _ = wait(set(in_flight.keys()), return_when=FIRST_COMPLETED)
            for future in done:
                batch = in_flight.pop(future)
                all_rows.extend(future.result())
                if batch_index < total_batches:
                    next_batch = batches[batch_index]
                    in_flight[executor.submit(run_one, next_batch)] = next_batch
                    batch_index += 1

    all_rows.sort(key=lambda row: (int(row["row_index"]), str(row["book_id"])))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary_csv = logs_root / f"book_subtitle_coverage_{timestamp}.csv"
    write_summary_csv(summary_csv, all_rows)
    run_summary = logs_root / f"run_summary_{timestamp}.txt"
    write_run_summary(run_summary, all_rows, seed_count=len(seeds), batch_count=total_batches)
    result_xlsx = logs_root / f"book_subtitle_coverage_{timestamp}.xlsx"
    write_result_workbook(result_xlsx, all_rows)

    first5_real_csv = logs_root / f"first5_real_subtitles_{timestamp}.csv"
    first5_rows = [row for row in all_rows if int(row["first5_has_real_subtitles"]) > 0]
    no_subtitle_csv = logs_root / f"no_subtitle_books_{timestamp}.csv"
    no_subtitle_rows = [row for row in all_rows if int(row["no_subtitle_file"]) > 0]
    if first5_rows:
        write_summary_csv(first5_real_csv, first5_rows)
    if no_subtitle_rows:
        write_summary_csv(no_subtitle_csv, no_subtitle_rows)

    elapsed = time.perf_counter() - started_at
    print(f"OK seed_count={len(seeds)} batch_count={total_batches} elapsed_seconds={elapsed:.2f}")
    print(f"OK summary_csv={summary_csv}")
    print(f"OK result_xlsx={result_xlsx}")
    print(f"OK run_summary={run_summary}")
    if first5_rows:
        print(f"OK first5_real_csv={first5_real_csv}")
    if no_subtitle_rows:
        print(f"OK no_subtitle_csv={no_subtitle_csv}")
    print(f"OK raw_root={raw_root}")


if __name__ == "__main__":
    main()
