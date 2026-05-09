from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_PATH = Path("dialogue_quote_compare_results.xlsx")
OUTPUT_PATH = Path("data_samples/batch_uploads/batch_report_samples_9rows_2026-05-08.xlsx")
TARGET_ROWS = [6, 11, 12, 13, 16, 22, 26, 28, 8]


def main() -> None:
    wb = load_workbook(SOURCE_PATH, read_only=True, data_only=True)
    try:
        ws = wb["Sheet1"]
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.title = "batch_input"

        headers = ["source_ref", "excel_row", "novel_name", "query_text"]
        out_ws.append(headers)

        for row_idx in TARGET_ROWS:
            novel_name = ws.cell(row_idx, 2).value or ""
            query_text = ws.cell(row_idx, 28).value or ""
            out_ws.append(
                [
                    f"excel_row_{row_idx}",
                    row_idx,
                    str(novel_name).strip(),
                    str(query_text).strip(),
                ]
            )

        header_fill = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
        header_font = Font(name="Arial", bold=True, color="1F3A5F")
        body_font = Font(name="Arial")
        wrap_alignment = Alignment(vertical="top", wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center")

        for cell in out_ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment

        for row in out_ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = wrap_alignment

        out_ws.column_dimensions["A"].width = 18
        out_ws.column_dimensions["B"].width = 10
        out_ws.column_dimensions["C"].width = 26
        out_ws.column_dimensions["D"].width = 180

        for row_idx in range(2, out_ws.max_row + 1):
            out_ws.row_dimensions[row_idx].height = 72

        out_ws.freeze_panes = "A2"
        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        out_wb.save(OUTPUT_PATH)
    finally:
        wb.close()

    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
