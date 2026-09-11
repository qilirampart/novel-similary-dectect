from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Enrich a cloud rerun workbook with payload-level candidate metrics.")
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--raw", required=True)
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def as_dict(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    try:
        parsed = json.loads(str(value or "{}"))
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def main() -> int:
    args = parse_args()
    workbook = openpyxl.load_workbook(args.workbook)
    worksheet = workbook.active
    raw = json.loads(Path(args.raw).read_text(encoding="utf-8"))
    raw_items = raw.get("items") if isinstance(raw, dict) else []
    payload_by_order: dict[int, dict[str, Any]] = {}
    for item in raw_items or []:
        try:
            order = int(item.get("item_order"))
        except (TypeError, ValueError):
            continue
        payload_by_order[order] = as_dict(item.get("result_payload_json"))

    headers = [cell.value for cell in worksheet[1]]
    header_index = {str(value): index + 1 for index, value in enumerate(headers) if value}
    new_headers = [
        "语义状态（payload）",
        "第一候选词法召回分数",
        "第一候选查询覆盖率",
        "第一候选证据覆盖率",
        "第一候选共享三元词数",
    ]
    start_column = worksheet.max_column + 1
    for offset, header in enumerate(new_headers):
        column = start_column + offset
        cell = worksheet.cell(1, column, header)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    order_column = header_index["序号"]
    json_column = header_index["第一候选完整JSON"]
    semantic_column = header_index["语义状态"]
    for row_number in range(2, worksheet.max_row + 1):
        order = int(worksheet.cell(row_number, order_column).value)
        payload = payload_by_order.get(order, {})
        candidate = as_dict(worksheet.cell(row_number, json_column).value)
        metrics = candidate.get("match_metrics") if isinstance(candidate.get("match_metrics"), dict) else {}
        semantic_status = payload.get("semantic_status") or ""
        worksheet.cell(row_number, semantic_column, semantic_status)
        values = [
            semantic_status,
            candidate.get("best_lexical_score", ""),
            metrics.get("query_coverage_rate", ""),
            metrics.get("evidence_coverage_rate", ""),
            metrics.get("shared_trigram_count", ""),
        ]
        for offset, value in enumerate(values):
            cell = worksheet.cell(row_number, start_column + offset, value)
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
    widths = [22, 22, 22, 22, 24]
    for offset, width in enumerate(widths):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(start_column + offset)].width = width
    for table in worksheet.tables.values():
        table.ref = f"A1:{openpyxl.utils.get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
