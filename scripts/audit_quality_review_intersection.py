from __future__ import annotations

import argparse
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_SCAN_ROWS = 6
HEADER_SCAN_COLS = 24

HAS_ORIGINAL_POSITIVE_EXACT = {
    "是",
    "有",
    "有原著",
    "已导入中台",
}
HAS_ORIGINAL_NEGATIVE_SUBSTRINGS = (
    "否",
    "无",
    "搜不到",
    "未找到",
    "未检索到",
    "已下架",
    "重复",
)

MID_FOUND_POSITIVE_EXACT = {
    "是",
    "有",
    "能",
    "已导入中台",
}
MID_FOUND_NEGATIVE_SUBSTRINGS = (
    "否",
    "无",
    "搜不到",
    "未找到",
    "未检索到",
    "未侵权",
)

CURRENT_WORKBOOK_COLUMN_ALIASES = {
    "source_sheet": ("来源sheet",),
    "source_row": ("原表Excel行号",),
    "original_novel_name": ("原著小说名", "原表_原著小说名称"),
    "dataset_key": ("小说库数据集",),
    "book_ext_id": ("小说book_ext_id",),
    "book_name": ("库内书名",),
    "short_drama_name": ("短剧名",),
}


@dataclass
class SourceRow:
    source_sheet: str
    source_row: int
    drama_name: str
    original_novel_name: str
    has_original_raw: str
    has_original_class: str
    mid_found_raw: str
    mid_found_class: str
    original_cp_name: str
    header_row: int


@dataclass
class DbBook:
    dataset_key: str
    book_ext_id: str
    book_name: str


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    return text.strip()


def normalize_header(value: Any) -> str:
    text = normalize_text(value).lower()
    text = re.sub(r"\s+", "", text)
    return text.replace("（", "(").replace("）", ")")


def strict_title_key(value: Any) -> str:
    text = normalize_text(value).lower()
    text = re.sub(r"\s+", "", text)
    return re.sub(r"[《》「」『』“”‘’\"'`]", "", text)


def loose_title_key(value: Any) -> str:
    text = strict_title_key(value)
    return re.sub(r"[^\u4e00-\u9fff\w]", "", text)


def classify_has_original(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        return "blank"
    if text.startswith("是") or text in HAS_ORIGINAL_POSITIVE_EXACT:
        return "yes"
    if any(token in text for token in HAS_ORIGINAL_NEGATIVE_SUBSTRINGS):
        return "no"
    return "uncertain"


def classify_mid_found(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        return "blank"
    if text.startswith("是") or text in MID_FOUND_POSITIVE_EXACT:
        return "yes"
    if any(token in text for token in MID_FOUND_NEGATIVE_SUBSTRINGS):
        return "no"
    return "uncertain"


def parse_row_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else normalize_text(value)
    text = normalize_text(value)
    if re.fullmatch(r"\d+(\.0+)?", text):
        return str(int(float(text)))
    return text


def detect_header_row(ws) -> int | None:
    best_row = None
    best_score = -1

    for row_idx in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
        cells = [
            normalize_header(ws.cell(row_idx, col_idx).value)
            for col_idx in range(1, min(ws.max_column, HEADER_SCAN_COLS) + 1)
        ]
        joined = "|".join(cells)
        score = sum(
            token in joined
            for token in (
                "短剧名称",
                "剧名",
                "是否有原著小说",
                "原著小说名称",
                "原著小说是否在中台能查找到",
                "原著小说cp名称",
            )
        )
        if score > best_score:
            best_score = score
            best_row = row_idx

    if best_score < 3:
        return None
    return best_row


def build_header_map(ws, header_row: int) -> dict[str, int]:
    header_map: dict[str, int] = {}

    for col_idx in range(1, ws.max_column + 1):
        header = normalize_header(ws.cell(header_row, col_idx).value)
        if not header:
            continue

        if "原著小说是否在中台能查找到" in header or ("中台" in header and "查找" in header):
            header_map["mid_found"] = col_idx
        elif "原著小说名称" in header:
            header_map["original_novel_name"] = col_idx
        elif "是否有原著小说" in header:
            header_map["has_original"] = col_idx
        elif "短剧名称" in header or header == "剧名" or "短剧名" in header:
            header_map["drama_name"] = col_idx
        elif "原著小说cp名称" in header or "原著小说cp名" in header:
            header_map["original_cp_name"] = col_idx

    return header_map


def read_sheet_rows(
    workbook_path: Path,
) -> tuple[list[SourceRow], list[dict[str, Any]], dict[str, Counter], dict[str, dict[str, Any]], int]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows: list[SourceRow] = []
        issues: list[dict[str, Any]] = []
        sheet_meta: dict[str, dict[str, Any]] = {}
        stats = {
            "header_rows": Counter(),
            "has_original_values": Counter(),
            "mid_found_values": Counter(),
        }

        total_sheet_count = len(wb.sheetnames)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row = detect_header_row(ws)
            sheet_meta[sheet_name] = {
                "sheet_name": sheet_name,
                "header_row": header_row or "",
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "rows_extracted": 0,
                "eligible_rows": 0,
            }

            if header_row is None:
                issues.append(
                    {
                        "sheet_name": sheet_name,
                        "issue_type": "unmatched_header",
                        "detail": f"rows={ws.max_row}, cols={ws.max_column}",
                    }
                )
                continue

            stats["header_rows"][header_row] += 1
            header_map = build_header_map(ws, header_row)
            sheet_meta[sheet_name]["header_row"] = header_row
            sheet_meta[sheet_name]["header_map"] = str(header_map)

            if "original_novel_name" not in header_map or "mid_found" not in header_map:
                issues.append(
                    {
                        "sheet_name": sheet_name,
                        "issue_type": "missing_required_columns",
                        "detail": str(header_map),
                    }
                )
                continue

            for row_idx in range(header_row + 1, ws.max_row + 1):
                drama_name = normalize_text(ws.cell(row_idx, header_map.get("drama_name", 0)).value)
                original_novel_name = normalize_text(ws.cell(row_idx, header_map["original_novel_name"]).value)
                has_original_raw = normalize_text(ws.cell(row_idx, header_map.get("has_original", 0)).value)
                mid_found_raw = normalize_text(ws.cell(row_idx, header_map["mid_found"]).value)
                original_cp_name = normalize_text(ws.cell(row_idx, header_map.get("original_cp_name", 0)).value)

                if not any([drama_name, original_novel_name, has_original_raw, mid_found_raw, original_cp_name]):
                    continue

                has_original_class = classify_has_original(has_original_raw)
                mid_found_class = classify_mid_found(mid_found_raw)
                stats["has_original_values"][has_original_raw or "<blank>"] += 1
                stats["mid_found_values"][mid_found_raw or "<blank>"] += 1

                row = SourceRow(
                    source_sheet=sheet_name,
                    source_row=row_idx,
                    drama_name=drama_name,
                    original_novel_name=original_novel_name,
                    has_original_raw=has_original_raw,
                    has_original_class=has_original_class,
                    mid_found_raw=mid_found_raw,
                    mid_found_class=mid_found_class,
                    original_cp_name=original_cp_name,
                    header_row=header_row,
                )
                rows.append(row)
                sheet_meta[sheet_name]["rows_extracted"] += 1

                if row.original_novel_name and row.mid_found_class == "yes":
                    sheet_meta[sheet_name]["eligible_rows"] += 1

                if not original_novel_name and (has_original_raw or mid_found_raw):
                    issues.append(
                        {
                            "sheet_name": sheet_name,
                            "issue_type": "missing_original_title",
                            "detail": f"row={row_idx}, has_original={has_original_raw}, mid_found={mid_found_raw}",
                        }
                    )
    finally:
        wb.close()

    return rows, issues, stats, sheet_meta, total_sheet_count


def require_current_column(index: dict[str, int], field: str) -> int:
    aliases = CURRENT_WORKBOOK_COLUMN_ALIASES[field]
    for alias in aliases:
        if alias in index:
            return index[alias]
    raise KeyError(f"current workbook missing required column: {field} / {aliases}")


def load_current_intersection(path: Path) -> tuple[list[dict[str, Any]], set[tuple[str, str, str]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [normalize_text(ws.cell(1, col_idx).value) for col_idx in range(1, ws.max_column + 1)]
        index = {header: col_idx for col_idx, header in enumerate(headers, start=1) if header}

        rows: list[dict[str, Any]] = []
        keys: set[tuple[str, str, str]] = set()

        for row_idx in range(2, ws.max_row + 1):
            source_sheet = normalize_text(ws.cell(row_idx, require_current_column(index, "source_sheet")).value)
            source_row = parse_row_id(ws.cell(row_idx, require_current_column(index, "source_row")).value)
            original_novel_name = normalize_text(
                ws.cell(row_idx, require_current_column(index, "original_novel_name")).value
            )

            row = {
                "source_sheet": source_sheet,
                "source_row": source_row,
                "original_novel_name": original_novel_name,
                "dataset_key": normalize_text(ws.cell(row_idx, require_current_column(index, "dataset_key")).value),
                "book_ext_id": normalize_text(ws.cell(row_idx, require_current_column(index, "book_ext_id")).value),
                "book_name": normalize_text(ws.cell(row_idx, require_current_column(index, "book_name")).value),
                "short_drama_name": normalize_text(
                    ws.cell(row_idx, require_current_column(index, "short_drama_name")).value
                ),
            }
            rows.append(row)
            keys.add((source_sheet, source_row, loose_title_key(original_novel_name)))
    finally:
        wb.close()

    return rows, keys


def load_db_books(db_path: Path) -> list[DbBook]:
    conn = sqlite3.connect(db_path)
    try:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT dataset_key, book_ext_id, book_name
            FROM books
            WHERE book_name IS NOT NULL
              AND TRIM(book_name) <> ''
            """
        ).fetchall()
        return [
            DbBook(
                dataset_key=normalize_text(row["dataset_key"]),
                book_ext_id=normalize_text(row["book_ext_id"]),
                book_name=normalize_text(row["book_name"]),
            )
            for row in rows
        ]
    finally:
        conn.close()


def build_book_indexes(books: list[DbBook]) -> tuple[dict[str, list[DbBook]], dict[str, list[DbBook]]]:
    strict_index: dict[str, list[DbBook]] = defaultdict(list)
    loose_index: dict[str, list[DbBook]] = defaultdict(list)

    for book in books:
        strict_index[strict_title_key(book.book_name)].append(book)
        loose_index[loose_title_key(book.book_name)].append(book)

    return strict_index, loose_index


def autosize(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), 60)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def excel_safe_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def write_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["empty"])
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe_value(row.get(header, "")) for header in headers])

    header_fill = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
    header_font = Font(name="Arial", bold=True, color="1F3A5F")
    body_font = Font(name="Arial")
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap_alignment
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap_alignment

    ws.freeze_panes = "A2"
    autosize(ws)


def build_summary_rows(
    total_sheet_count: int,
    source_rows: list[SourceRow],
    eligible_rows: list[SourceRow],
    rebuilt_matches: list[dict[str, Any]],
    suspected_missing: list[dict[str, Any]],
    current_rows: list[dict[str, Any]],
    eligible_unmatched: list[dict[str, Any]],
    review_rows: list[dict[str, Any]],
    current_only: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    books: list[DbBook],
    stats: dict[str, Counter],
) -> list[dict[str, Any]]:
    return [
        {"metric": "线索表总 sheet 数", "value": total_sheet_count, "note": ""},
        {"metric": "识别到表头的 sheet 数", "value": sum(stats["header_rows"].values()), "note": dict(stats["header_rows"])},
        {"metric": "线索表抽取到的有效行数", "value": len(source_rows), "note": ""},
        {"metric": "满足中台可查=正例的候选行数", "value": len(eligible_rows), "note": ""},
        {"metric": "当前交集结果表行数", "value": len(current_rows), "note": ""},
        {"metric": "数据库 books 表标题数", "value": len(books), "note": ""},
        {"metric": "重建后匹配到数据库的行数", "value": len(rebuilt_matches), "note": ""},
        {"metric": "疑似漏交集行数", "value": len(suspected_missing), "note": ""},
        {"metric": "中台可查=正例但未匹配数据库", "value": len(eligible_unmatched), "note": ""},
        {"metric": "原著存在但需要人工复核的行数", "value": len(review_rows), "note": ""},
        {"metric": "当前结果表独有行数", "value": len(current_only), "note": ""},
        {"metric": "结构/表头/缺字段问题数", "value": len(issues), "note": ""},
        {"metric": "has_original 原始值分布", "value": "", "note": dict(stats["has_original_values"].most_common(30))},
        {"metric": "mid_found 原始值分布", "value": "", "note": dict(stats["mid_found_values"].most_common(30))},
    ]


def build_sheet_overview(
    sheet_meta: dict[str, dict[str, Any]],
    issues: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    issue_counter = Counter(issue["sheet_name"] for issue in issues)
    rows: list[dict[str, Any]] = []

    for sheet_name, meta in sheet_meta.items():
        rows.append(
            {
                "sheet_name": sheet_name,
                "header_row": meta.get("header_row", ""),
                "max_row": meta.get("max_row", ""),
                "max_column": meta.get("max_column", ""),
                "rows_extracted": meta.get("rows_extracted", 0),
                "eligible_rows": meta.get("eligible_rows", 0),
                "issue_count": issue_counter.get(sheet_name, 0),
                "header_map": meta.get("header_map", ""),
            }
        )

    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--clue-workbook", required=True)
    parser.add_argument("--current-workbook", required=True)
    parser.add_argument("--db-path", default="data/novel_similarity_v2.sqlite3")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    clue_path = Path(args.clue_workbook)
    current_path = Path(args.current_workbook)
    db_path = Path(args.db_path)
    output_path = Path(args.output)

    source_rows, issues, stats, sheet_meta, total_sheet_count = read_sheet_rows(clue_path)
    current_rows, current_keys = load_current_intersection(current_path)
    books = load_db_books(db_path)
    strict_index, loose_index = build_book_indexes(books)

    eligible_rows = [row for row in source_rows if row.original_novel_name and row.mid_found_class == "yes"]
    review_rows: list[dict[str, Any]] = []
    rebuilt_matches: list[dict[str, Any]] = []
    suspected_missing: list[dict[str, Any]] = []
    eligible_unmatched: list[dict[str, Any]] = []
    rebuilt_keys: set[tuple[str, str, str]] = set()

    for row in source_rows:
        if row.original_novel_name and row.mid_found_class != "yes":
            review_rows.append(
                {
                    "source_sheet": row.source_sheet,
                    "source_row": row.source_row,
                    "drama_name": row.drama_name,
                    "original_novel_name": row.original_novel_name,
                    "has_original_raw": row.has_original_raw,
                    "has_original_class": row.has_original_class,
                    "mid_found_raw": row.mid_found_raw,
                    "mid_found_class": row.mid_found_class,
                    "reason": "原著标题存在，但中台可查字段不是明确正例",
                }
            )

    for row in eligible_rows:
        exact_matches = strict_index.get(strict_title_key(row.original_novel_name), [])
        loose_matches: list[DbBook] = []
        match_type = "exact"
        if not exact_matches:
            loose_matches = loose_index.get(loose_title_key(row.original_novel_name), [])
            match_type = "loose"
        matches = exact_matches or loose_matches

        if not matches:
            eligible_unmatched.append(
                {
                    "source_sheet": row.source_sheet,
                    "source_row": row.source_row,
                    "drama_name": row.drama_name,
                    "original_novel_name": row.original_novel_name,
                    "has_original_raw": row.has_original_raw,
                    "mid_found_raw": row.mid_found_raw,
                    "reason": "线索表明确标记中台可查，但 books 表未命中该标题",
                }
            )
            continue

        source_key = (row.source_sheet, str(row.source_row), loose_title_key(row.original_novel_name))
        rebuilt_keys.add(source_key)

        for book in matches:
            rebuilt_row = {
                "source_sheet": row.source_sheet,
                "source_row": row.source_row,
                "drama_name": row.drama_name,
                "original_novel_name": row.original_novel_name,
                "has_original_raw": row.has_original_raw,
                "has_original_class": row.has_original_class,
                "mid_found_raw": row.mid_found_raw,
                "mid_found_class": row.mid_found_class,
                "match_type": match_type,
                "match_count": len(matches),
                "dataset_key": book.dataset_key,
                "book_ext_id": book.book_ext_id,
                "book_name": book.book_name,
                "in_current_workbook": "Y" if source_key in current_keys else "N",
            }
            rebuilt_matches.append(rebuilt_row)
            if source_key not in current_keys:
                suspected_missing.append(rebuilt_row)

    current_only: list[dict[str, Any]] = []
    for row in current_rows:
        key = (row["source_sheet"], row["source_row"], loose_title_key(row["original_novel_name"]))
        if key not in rebuilt_keys:
            current_only.append(
                {
                    **row,
                    "reason": "当前交集结果中存在，但本次按线索表规则未重建出来",
                }
            )

    summary_rows = build_summary_rows(
        total_sheet_count=total_sheet_count,
        source_rows=source_rows,
        eligible_rows=eligible_rows,
        rebuilt_matches=rebuilt_matches,
        suspected_missing=suspected_missing,
        current_rows=current_rows,
        eligible_unmatched=eligible_unmatched,
        review_rows=review_rows,
        current_only=current_only,
        issues=issues,
        books=books,
        stats=stats,
    )
    sheet_overview_rows = build_sheet_overview(sheet_meta, issues)

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)
    write_sheet(wb, "summary", summary_rows)
    write_sheet(wb, "rebuilt_matches", rebuilt_matches)
    write_sheet(wb, "suspected_missing", suspected_missing)
    write_sheet(wb, "eligible_unmatched", eligible_unmatched)
    write_sheet(wb, "review_rows", review_rows)
    write_sheet(wb, "current_only", current_only)
    write_sheet(wb, "sheet_issues", issues)
    write_sheet(wb, "sheet_overview", sheet_overview_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(output_path)
    print(f"total_sheet_count={total_sheet_count}")
    print(f"recognized_sheets={sum(stats['header_rows'].values())}")
    print(f"source_rows={len(source_rows)}")
    print(f"eligible_rows={len(eligible_rows)}")
    print(f"rebuilt_matches={len(rebuilt_matches)}")
    print(f"suspected_missing={len(suspected_missing)}")
    print(f"eligible_unmatched={len(eligible_unmatched)}")
    print(f"review_rows={len(review_rows)}")
    print(f"current_only={len(current_only)}")
    print(f"issues={len(issues)}")


if __name__ == "__main__":
    main()
