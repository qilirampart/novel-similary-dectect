# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT_DIR / "docs" / "39_项目介绍_技术栈与功能说明.docx"


TITLE = "小说文本相似度比对系统项目介绍"
SUBTITLE = "技术栈、核心方法与功能能力说明"


TECH_STACK_ROWS = [
    ("前端工作台", "React 18、TypeScript、Vite、React Router DOM、原生 CSS", "承载单条检测、批量任务、结果复核、任务记录和系统状态页面"),
    ("服务层", "Python、FastAPI", "提供单条比对、任务创建、结果查询、复核回写、导出下载和系统健康检查接口"),
    ("词面检索层", "SQLite、FTS5、自研 3-gram seed 检索与粗排策略", "负责在多来源正文索引中做高效率候选召回"),
    ("语义检索层", "Qdrant、Gitee Embedding API、Qwen3-Embedding-8B", "负责章节级和正文切块级语义召回，补足改写检测场景"),
    ("任务与结果沉淀", "SQLite 业务库、后台 worker、CSV/JSON 导出", "负责任务队列、状态追踪、结果落库、复核记录和交付物输出"),
]


SECTIONS = [
    (
        "项目概述",
        [
            "这套系统面向的是一个很具体也很难做干净的问题：给一段待检测文本，在大体量小说正文库里找出最可能对应的作品、章节和证据片段，并把结果组织成可以复核、可以追踪、可以导出的结构。它不是单纯的全文搜索，也不是一个只吐相似度分数的黑盒模型，而是一条从召回、比对到证据输出都可解释的工程链路。",
            "从现有实现看，项目已经具备 Web 化的完整闭环。前台可以做单条即时检测和批量任务提交，后台可以自动消费队列、执行检索和精排、沉淀结果并导出复核材料。真正有价值的地方，不是把几种技术堆在一起，而是把“找得到、比得准、看得懂、复得了”这四件事接成了一套稳定流程。",
        ],
    ),
    (
        "技术架构",
        [
            "整体架构采用前后端分离的轻量部署方案。前端使用 React 18、TypeScript 和 Vite 构建，页面路由由 React Router DOM 承担，重点不是做一个花哨的展示层，而是把比对结果、证据窗口、任务状态和系统健康信息清晰地摆到操作台上。当前 Web 工作台已经覆盖总览看板、单条比对、批量任务、结果复核、任务记录和系统状态等核心页面。",
            "后端以 FastAPI 为主，直接承接文本检测、任务管理、结果查询、复核回写和导出下载等接口。业务侧的在线任务与复核数据使用独立的 SQLite 业务库保存，并开启 WAL、busy_timeout 等设置来保证本地并发读写的稳定性。检索侧则使用另一套 SQLite 数据库承载书目、章节、正文、证据窗口和全文检索索引，在线服务与离线建库链路分开，职责边界比较清楚。",
            "语义层采用 Qdrant 作为向量检索底座，当前接入的是基于 Gitee Embedding API 的 Qwen3-Embedding-8B 向量化链路，向量集合分为章节级和正文切块级两层。这个设计很实用：章节级召回负责兜大范围语义相似，切块级召回负责把真正有证据价值的局部正文拉出来，后续再交给细粒度比对模块做精排。",
        ],
    ),
    (
        "核心方法",
        [
            "项目在方法上没有把问题粗暴地交给单一模型，而是拆成了词面召回、语义召回、候选融合和窗口级精排四个层次。这样做不是为了显得复杂，而是因为这类任务天然同时要求召回范围、证据精度和结果可解释性，少了任何一层，最后的体验都会明显掉档。",
        ],
    ),
]


METHOD_SUBSECTIONS = [
    (
        "词面召回",
        [
            "词面召回不是传统的整句关键词检索，而是先对查询文本做规范化处理，再抽取有顺序去重的中文 3-gram 片段，并利用 FTS5 词表里的文档频率挑出更稀缺、更有区分度的 seed term。系统先用这些 seed term 在多个全文索引中快速拉出候选池，再结合 seed hit 数、seed hit 权重和 n-gram overlap 做粗排。这样做的好处是速度快，而且对中文短句、台词类文本和轻度改写文本都比简单分词更稳。",
        ],
    ),
    (
        "语义召回",
        [
            "改写检测这一路并不会把文本一次性塞进向量库里碰运气。系统会先把长 query 按窗口切成多个语义块，再分别做 embedding，然后同时查询章节级集合和正文 chunk 集合。每个 query chunk 的命中结果都会保留下来，最后按命中次数、最佳 chunk 分数、章节桥接分数和覆盖到的 query chunk 数做聚合。这一步解决的是“字面不像，但剧情表达靠得很近”的问题，同时也避免长文本只靠第一段向量决定结果。",
        ],
    ),
    (
        "候选融合与窗口级精排",
        [
            "词面候选和语义候选不会简单拼接。系统会先处理两个候选池的交集，再对各自独有结果做交替补位，尽量兼顾精确文本命中和语义近邻。进入精排阶段后，候选章节会被切成固定大小的证据窗口，查询文本也会做同样的窗口化处理，随后逐窗口计算最长公共片段比例、n-gram recall、n-gram precision、Jaccard、SequenceMatcher ratio 以及 exact substring hit 等指标，最后得到 fine score、confidence label 和 review label。这个阶段本质上是在做可解释的文本对齐，所以系统给出的不是一个抽象结论，而是能直接拿去复核的命中片段。",
        ],
    ),
    (
        "检测口径设计",
        [
            "项目目前内置两条检测口径。面向直接复用的口径更强调精确文本证据，适合查近重复、直接摘抄和局部拼接；面向改写风险的口径则保留前者的判断能力，同时打开语义召回和可疑结果输出结构，用来筛查剧情保持但表达方式发生变化的文本。这个设计比较务实，因为实际业务里“完全复用”和“改写复用”往往不是两套系统，而是同一条链路上不同强度的判断。",
        ],
    ),
]


TAIL_SECTIONS = [
    (
        "功能能力",
        [
            "从功能上看，单条检测已经是一条完整的在线链路。用户输入一段文本后，系统会返回 top 命中的作品、章节、命中窗口、证据强度和复核所需的核心指标，前端可以直接联动展示候选列表和证据详情。对使用者来说，最重要的不是看见一个分数，而是能马上知道为什么命中这本书、为什么落在这一章、证据具体落在哪一段。",
            "批量任务部分更偏生产化。系统支持上传 txt、csv、tsv 和 xlsx 文件，能够自动识别文本字段与来源字段，把每一条文案拆成独立 task item 进入队列。后台 worker 会顺序消费待处理任务，持续回写心跳、单项完成状态和失败原因。任务跑完之后，系统会自动生成 summary CSV、review CSV 和 JSON 摘要，同时把 top1 结果、语义状态、复核状态和完整 payload 一起入库。这样一来，批量检测、人工复核和结果追踪就成了一套闭环，而不是几段孤立脚本。",
            "结果复核和系统管理这两块也已经成型。复核侧可以按任务、状态和复核结果筛选记录，并对单条结果回写 reviewer、review status 和备注；运维侧可以直接看到检索库规模、semantic chunk 规模、运行中任务、回退到 lexical only 的结果数量以及最近任务异常。换句话说，这个项目现在不只是“能跑”，而是已经具备了作为内部工作台长期使用的基础形态。",
        ],
    ),
    (
        "工程特点与可扩展性",
        [
            "这个项目真正难的地方，从来不是把模型接口接上，而是把离线建库、在线召回、窗口级证据、人工复核和任务状态放进同一条稳定链路里。当前实现比较可取的一点，是把离线索引构建和在线比对服务彻底拆开了：离线脚本负责清洗、切块、建索引和向量回填，在线 API 只负责接收请求、调度任务和组织结果，这让系统在维护时清楚得多。",
            "它的扩展空间也比较明确。检索侧现在已经支持多索引并行召回，后面可以继续扩展更多数据域；语义侧当前使用远程 embedding 加本地 Qdrant，后续可以平滑切换到本地 embedding 或其他模型；业务侧虽然当前使用 SQLite 作为轻量运行底座，但任务、结果和复核表结构已经拆开，未来迁移到 PostgreSQL 也不会重做整条业务链。这个版本没有为了追求“架构感”做过度设计，但把后续扩展最值钱的接口都留出来了。",
        ],
    ),
    (
        "结语",
        [
            "如果只用一句话概括，这个项目更像一套可解释的文本取证系统，而不只是一个小说搜索页面。它把全文检索、向量召回、窗口级对齐、证据评分、任务编排和人工复核接成了一条能落地的工程链路。就当前完成度来看，核心价值已经很清楚：面对一段待检测文本，系统不仅能给出答案，还能把答案背后的证据一起交出来。",
        ],
    ),
]


def set_run_font(run, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.start_type = WD_SECTION.NEW_PAGE

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True


def add_title(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run(TITLE)
    set_run_font(title_run, size=22, bold=True, color=RGBColor(30, 30, 30))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run(SUBTITLE)
    set_run_font(subtitle_run, size=11, color=RGBColor(96, 96, 96))


def add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(text)


def add_section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_paragraph(heading, style="Heading 1")
    for text in paragraphs:
        add_paragraph(doc, text)


def add_method_subsection(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_paragraph(heading, style="Heading 2")
    for text in paragraphs:
        add_paragraph(doc, text)


def add_tech_table(doc: Document) -> None:
    lead = doc.add_paragraph(style="Normal")
    lead.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead.add_run("当前版本的主干技术栈可以概括为下面这张表。它不算花哨，但很扎实，每一层都和实际职责对得上。")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("层级", "技术栈", "承担的职责")
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=10, bold=True)

    for row_values in TECH_STACK_ROWS:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_font(run, size=10)


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    add_title(doc)

    for heading, paragraphs in SECTIONS[:2]:
        add_section(doc, heading, paragraphs)

    add_tech_table(doc)

    add_section(doc, SECTIONS[2][0], SECTIONS[2][1])
    for heading, paragraphs in METHOD_SUBSECTIONS:
        add_method_subsection(doc, heading, paragraphs)

    for heading, paragraphs in TAIL_SECTIONS:
        add_section(doc, heading, paragraphs)

    return doc


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = "OpenAI Codex"
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
