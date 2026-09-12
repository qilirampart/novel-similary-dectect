from __future__ import annotations

import argparse
from dataclasses import asdict
from datetime import datetime
import json
from pathlib import Path
import sys
import time
from typing import Any

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from service.cover_monitor.collector import YouTubeChannelCollector
from service.cover_monitor.downloader import YouTubeCoverDownloader
from service.cover_monitor.reviewer import CoverVisionProfile, CoverVisionReviewer


def _first_channel_url(path: Path) -> tuple[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        urls = []
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                text = str(value or "").strip()
                if "youtube.com/" in text:
                    urls.append(text)
                    break
        if not urls:
            raise ValueError("工作簿中没有 YouTube 频道链接")
        return urls[0], len(urls)
    finally:
        workbook.close()


def _vision_profile_profile(config_path: Path) -> dict[str, Any]:
    payload = json.loads(config_path.read_text(encoding="utf-8"))
    profiles = payload.get("llm", {}).get("profiles", [])
    cover = payload.get("cover_review") or payload.get("video_review") or {}
    profile_id = str(cover.get("profile_id") or "").strip()
    ready = [
        item for item in profiles
        if isinstance(item, dict)
        and item.get("enabled", True)
        and item.get("api_base")
        and item.get("api_key")
        and item.get("model")
    ]
    selected = next((item for item in ready if str(item.get("id")) == profile_id), None)
    if selected is None and ready:
        selected = ready[0]
    if selected is None:
        raise ValueError("没有可用的视觉模型配置")
    return selected


def _md_cell(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run a ten-cover real network sample")
    parser.add_argument("--input", required=True)
    parser.add_argument("--assistant-runtime", required=True)
    parser.add_argument("--limit", type=int, default=10)
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    assistant_runtime = Path(args.assistant_runtime).resolve()
    limit = min(max(int(args.limit), 1), 10)
    channel_url, workbook_channel_count = _first_channel_url(input_path)
    proxy_url = (assistant_runtime / "youtube_proxy.txt").read_text(encoding="utf-8").strip()
    cookie_path = assistant_runtime / "youtube_cookies.txt"
    profile = _vision_profile_profile(assistant_runtime / "api_config.json")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = ROOT_DIR / "docs" / "cover_real_test_20260912"
    asset_root = ROOT_DIR / "runtime" / "cover_monitor" / "real_tests" / timestamp / "assets"
    output_dir.mkdir(parents=True, exist_ok=True)

    collector = YouTubeChannelCollector(
        proxy_url=proxy_url,
        cookie_path=cookie_path,
        timeout_seconds=30,
        max_attempts=2,
    )
    downloader = YouTubeCoverDownloader(asset_root, proxy_url=proxy_url, timeout_seconds=30)
    reviewer = CoverVisionReviewer(
        CoverVisionProfile(
            api_base=str(profile["api_base"]),
            api_key=str(profile["api_key"]),
            model=str(profile["model"]),
            timeout_seconds=90,
        )
    )

    total_started = time.perf_counter()
    collection_started = time.perf_counter()
    collection = collector.collect(
        channel_url,
        max_items_per_scope=limit,
        include_shorts=False,
    )
    collection_seconds = time.perf_counter() - collection_started
    videos = list(collection.videos)[:limit]
    results: list[dict[str, Any]] = []
    for index, video in enumerate(videos, start=1):
        item: dict[str, Any] = {
            "index": index,
            "video": asdict(video),
            "download_status": "failed",
            "review_status": "not_started",
        }
        download_started = time.perf_counter()
        try:
            downloaded = downloader.download(video.video_id, video.thumbnail_url)
            item["download_seconds"] = round(time.perf_counter() - download_started, 3)
            item["download_status"] = "succeeded"
            item["asset"] = asdict(downloaded)
        except Exception as exc:
            item["download_seconds"] = round(time.perf_counter() - download_started, 3)
            item["download_error"] = f"{type(exc).__name__}: {str(exc)[:500]}"
            results.append(item)
            print(f"[{index}/{len(videos)}] download failed {video.video_id}", flush=True)
            continue

        outcome = reviewer.review(
            image_path=item["asset"]["local_path"],
            video_title=video.title,
            intensity="standard",
        )
        item["review_status"] = outcome.status
        item["review_seconds"] = round(outcome.duration_seconds, 3)
        item["review_error_type"] = outcome.error_type
        item["review_error"] = outcome.error_message
        if outcome.result is not None:
            result = outcome.result.to_dict()
            result.pop("raw_response", None)
            item["detection"] = result
        item["total_item_seconds"] = round(
            float(item["download_seconds"]) + float(item["review_seconds"]), 3
        )
        results.append(item)
        print(
            f"[{index}/{len(videos)}] {video.video_id} download={item['download_seconds']}s "
            f"review={item['review_seconds']}s status={item['review_status']}",
            flush=True,
        )

    total_seconds = time.perf_counter() - total_started
    report = {
        "tested_at": datetime.now().isoformat(timespec="seconds"),
        "source_file": str(input_path),
        "source_channel_count": workbook_channel_count,
        "selected_channel_url": channel_url,
        "channel_id": collection.channel_id,
        "channel_name": collection.channel_name,
        "collection_seconds": round(collection_seconds, 3),
        "collection_completeness": collection.completeness,
        "scope_item_counts": collection.scope_item_counts,
        "scope_errors": collection.scope_errors,
        "requested_video_count": limit,
        "collected_video_count": len(videos),
        "download_success_count": sum(item["download_status"] == "succeeded" for item in results),
        "review_success_count": sum(item["review_status"] == "succeeded" for item in results),
        "total_seconds": round(total_seconds, 3),
        "vision_model": str(profile["model"]),
        "asset_root": str(asset_root.resolve()),
        "items": results,
    }
    json_path = output_dir / f"封面真实样本测试_{timestamp}.json"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "# 封面采集、下载与检测真实样本测试",
        "",
        f"- 测试时间：{report['tested_at']}",
        f"- 来源文件：`{input_path}`，共识别 {workbook_channel_count} 条频道链接",
        f"- 抽样频道：{collection.channel_name or '未返回频道名'}（`{collection.channel_id or '未返回频道ID'}`）",
        f"- 频道采集耗时：`{collection_seconds:.3f}` 秒，完整性：`{collection.completeness}`",
        f"- 封面下载：`{report['download_success_count']}/{len(videos)}` 成功",
        f"- 视觉检测：`{report['review_success_count']}/{len(videos)}` 成功，模型：`{profile['model']}`",
        f"- 端到端总耗时：`{total_seconds:.3f}` 秒",
        f"- 隔离封面目录：`{asset_root.resolve()}`",
        "",
        "| 序号 | 视频 ID | 视频标题 | 下载耗时(s) | 检测耗时(s) | 总耗时(s) | 结论 | 置信度 | 摘要/错误 |",
        "|---:|---|---|---:|---:|---:|---|---:|---|",
    ]
    for item in results:
        detection = item.get("detection") or {}
        summary = detection.get("summary") or item.get("review_error") or item.get("download_error") or ""
        lines.append(
            "| " + " | ".join(_md_cell(value) for value in (
                item["index"], item["video"]["video_id"], item["video"]["title"],
                item.get("download_seconds", ""), item.get("review_seconds", ""),
                item.get("total_item_seconds", ""), detection.get("overall_risk") or item["review_status"],
                detection.get("confidence", ""), summary,
            )) + " |"
        )
    md_path = output_dir / f"封面真实样本测试_{timestamp}.md"
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    # ASCII escaping keeps unattended Windows runs from failing on channel names outside GBK.
    print(json.dumps({"report": str(md_path), "raw": str(json_path), **{key: report[key] for key in ("channel_name", "collection_seconds", "download_success_count", "review_success_count", "total_seconds", "asset_root")}}, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()
