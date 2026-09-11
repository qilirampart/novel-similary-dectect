from __future__ import annotations

import csv
import gc
import json
import sqlite3
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))
SUBTITLE_DB = ROOT_DIR / "data" / "drama_subtitle_similarity_v1.sqlite3"

from service.business_store import init_business_db, list_users
from service.drama_subtitle_export import build_drama_subtitle_review_export_xlsx
from service.drama_subtitle_task_executor import execute_claimed_drama_subtitle_task
from service.drama_subtitle_task_store import (
    create_drama_subtitle_task,
    get_drama_subtitle_task,
    list_drama_subtitle_task_items,
    upsert_drama_subtitle_task_review,
)


def main() -> None:
    with sqlite3.connect(SUBTITLE_DB) as source_conn:
        source_conn.row_factory = sqlite3.Row
        source_rows = source_conn.execute(
            """
            SELECT book_id, book_name, episode_order, time_start, time_end, window_text
              FROM drama_subtitle_windows
             WHERE length(window_text) BETWEEN 320 AND 500
             ORDER BY book_id, episode_order, line_start
             LIMIT 3
            """
        ).fetchall()
    if len(source_rows) != 3:
        raise RuntimeError("not enough real subtitle windows for smoke test")

    with tempfile.TemporaryDirectory(prefix="drama-subtitle-flow-") as temp_dir:
        temp_root = Path(temp_dir)
        input_path = temp_root / "youtube_segments.csv"
        with input_path.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.DictWriter(
                output,
                fieldnames=[
                    "source_ref", "source_video_id", "source_channel", "source_upload_date",
                    "source_caption_language", "source_caption_source", "source_segment_order",
                    "source_time_start", "source_time_end", "source_display_title", "query_text",
                    "source_text_original",
                ],
            )
            writer.writeheader()
            for index, row in enumerate(source_rows, start=1):
                writer.writerow({
                    "source_ref": f"local-test-video#segment-{index}",
                    "source_video_id": "local-test-video",
                    "source_channel": "existing subtitle fixture",
                    "source_upload_date": "2026-08-03",
                    "source_caption_language": "en",
                    "source_caption_source": "fixture",
                    "source_segment_order": index,
                    "source_time_start": row["time_start"],
                    "source_time_end": row["time_end"],
                    "source_display_title": row["book_name"],
                    "query_text": row["window_text"],
                    "source_text_original": row["window_text"],
                })

        business_db = temp_root / "business.sqlite3"
        init_business_db(business_db)
        owner = list_users(business_db)[0]
        task_id = "existing-subtitle-flow-smoke"
        create_drama_subtitle_task(
            db_path=business_db,
            task_id=task_id,
            source_file_name=input_path.name,
            source_file_ext=input_path.suffix,
            source_file_path=str(input_path),
            source_file_sha256="fixture",
            source_file_size=input_path.stat().st_size,
            owner_user_id=int(owner["user_id"]),
            created_by="smoke-test",
            params={"top_k": 5, "window_limit": 50},
        )
        result = execute_claimed_drama_subtitle_task(
            task_id=task_id,
            business_db_path=business_db,
            subtitle_db_path=SUBTITLE_DB,
        )
        items = list_drama_subtitle_task_items(business_db, task_id)
        if result.get("status") != "completed" or len(items) != 3:
            raise AssertionError({"task": result, "items": items})
        if not all(item["status"] == "completed" and item["matched_book_id"] for item in items):
            raise AssertionError("one or more real subtitle windows were not recalled")

        reviewed = upsert_drama_subtitle_task_review(
            db_path=business_db,
            task_item_id=int(items[0]["task_item_id"]),
            review_status="confirmed_high_risk",
            reviewer_name="smoke-test",
            review_note="real subtitle fixture flow",
        )
        if reviewed is None:
            raise AssertionError("review update failed")
        export_path, _ = build_drama_subtitle_review_export_xlsx(
            business_db_path=business_db,
            export_root=temp_root / "exports",
            owner_user_id=int(owner["user_id"]),
            task_id=task_id,
        )
        workbook = load_workbook(export_path, read_only=True, data_only=True)
        try:
            first_sheet = workbook[workbook.sheetnames[0]]
            headers = list(next(first_sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            if "source_video_id" not in headers or "source_time_start" not in headers:
                raise AssertionError("provenance fields missing from review export")
        finally:
            workbook.close()
            del workbook
            gc.collect()

        summary = {
            "task_id": task_id,
            "status": get_drama_subtitle_task(business_db, task_id)["status"],
            "input_count": len(items),
            "matched_book_ids": [item["matched_book_id"] for item in items],
            "durations_seconds": [item["duration_seconds"] for item in items],
            "source_video_id": items[0]["source_video_id"],
            "review_export": str(export_path),
        }
        print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
