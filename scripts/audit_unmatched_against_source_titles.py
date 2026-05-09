from __future__ import annotations

import argparse
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    return text.strip()


def strict_title_key(value: Any) -> str:
    text = normalize_text(value).lower()
    text = re.sub(r"\s+", "", text)
    return re.sub(r"[《》「」『』“”‘’\"'`]", "", text)


def loose_title_key(value: Any) -> str:
    text = strict_title_key(value)
    return re.sub(r"[^\u4e00-\u9fff\w]", "", text)


def excel_safe_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def autosize(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), 60)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_sheet(ws) -> None:
    header_fill = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
    header_font = Font(name="Arial", bold=True, color="1F3A5F")
    body_font = Font(name="Arial")
    wrap = Alignment(vertical="top", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    autosize(ws)


def write_rows_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["empty"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe_value(row.get(header, "")) for header in headers])
    style_sheet(ws)


def load_unmatched_titles(path: Path) -> list[dict[str, str]]:
    df = pd.read_excel(path, sheet_name=0, header=6)
    df = df.fillna("")
    rows: list[dict[str, str]] = []
    for record in df.to_dict(orient="records"):
        row = {str(k): normalize_text(v) for k, v in record.items()}
        if any(row.values()):
            rows.append(row)
    return rows


def load_source_titles(path: Path, source_label: str) -> list[dict[str, str]]:
    df = pd.read_excel(path, sheet_name=0, usecols=[0])
    title_col = df.columns[0]
    titles = df[title_col].fillna("").astype(str).map(normalize_text)
    unique_titles = sorted({title for title in titles if title})
    return [{"source_dataset": source_label, "source_title": title} for title in unique_titles]


def load_books(db_path: Path) -> list[dict[str, str]]:
    conn = sqlite3.connect(db_path)
    try:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT dataset_key, book_ext_id, book_name
            FROM books
            WHERE book_name IS NOT NULL
              AND TRIM(book_name) <> ''
            """
        ).fetchall()
        return [
            {
                "dataset_key": normalize_text(row["dataset_key"]),
                "book_ext_id": normalize_text(row["book_ext_id"]),
                "book_name": normalize_text(row["book_name"]),
            }
            for row in rows
        ]
    finally:
        conn.close()


def build_title_indexes(
    source_titles: list[dict[str, str]],
    books: list[dict[str, str]],
) -> dict[str, Any]:
    source_strict: dict[str, list[dict[str, str]]] = defaultdict(list)
    source_loose: dict[str, list[dict[str, str]]] = defaultdict(list)
    book_strict: dict[str, list[dict[str, str]]] = defaultdict(list)
    book_loose: dict[str, list[dict[str, str]]] = defaultdict(list)

    for item in source_titles:
        source_strict[strict_title_key(item["source_title"])].append(item)
        source_loose[loose_title_key(item["source_title"])].append(item)

    for book in books:
        book_strict[strict_title_key(book["book_name"])].append(book)
        book_loose[loose_title_key(book["book_name"])].append(book)

    return {
        "source_strict": source_strict,
        "source_loose": source_loose,
        "book_strict": book_strict,
        "book_loose": book_loose,
    }


def match_source_title(title: str, indexes: dict[str, Any]) -> tuple[str, list[dict[str, str]], str]:
    strict_matches = indexes["source_strict"].get(strict_title_key(title), [])
    if strict_matches:
        return "source_exact", strict_matches, "strict_equal"

    loose_matches = indexes["source_loose"].get(loose_title_key(title), [])
    if loose_matches:
        return "source_normalized", loose_matches, "loose_equal"

    return "source_missing", [], "none"


def match_imported_book(title: str, indexes: dict[str, Any]) -> tuple[str, list[dict[str, str]], str]:
    strict_matches = indexes["book_strict"].get(strict_title_key(title), [])
    if strict_matches:
        return "imported_exact", strict_matches, "strict_equal"

    loose_matches = indexes["book_loose"].get(loose_title_key(title), [])
    if loose_matches:
        return "imported_normalized", loose_matches, "loose_equal"

    return "imported_missing", [], "none"


def classify_row(row: dict[str, str], indexes: dict[str, Any]) -> dict[str, Any]:
    title = row.get("original_novel_name", "")
    source_status, source_matches, source_rule = match_source_title(title, indexes)
    imported_status, imported_matches, imported_rule = match_imported_book(title, indexes)

    if source_status != "source_missing" and imported_status != "imported_missing":
        conclusion = "源数据中有该书名，且已导入库中也有该书名"
        final_category = "in_source_and_imported"
    elif source_status != "source_missing" and imported_status == "imported_missing":
        conclusion = "源数据中有该书名，但当前已导入库中没有，属于高疑似导入遗漏"
        final_category = "in_source_but_not_imported"
    else:
        conclusion = "两份源数据书名中均未找到，当前更像是源数据本身未覆盖"
        final_category = "not_in_source"

    return {
        **row,
        "source_presence_status": source_status,
        "source_match_rule": source_rule,
        "source_match_count": len(source_matches),
        "source_datasets": sorted({item["source_dataset"] for item in source_matches}),
        "source_titles": [item["source_title"] for item in source_matches[:10]],
        "imported_presence_status": imported_status,
        "imported_match_rule": imported_rule,
        "imported_match_count": len(imported_matches),
        "imported_dataset_keys": sorted({item["dataset_key"] for item in imported_matches}),
        "imported_titles": [item["book_name"] for item in imported_matches[:10]],
        "final_category": final_category,
        "conclusion": conclusion,
    }


def build_summary_rows(
    source_title_count: int,
    books_count: int,
    results: list[dict[str, Any]],
    source_paths: list[Path],
    db_path: Path,
) -> list[dict[str, Any]]:
    counter = Counter(item["final_category"] for item in results)
    return [
        {"metric": "待复核标题数", "value": len(results), "note": ""},
        {"metric": "源表唯一书名数", "value": source_title_count, "note": "两份源表合并去重后"},
        {"metric": "已导入 books 表书名数", "value": books_count, "note": ""},
        {"metric": "源表和库里都有", "value": counter["in_source_and_imported"], "note": ""},
        {"metric": "源表有但库里没有", "value": counter["in_source_but_not_imported"], "note": "高疑似导入遗漏"},
        {"metric": "源表本身没有", "value": counter["not_in_source"], "note": ""},
        {"metric": "源表1", "value": str(source_paths[0]), "note": ""},
        {"metric": "源表2", "value": str(source_paths[1]), "note": ""},
        {"metric": "比对数据库", "value": str(db_path), "note": ""},
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--source-online", required=True)
    parser.add_argument("--source-self-short", required=True)
    parser.add_argument("--db-path", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    input_path = Path(args.input_workbook)
    source_online = Path(args.source_online)
    source_self_short = Path(args.source_self_short)
    db_path = Path(args.db_path)
    output_path = Path(args.output)

    unmatched_rows = load_unmatched_titles(input_path)
    source_titles = load_source_titles(source_online, "online_owned_source") + load_source_titles(
        source_self_short, "self_short_source"
    )
    books = load_books(db_path)
    indexes = build_title_indexes(source_titles, books)

    results = [classify_row(row, indexes) for row in unmatched_rows]
    in_source_and_imported = [row for row in results if row["final_category"] == "in_source_and_imported"]
    in_source_but_not_imported = [row for row in results if row["final_category"] == "in_source_but_not_imported"]
    not_in_source = [row for row in results if row["final_category"] == "not_in_source"]
    summary_rows = build_summary_rows(
        source_title_count=len({item["source_title"] for item in source_titles}),
        books_count=len(books),
        results=results,
        source_paths=[source_online, source_self_short],
        db_path=db_path,
    )

    wb = Workbook()
    wb.remove(wb.active)
    write_rows_sheet(wb, "summary", summary_rows)
    write_rows_sheet(wb, "all_results", results)
    write_rows_sheet(wb, "in_source_and_imported", in_source_and_imported)
    write_rows_sheet(wb, "in_source_but_not_imported", in_source_but_not_imported)
    write_rows_sheet(wb, "not_in_source", not_in_source)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    counter = Counter(item["final_category"] for item in results)
    print(output_path)
    print(f"input_rows={len(results)}")
    print(f"in_source_and_imported={counter['in_source_and_imported']}")
    print(f"in_source_but_not_imported={counter['in_source_but_not_imported']}")
    print(f"not_in_source={counter['not_in_source']}")


if __name__ == "__main__":
    main()
