from __future__ import annotations

import argparse
import http.cookiejar
import json
import mimetypes
import time
import uuid
from pathlib import Path
from typing import Any
from urllib import error, parse, request

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


TERMINAL_STATUSES = {"completed", "partial_failed", "failed", "cancelled", "paused"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rerun the original-ID-in-library no-match subset through the cloud API.")
    parser.add_argument("--base-url", default="http://novel-similarity-dev.dzkjm.cn")
    parser.add_argument("--username", required=True)
    parser.add_argument("--password", required=True)
    parser.add_argument("--file", required=True)
    parser.add_argument("--task-id", default="", help="Reuse an existing cloud task instead of creating a new one.")
    parser.add_argument("--out-dir", default=r"E:\点众\YouTube字幕核验助手工作区\output\结果比对文件夹")
    parser.add_argument("--poll-seconds", type=float, default=3.0)
    parser.add_argument("--timeout-seconds", type=float, default=3600.0)
    return parser.parse_args()


def build_opener() -> request.OpenerDirector:
    return request.build_opener(request.HTTPCookieProcessor(http.cookiejar.CookieJar()))


def json_request(
    opener: request.OpenerDirector,
    url: str,
    *,
    method: str = "GET",
    payload: dict[str, Any] | None = None,
    timeout: float = 60.0,
) -> dict[str, Any]:
    data = None
    headers = {"Accept": "application/json"}
    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = request.Request(url, data=data, headers=headers, method=method)
    try:
        with opener.open(req, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} {url}: {body[:1000]}") from exc


def multipart_task_request(
    opener: request.OpenerDirector,
    url: str,
    file_path: Path,
    *,
    fields: dict[str, str],
) -> dict[str, Any]:
    boundary = f"----CodexBoundary{uuid.uuid4().hex}"
    body = bytearray()
    for name, value in fields.items():
        body.extend(f"--{boundary}\r\n".encode("utf-8"))
        body.extend(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode("utf-8"))
        body.extend(value.encode("utf-8"))
        body.extend(b"\r\n")
    mime_type = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
    body.extend(f"--{boundary}\r\n".encode("utf-8"))
    body.extend(
        (
            f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'
            f"Content-Type: {mime_type}\r\n\r\n"
        ).encode("utf-8")
    )
    body.extend(file_path.read_bytes())
    body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode("utf-8"))
    req = request.Request(
        url,
        data=bytes(body),
        headers={
            "Accept": "application/json",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
        method="POST",
    )
    try:
        with opener.open(req, timeout=180.0) as response:
            return json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        body_text = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} {url}: {body_text[:1000]}") from exc


def load_input_rows(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(iterator)]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if any(value is not None and str(value) != "" for value in values):
            rows.append(dict(zip(headers, values)))
    return rows


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def compact_json(value: Any) -> str:
    if value in (None, ""):
        return ""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def payload_from_item(item: dict[str, Any]) -> dict[str, Any]:
    raw = item.get("result_payload_json")
    if isinstance(raw, dict):
        return raw
    try:
        parsed = json.loads(text(raw) or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def first_candidate(payload: dict[str, Any]) -> dict[str, Any]:
    candidates = payload.get("candidates")
    if isinstance(candidates, list) and candidates and isinstance(candidates[0], dict):
        return candidates[0]
    fine = payload.get("fine")
    fine_results = fine.get("results") if isinstance(fine, dict) else None
    if isinstance(fine_results, list) and fine_results and isinstance(fine_results[0], dict):
        item = fine_results[0]
        best_match = item.get("best_match")
        if isinstance(best_match, dict):
            merged = dict(item)
            merged["evidence"] = best_match.get("evidence") or best_match
            merged["candidate_text"] = best_match.get("candidate_text") or best_match.get("candidate_text_preview")
            return merged
        return item
    results = payload.get("results")
    if isinstance(results, list) and results and isinstance(results[0], dict):
        return results[0]
    return {}


def candidate_value(candidate: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in candidate and candidate[key] not in (None, ""):
            return candidate[key]
    return ""


def candidate_evidence_text(candidate: dict[str, Any]) -> str:
    evidence = candidate.get("evidence")
    if isinstance(evidence, dict):
        return text(evidence.get("window_text") or evidence.get("text") or evidence.get("candidate_text"))
    return text(evidence or candidate.get("candidate_text") or candidate.get("candidate_text_preview"))


def fetch_task_detail(opener: request.OpenerDirector, base_url: str, task_id: str) -> dict[str, Any]:
    url = f"{base_url}/api/v1/drama-subtitles/tasks/{parse.quote(task_id)}"
    return json_request(opener, url, timeout=90.0)


def write_output(
    output_path: Path,
    input_rows: list[dict[str, Any]],
    detail: dict[str, Any],
) -> int:
    task = detail.get("task") if isinstance(detail.get("task"), dict) else {}
    cloud_items = detail.get("items") if isinstance(detail.get("items"), list) else []
    by_order = {int(item.get("item_order")): item for item in cloud_items if text(item.get("item_order")).isdigit()}
    headers = [
        "序号", "结果源行号", "视频 ID", "来源频道", "来源剧名", "原版剧ID", "原版剧名称",
        "云端任务项ID", "云端状态", "最终判定", "判定原因", "用户结论", "耗时(秒)", "语义状态",
        "第一候选剧名", "第一候选Book ID", "第一候选集数", "第一候选字幕语言", "第一候选语义分数",
        "第一候选精排分数", "第一候选复核标签", "第一候选召回来源", "第一候选证据文本",
        "第一候选完整JSON", "原始字幕", "译文", "翻译回退", "服务端执行", "错误信息",
    ]
    rows: list[list[Any]] = []
    for order, source in enumerate(input_rows, start=1):
        item = by_order.get(order, {})
        payload = payload_from_item(item)
        decision = payload.get("decision") if isinstance(payload.get("decision"), dict) else {}
        candidate = first_candidate(payload)
        evidence = candidate_evidence_text(candidate)
        rows.append([
            order,
            source.get("结果源行号", ""),
            source.get("视频 ID", ""),
            source.get("来源频道", ""),
            source.get("来源剧名", ""),
            source.get("基准原版剧ID", ""),
            source.get("基准原版剧名称", ""),
            item.get("task_item_id", ""),
            item.get("status", ""),
            decision.get("outcome") or decision.get("status") or item.get("status", ""),
            decision.get("reason", ""),
            decision.get("user_message", ""),
            item.get("duration_seconds", ""),
            item.get("semantic_status", ""),
            candidate_value(candidate, "book_name", "matched_book_name"),
            candidate_value(candidate, "book_id", "matched_book_id", "book_ext_id"),
            candidate_value(candidate, "episode_order", "matched_episode_order"),
            candidate_value(candidate, "language_code"),
            candidate_value(candidate, "semantic_score"),
            candidate_value(candidate, "fine_score"),
            candidate_value(candidate, "review_label", "confidence_label"),
            compact_json(candidate.get("retrieval_sources")),
            evidence,
            compact_json(candidate),
            source.get("完整字幕", ""),
            source.get("译文", ""),
            source.get("翻译回退", ""),
            source.get("服务端执行", ""),
            item.get("error_message", ""),
        ])
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "云端重跑结果"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    text_headers = {
        "来源频道", "来源剧名", "原版剧名称", "最终判定", "判定原因", "用户结论",
        "第一候选剧名", "第一候选证据文本", "第一候选完整JSON", "原始字幕", "译文",
        "错误信息",
    }
    for row in worksheet.iter_rows(min_row=2):
        for index, cell in enumerate(row):
            cell.alignment = Alignment(
                horizontal="left" if headers[index] in text_headers else "center",
                vertical="top",
                wrap_text=headers[index] in text_headers,
            )
    widths = {
        "来源剧名": 28, "原版剧名称": 28, "原始字幕": 60, "译文": 60,
        "判定原因": 34, "用户结论": 42, "第一候选剧名": 32, "第一候选证据文本": 70,
        "第一候选完整JSON": 90, "错误信息": 35,
    }
    for index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = widths.get(header, 18)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    if rows:
        table = Table(displayName="CloudRerunFirstCandidate", ref=f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        worksheet.add_table(table)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(rows)


def main() -> int:
    args = parse_args()
    base_url = args.base_url.rstrip("/")
    input_path = Path(args.file).resolve()
    out_dir = Path(args.out_dir).resolve()
    input_rows = load_input_rows(input_path)
    if len(input_rows) != 54:
        raise RuntimeError(f"expected 54 input rows, got {len(input_rows)}: {input_path}")
    opener = build_opener()
    login = json_request(
        opener,
        f"{base_url}/api/v1/auth/login",
        method="POST",
        payload={"username": args.username, "password": args.password},
        timeout=30.0,
    )
    task_id = text(args.task_id)
    if not task_id:
        task_create = multipart_task_request(
            opener,
            f"{base_url}/api/v1/drama-subtitles/tasks",
            input_path,
            fields={
                "top_k": "10",
                "window_limit": "200",
                "semantic_enabled": "true",
                "semantic_window_limit": "100",
                "translation_fallback": "true",
            },
        )
        task_id = text(task_create.get("task_id"))
        if not task_id:
            raise RuntimeError(f"cloud task creation did not return task_id: {task_create}")
    print(json.dumps({"login_user": login.get("user", {}).get("username"), "task_id": task_id, "input_rows": len(input_rows)}, ensure_ascii=False), flush=True)
    started = time.monotonic()
    last_detail: dict[str, Any] = {}
    poll_log: list[dict[str, Any]] = []
    while True:
        last_detail = fetch_task_detail(opener, base_url, task_id)
        task = last_detail.get("task") if isinstance(last_detail.get("task"), dict) else {}
        status = text(task.get("status"))
        counts = task.get("counts") if isinstance(task.get("counts"), dict) else {}
        snapshot = {"elapsed_seconds": round(time.monotonic() - started, 3), "status": status, "counts": counts, "item_count": len(last_detail.get("items") or [])}
        poll_log.append(snapshot)
        print(json.dumps(snapshot, ensure_ascii=False), flush=True)
        if status in TERMINAL_STATUSES:
            break
        if time.monotonic() - started > args.timeout_seconds:
            raise TimeoutError(f"cloud task timed out after {args.timeout_seconds}s: {task_id}")
        time.sleep(args.poll_seconds)
    safe_id = task_id.replace("-", "")
    out_dir.mkdir(parents=True, exist_ok=True)
    raw_path = out_dir / f"54条云端重跑_{safe_id}_raw.json"
    raw_path.write_text(json.dumps({"task": last_detail.get("task"), "items": last_detail.get("items", []), "poll_log": poll_log}, ensure_ascii=False, indent=2), encoding="utf-8")
    output_path = out_dir / "54条原版ID在库未命中_云端重跑_第一候选.xlsx"
    row_count = write_output(output_path, input_rows, last_detail)
    summary = {
        "task_id": task_id,
        "status": (last_detail.get("task") or {}).get("status"),
        "counts": (last_detail.get("task") or {}).get("counts"),
        "item_count": len(last_detail.get("items") or []),
        "output_rows": row_count,
        "raw_path": str(raw_path),
        "output_path": str(output_path),
    }
    summary_path = out_dir / "54条原版ID在库未命中_云端重跑_第一候选_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
