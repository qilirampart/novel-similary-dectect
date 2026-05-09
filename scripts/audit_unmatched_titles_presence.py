from __future__ import annotations

import argparse
import json
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_DATE = "2026-05-08"


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    return text.strip()


def strict_title_key(value: Any) -> str:
    text = normalize_text(value).lower()
    text = re.sub(r"\s+", "", text)
    return re.sub(r"[《》「」『』“”‘’\"'`]", "", text)


def loose_title_key(value: Any) -> str:
    text = strict_title_key(value)
    return re.sub(r"[^\u4e00-\u9fff\w]", "", text)


def ultra_loose_title_key(value: Any) -> str:
    text = loose_title_key(value)
    text = text.lower()
    text = re.sub(r"\d+", "", text)
    noise_tokens = (
        "全文",
        "完整版",
        "全本",
        "小说",
        "原著",
        "短剧",
        "漫剧",
    )
    for token in noise_tokens:
        text = text.replace(token, "")
    return text


def excel_safe_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def autosize(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), 60)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_sheet(ws, freeze_panes: str | None = "A2") -> None:
    header_fill = PatternFill(fill_type="solid", start_color="DCEBFA", end_color="DCEBFA")
    header_font = Font(name="Arial", bold=True, color="1F3A5F")
    body_font = Font(name="Arial")
    wrap = Alignment(vertical="top", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap

    if freeze_panes:
        ws.freeze_panes = freeze_panes
    autosize(ws)


def write_rows_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["empty"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([excel_safe_value(row.get(header, "")) for header in headers])
    style_sheet(ws)


def load_unmatched_rows(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        header_row = 7
        headers = [normalize_text(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)]
        rows: list[dict[str, Any]] = []
        for r in range(header_row + 1, ws.max_row + 1):
            row = {headers[c - 1]: normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
            if any(row.values()):
                rows.append(row)
        return rows
    finally:
        wb.close()


def load_books(db_path: Path) -> list[dict[str, str]]:
    conn = sqlite3.connect(db_path)
    try:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """
            SELECT dataset_key, book_ext_id, book_name
            FROM books
            WHERE book_name IS NOT NULL
              AND TRIM(book_name) <> ''
            ORDER BY book_name
            """
        ).fetchall()
        return [
            {
                "dataset_key": normalize_text(row["dataset_key"]),
                "book_ext_id": normalize_text(row["book_ext_id"]),
                "book_name": normalize_text(row["book_name"]),
            }
            for row in rows
        ]
    finally:
        conn.close()


def build_indexes(books: list[dict[str, str]]) -> dict[str, Any]:
    strict_index: dict[str, list[dict[str, str]]] = defaultdict(list)
    loose_index: dict[str, list[dict[str, str]]] = defaultdict(list)
    ultra_loose_index: dict[str, list[dict[str, str]]] = defaultdict(list)

    for book in books:
        strict_index[strict_title_key(book["book_name"])].append(book)
        loose_index[loose_title_key(book["book_name"])].append(book)
        ultra_loose_index[ultra_loose_title_key(book["book_name"])].append(book)

    return {
        "strict": strict_index,
        "loose": loose_index,
        "ultra_loose": ultra_loose_index,
        "books": books,
    }


def find_contains_candidates(target: str, books: list[dict[str, str]], limit: int = 10) -> list[dict[str, str]]:
    normalized_target = ultra_loose_title_key(target)
    if not normalized_target:
        return []

    candidates: list[dict[str, str]] = []
    for book in books:
        book_key = ultra_loose_title_key(book["book_name"])
        if not book_key:
            continue
        if normalized_target in book_key or book_key in normalized_target:
            candidates.append(book)
            if len(candidates) >= limit:
                break
    return candidates


def classify_row(
    row: dict[str, str],
    indexes: dict[str, Any],
) -> dict[str, Any]:
    title = row.get("original_novel_name", "")
    strict_matches = indexes["strict"].get(strict_title_key(title), [])
    if strict_matches:
        return {
            **row,
            "presence_status": "exact_present",
            "match_rule": "strict_equal",
            "match_count": len(strict_matches),
            "matched_titles": [book["book_name"] for book in strict_matches[:10]],
            "matched_book_ext_ids": [book["book_ext_id"] for book in strict_matches[:10]],
            "conclusion": "书名已明确存在于已导入 books 表",
        }

    loose_matches = indexes["loose"].get(loose_title_key(title), [])
    if loose_matches:
        return {
            **row,
            "presence_status": "normalized_present",
            "match_rule": "loose_equal",
            "match_count": len(loose_matches),
            "matched_titles": [book["book_name"] for book in loose_matches[:10]],
            "matched_book_ext_ids": [book["book_ext_id"] for book in loose_matches[:10]],
            "conclusion": "书名经标点/空白归一化后存在于已导入 books 表",
        }

    ultra_loose_matches = indexes["ultra_loose"].get(ultra_loose_title_key(title), [])
    if ultra_loose_matches:
        return {
            **row,
            "presence_status": "ultra_normalized_present",
            "match_rule": "ultra_loose_equal",
            "match_count": len(ultra_loose_matches),
            "matched_titles": [book["book_name"] for book in ultra_loose_matches[:10]],
            "matched_book_ext_ids": [book["book_ext_id"] for book in ultra_loose_matches[:10]],
            "conclusion": "书名经更宽松去噪后存在于已导入 books 表",
        }

    contains_candidates = find_contains_candidates(title, indexes["books"])
    if contains_candidates:
        return {
            **row,
            "presence_status": "possible_variant",
            "match_rule": "contains_or_contained",
            "match_count": len(contains_candidates),
            "matched_titles": [book["book_name"] for book in contains_candidates[:10]],
            "matched_book_ext_ids": [book["book_ext_id"] for book in contains_candidates[:10]],
            "conclusion": "库内存在疑似别名/副标题变体，需人工复核",
        }

    return {
        **row,
        "presence_status": "not_found_in_imported_books",
        "match_rule": "none",
        "match_count": 0,
        "matched_titles": [],
        "matched_book_ext_ids": [],
        "conclusion": "按当前 books 表及宽松标题规则，未发现该书名已导入证据",
    }


def build_summary_rows(results: list[dict[str, Any]], db_path: Path, input_path: Path) -> list[dict[str, Any]]:
    counter = Counter(row["presence_status"] for row in results)
    return [
        {"metric": "统计日期", "value": SUMMARY_DATE, "note": ""},
        {"metric": "输入文件", "value": str(input_path), "note": ""},
        {"metric": "比对数据库", "value": str(db_path), "note": ""},
        {"metric": "待复核标题数", "value": len(results), "note": ""},
        {"metric": "明确已在库中", "value": counter["exact_present"], "note": "strict_equal"},
        {"metric": "归一化后可确认在库中", "value": counter["normalized_present"], "note": "loose_equal"},
        {"metric": "宽松去噪后可确认在库中", "value": counter["ultra_normalized_present"], "note": "ultra_loose_equal"},
        {"metric": "疑似别名/副标题变体", "value": counter["possible_variant"], "note": "contains_or_contained"},
        {"metric": "当前仍未找到已导入证据", "value": counter["not_found_in_imported_books"], "note": "none"},
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--db-path", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    input_path = Path(args.input_workbook)
    db_path = Path(args.db_path)
    output_path = Path(args.output)

    unmatched_rows = load_unmatched_rows(input_path)
    books = load_books(db_path)
    indexes = build_indexes(books)

    results = [classify_row(row, indexes) for row in unmatched_rows]

    exact_rows = [row for row in results if row["presence_status"] == "exact_present"]
    normalized_rows = [row for row in results if row["presence_status"] == "normalized_present"]
    ultra_rows = [row for row in results if row["presence_status"] == "ultra_normalized_present"]
    variant_rows = [row for row in results if row["presence_status"] == "possible_variant"]
    not_found_rows = [row for row in results if row["presence_status"] == "not_found_in_imported_books"]
    summary_rows = build_summary_rows(results, db_path, input_path)

    wb = Workbook()
    wb.remove(wb.active)
    write_rows_sheet(wb, "summary", summary_rows)
    write_rows_sheet(wb, "all_results", results)
    write_rows_sheet(wb, "exact_present", exact_rows)
    write_rows_sheet(wb, "normalized_present", normalized_rows)
    write_rows_sheet(wb, "ultra_normalized_present", ultra_rows)
    write_rows_sheet(wb, "possible_variant", variant_rows)
    write_rows_sheet(wb, "not_found", not_found_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    counter = Counter(row["presence_status"] for row in results)
    print(output_path)
    print(f"input_rows={len(results)}")
    for key in (
        "exact_present",
        "normalized_present",
        "ultra_normalized_present",
        "possible_variant",
        "not_found_in_imported_books",
    ):
        print(f"{key}={counter[key]}")


if __name__ == "__main__":
    main()
